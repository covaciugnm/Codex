import pathlib,zipfile,json,xml.etree.ElementTree as E,shutil,copy
import openpyxl
work=pathlib.Path(__file__).parent
meta=json.loads((work/'comparison_meta.json').read_text(encoding='utf-8'))
dest=pathlib.Path(r'Z:\00. Proiecte 2025\AFIR FotoVoltaic 1\50. Fabrica Nuci\nuttechnology.com\07_Centralizatoare\Echipamente_si_preturi_NutTechnology_2026-10-16.xlsx')
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',N);E.register_namespace('r',R)
with zipfile.ZipFile(work/'comparison_export.xlsx') as z: files={n:z.read(n) for n in z.namelist()}
wb=E.fromstring(files['xl/workbook.xml']);sheets=wb.find('{'+N+'}sheets')
byname={s.attrib['name']:s for s in sheets}
for s in list(sheets):sheets.remove(s)
for n in meta['order']:sheets.append(byname[n])
for v in wb.findall('.//{'+N+'}workbookView'):v.set('activeTab','0');v.set('firstSheet','0')
files['xl/workbook.xml']=E.tostring(wb,encoding='utf-8',xml_declaration=True)
rels=E.fromstring(files['xl/_rels/workbook.xml.rels'])
targets={r.attrib['Id']:r.attrib['Target'] for r in rels}
for n in meta['order']:
 target=targets[byname[n].attrib['{'+R+'}id']]
 target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
 sheet=E.fromstring(files[target])
 for v in sheet.findall('.//{'+N+'}sheetView'):
  v.set('tabSelected','1' if n==meta['order'][0] else '0')
 # Simple filters support sorting the requested data area without adding duplicate table headers.
 if n=='Matrice echipamente':filter_range='A7:Z'+str(7+len(meta['models']))
 elif n=='Comparativ preturi':filter_range='A6:G20'
 elif n in [o['name'] for o in meta['offers']]:
  o=next(o for o in meta['offers'] if o['name']==n);filter_range='A7:N'+str(7+len(o['rows']))
 else:filter_range=None
 if filter_range:
  for old in sheet.findall('{'+N+'}autoFilter'):sheet.remove(old)
  idx=list(sheet).index(sheet.find('{'+N+'}sheetData'))+1
  sheet.insert(idx,E.Element('{'+N+'}autoFilter',{'ref':filter_range}))
 files[target]=E.tostring(sheet,encoding='utf-8',xml_declaration=True)
# Restore local hyperlink target spelling changed by the import/export normalizer.
for rp in list(files):
 if rp.endswith('.rels'):
  rr=E.fromstring(files[rp]);changed=False
  for rel in rr:
   t=rel.attrib.get('Target','')
   if t.startswith('file:///Z:'):
    rel.set('Target',t[8:]);changed=True
  if changed:files[rp]=E.tostring(rr,encoding='utf-8',xml_declaration=True)
# Sheet removal/reordering is performed at package level; content is authored and calculated with Artifact Tool.
candidate=work/'comparison_final.xlsx'
with zipfile.ZipFile(candidate,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in files.items():z.writestr(name,data)
a=openpyxl.load_workbook(work/'before_edit.xlsx',data_only=False)
b=openpyxl.load_workbook(candidate,data_only=False)
assert b.sheetnames==meta['order']
for name in ['Catalog echipamente','Index fise','Conditii si diferente']:
 for row in a[name]:
  for cell in row:
   dst=b[name][cell.coordinate]
   assert cell.value==dst.value,(name,cell.coordinate,'value')
   assert (cell.hyperlink.target if cell.hyperlink else None)==(dst.hyperlink.target if dst.hyperlink else None),(name,cell.coordinate,'link')
   assert cell.number_format==dst.number_format,(name,cell.coordinate,'number format')
   assert copy.copy(cell.font)==copy.copy(dst.font),(name,cell.coordinate,'font')
   assert copy.copy(cell.fill)==copy.copy(dst.fill),(name,cell.coordinate,'fill')
 assert set(a[name].tables)==set(b[name].tables)
 assert a[name].freeze_panes==b[name].freeze_panes
v=openpyxl.load_workbook(candidate,data_only=True)
for name,total in zip([o['name'] for o in meta['offers']],[85100,111200,314400]):
 assert v[name]['B24'].value==total
 assert v[name]['B29'].value=='ND'
 assert b[name]['A2'].value.startswith('Nr.')
assert v['179 11.10.2026 - 400']['B26'].value==6000
for s in v:
 for row in s:
  for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
# All intersections reconcile to offer rows, including inactive alternatives.
for i,m in enumerate(meta['models']):
 rr=8+i
 for j,o in enumerate(meta['offers']):
  sr=o['rowMap'].get(str(m['id']))
  if sr is None:continue
  os=v[o['name']];ms=v['Matrice echipamente'];start=3+j*8
  expected=[os.cell(sr,3).value*os.cell(sr,4).value,os.cell(sr,5).value,os.cell(sr,6).value,os.cell(sr,7).value,os.cell(sr,8).value,os.cell(sr,11).value,os.cell(sr,12).value,os.cell(sr,9).value if os.cell(sr,4).value else 0]
  assert [ms.cell(rr,start+k).value for k in range(8)]==expected,(m['id'],o['name'])
shutil.copy2(candidate,dest)
print(json.dumps({'saved':str(dest),'sheets':b.sheetnames,'models':len(meta['models']),'totals':[85100,111200,314400],'preserved_tabs':3,'verified':'formulas, values, native links, formats, tables, panes and all matrix intersections'},ensure_ascii=False))

