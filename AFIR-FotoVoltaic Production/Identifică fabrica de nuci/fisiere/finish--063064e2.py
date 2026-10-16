import pathlib,json,zipfile,xml.etree.ElementTree as E,shutil
import openpyxl
work=pathlib.Path(__file__).parent
root=pathlib.Path(r'Z:\00. Proiecte 2025\AFIR FotoVoltaic 1\50. Fabrica Nuci\nuttechnology.com')
links=json.loads((work/'links.json').read_text(encoding='utf-8'))
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',N);E.register_namespace('r',R)
for p in (root/'07_Centralizatoare').glob('*.xlsx'):
 with zipfile.ZipFile(p) as z: files={n:z.read(n) for n in z.namelist()}
 wb=E.fromstring(files['xl/workbook.xml'])
 rel=E.fromstring(files['xl/_rels/workbook.xml.rels'])
 relmap={r.attrib['Id']:r.attrib['Target'] for r in rel}
 count=0
 for sh in wb.find('{'+N+'}sheets'):
  sl=[l for l in links if l['sheet']==sh.attrib['name']]
  if not sl:continue
  target=relmap[sh.attrib['{'+R+'}id']]
  target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
  xml=E.fromstring(files[target])
  for old in list(xml.findall('{'+N+'}hyperlinks')):xml.remove(old)
  hs=E.Element('{'+N+'}hyperlinks')
  relpath=str(pathlib.PurePosixPath(target).parent/'_rels'/(pathlib.PurePosixPath(target).name+'.rels'))
  rs=E.fromstring(files[relpath]) if relpath in files else E.Element('{'+P+'}Relationships')
  for old in list(rs):
   if old.attrib.get('Id','').startswith('nativeLink'):rs.remove(old)
  for i,l in enumerate(sl):
   dest=l['target']
   if not dest.startswith('http'): assert pathlib.Path(dest).exists(),dest
   rid='nativeLink'+str(i+1)
   E.SubElement(hs,'{'+N+'}hyperlink',{'ref':l['cell'],'{'+R+'}id':rid})
   E.SubElement(rs,'{'+P+'}Relationship',{'Id':rid,'Type':R+'/hyperlink','Target':dest,'TargetMode':'External'})
   count+=1
  # hyperlinks precede print options and page settings in worksheet schema.
  position=len(xml)
  after_tags={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
  for i,node in enumerate(xml):
   if node.tag.split('}')[-1] in after_tags:position=i;break
  xml.insert(position,hs)
  files[target]=E.tostring(xml,encoding='utf-8',xml_declaration=True)
  files[relpath]=E.tostring(rs,encoding='utf-8',xml_declaration=True)
 tmp=work/p.name
 with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in files.items():z.writestr(n,b)
 shutil.copy2(tmp,p)
 r=openpyxl.load_workbook(p,data_only=True)
 assert all(c.data_type!='e' for s in r for row in s for c in row)
 actual=sum(bool(c.hyperlink) for s in r for row in s for c in row)
 assert actual==count,(actual,count)
 if 'Oferta 179' in r:
  assert r['Oferta 179']['F5'].value==314400
  assert r['Oferta 179']['G5'].value==6000
  assert sum(r['Oferta 179'].cell(i,5).value for i in range(9,18))==320400
 else:
  assert r['Comunicatii']['A6'].value.year==2024
 print(p.name, 'sheets',len(r.sheetnames),'links',actual,'checks OK')
# Move generated inspection logs out of deliverable folder.
for p in (root/'07_Centralizatoare').glob('*.ndjson'):
 shutil.move(str(p),str(work/p.name))
# Preserve a manifest for all final files.
rows=[]
for p in sorted(root.rglob('*')):
 if p.is_file():rows.append({'file':str(p.relative_to(root)),'bytes':p.stat().st_size})
(root/'01_Comunicatii'/'Inventar_fisiere.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')

