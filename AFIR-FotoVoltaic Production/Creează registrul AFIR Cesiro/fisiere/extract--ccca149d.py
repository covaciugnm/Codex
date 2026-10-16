import pathlib,json,zipfile,hashlib,re,io
from pypdf import PdfReader
from docx import Document
root=pathlib.Path(r'Z:\00. Proiecte 2025\AFIR FotoVoltaic 1')
out=pathlib.Path(__file__).parent
skip={'0. Intrare cont AFIR','.codex','outputs','60. Poze Hala Artfil','Nicos'}
rows=[]; archives=[]; cache={}
for p in root.rglob('*'):
 if not p.is_file() or any(s in skip for s in p.relative_to(root).parts) or p.name.startswith('~$'): continue
 rel=str(p.relative_to(root)); ext=p.suffix.lower()
 if ext not in {'.pdf','.docx','.doc','.xlsx','.xls','.zip','.txt','.html','.csv','.jpg','.png','.jpeg','.msg','.eml'}:continue
 row={'path':str(p),'relative':rel,'name':p.name,'ext':ext,'size':p.stat().st_size,'text':'','pages':0,'read':''}
 if ext in {'.pdf','.docx','.txt'}:
  try:
   b=p.read_bytes(); h=hashlib.sha256(b).hexdigest();row['hash']=h
   if h in cache:row.update(cache[h])
   else:
    if ext=='.pdf':
     reader=PdfReader(io.BytesIO(b));row['pages']=len(reader.pages);row['text']='\n'.join(f'[Pagina {i+1}]\n'+(x.extract_text() or '') for i,x in enumerate(reader.pages))
    elif ext=='.docx':
     d=Document(io.BytesIO(b));row['text']='\n'.join([x.text for x in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])
    else:row['text']=b.decode('utf-8-sig',errors='replace')
    row['read']='Text extras' if len(row['text'].strip())>60 else 'Scan / text insuficient'
    cache[h]={k:row[k] for k in ['text','pages','read']}
  except Exception as e:row['read']='Eroare citire: '+str(e)[:100]
 elif ext=='.zip':
  try:
   with zipfile.ZipFile(p) as z:
    for a in z.infolist():
     if not a.is_dir():archives.append({'archive':rel,'entry':a.filename,'size':a.file_size})
   row['read']='Inventar ZIP'
  except Exception as e:row['read']=str(e)[:100]
 else: row['read']='Inventariat; conținut neextras'
 rows.append(row)
(out/'local.json').write_text(json.dumps(rows,ensure_ascii=False),encoding='utf-8')
(out/'archives.json').write_text(json.dumps(archives,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'files':len(rows),'archive_entries':len(archives),'read':sum(bool(r['text']) for r in rows)},ensure_ascii=False))
