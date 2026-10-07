import json,re,shutil,hashlib,html,zipfile,datetime,unicodedata
from collections import Counter,defaultdict
from pathlib import Path
from urllib.parse import quote
from bs4 import BeautifulSoup
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
from catalogue import topic_program

BASE=Path(__file__).parent;OUT=BASE/'F';CACHE=BASE/'research_cache'
TODAY=datetime.date(2026,10,1);DATE='01.10.2026'
NETWORK=Path(r'\\192.168.100.169\Comun\00. Proiecte 2026 EU\Finantari directe RO - 2026.10.01')
NAME='Registru_finantari_directe_RO_2026-10-01.xlsx'
cat=json.loads((BASE/'catalogue.json').read_text(encoding='utf8'))
manifest=json.loads((BASE/'manifest.json').read_text(encoding='utf8'))
portal=json.loads((BASE/'portal_all.json').read_text(encoding='utf8'))
byid={x['id']:x for x in cat}
def patch(id,**kwargs):byid[id].update(kwargs)
patch('CH04',who='IMM-uri eligibile din producție, comerț și servicii; inclusiv microîntreprinderi în condițiile procedurii actualizate',source_key='Swiss_ECO_Update',source='https://www.swiss-contribution.ro/web/swiss/noutati/-/asset_publisher/31kQDN4YO3VO/content/noutati_20260701',apply='https://sme-ecotech.gov.ro/',deadline='2026-09-24',status='Termen publicat expirat; eventuală prelungire neconfirmată',amount='Grant maximum 50.000 CHF / 267.240 lei; investiție minimum 668.100 lei',budget='54.041.886 CHF / 288.843.072 lei – întregul program 2025–2027',person='Nu',note='Ultimul termen publicat în actualizarea MF: 24.09.2026, ora 20:00. O prelungire suplimentară nu a fost confirmată. Procedura consolidată și ordinul din iulie 2026 sunt salvate. Creditul bancar de minimum 60% este obligatoriu; nu este grant direct de la Comisia Europeană.')
patch('EYF',source_key='EYF_2027',source='https://www.coe.int/en/web/kyiv/-/announcing-the-eyf-call-for-proposals-for-international-one-off-youth-co-operation-activities-and-international-long-term-youth-co-operation-projects-with-the-deadline-of-1-october-2026',deadline='2026-10-01',status='Termen ASTĂZI – 01.10.2026, 23:59 CET, conform anunțului',amount='2027.C5.A: maximum 25.000 EUR; 2027.C6.B: conform ghidului',apply='https://support4youth.coe.int/',note='Două apeluri pentru activități / proiecte internaționale de tineret. Înregistrarea și eligibilitatea ONG-ului se verifică în ghid. C5.A cere participanți din minimum 7 state membre; cel puțin 75% au maximum 30 de ani. Nu finanțează investiții comerciale. Anunț verificat prin browser; serverul a refuzat descărcarea celor două ghiduri (403).')
patch('ESC',source='https://youth.europa.eu/solidarity/organisations/documents-resources_en',note='Ghidurile oficiale 2026 RO și EN sunt salvate. Acțiunile centralizate se depun la EACEA; alte acțiuni sunt gestionate de agenții naționale. Tinerii participă sau formează un grup eligibil, conform acțiunii.')
patch('CME',source_key='CME_2026_Call',source='https://culture.ec.europa.eu/culture-moves-europe/call-for-individual-mobility',deadline='2026-10-31',status='Calendar 2026–2027 publicat; prima rundă 31.10.2026',amount='85 EUR / zi + transport 400 / 800 EUR + suplimente eligibile',budget='25 mil. EUR – întregul mecanism 2025–2028',note='Artiști / profesioniști culturali cu rezidență legală în țări Creative Europe, minimum 18 ani; persoane sau grupuri de până la 5. Proiect în altă țară cu partener internațional. Runde 31.10, 30.11.2026 și 31.01, 28.02, 31.03, 30.04.2027. Implementare Goethe-Institut; ghidul versiunii curente prevalează.')
patch('CH01',note=byid['CH01']['note']+' Apelul pentru pilotarea calificărilor nivel 5 s-a închis la 29.07.2026, ora 16:30; buget 3,5 mil. CHF, contribuție elvețiană 85% și cofinanțare 15%. Ghidul și anexele sunt în 0.ARHIVA.')
patch('CH07',note=byid['CH07']['note']+' Selecția a 16 consultanți ROeea s-a închis la 21.08.2026; este selecție / formare profesională, nu grant de investiții pentru persoane. Ghidul, anexele și rezultatele sunt în 0.ARHIVA.')
patch('CH09',note=byid['CH09']['note']+' Apelul pentru 4 centre publice CSMPTA a avut termen 30.03.2026; dotări de aproximativ 120.000 CHF / centru, conform MF. Nu este apel pentru firme private.')
for id in ['CSF2','CSF5','CSF6']:
 patch(id,amount='Granturi mici / mari în benzi distincte; extreme: 15.000 și 350.000 EUR. Vezi ghid.',note=byid[id]['note']+' Sumele nu formează un interval continuu: fiecare categorie are limite proprii în ghid.')
for r in cat:
 if r['id'].startswith('NA'):r['company']='Conform vechiului ghid';r['ngo']='Conform vechiului ghid';r['person']='Conform vechiului ghid';r['public']='Conform vechiului ghid'
def append_program(id,title,source,what,who,**kw):
 r={k:'' for k in ['source_key','budget','deadline','note','parent']}
 r.update(id=id,title=title,source=source,what=what,who=who,group='ALTE',company='Condiționat',ngo='Nu',person='Nu',public='Nu',mode='Direct Consiliul Europei',instrument='Sprijin financiar',amount='Conform regulilor',rate='Conform regulilor',partnership='Conform regulilor',status='Program activ; verificare la nivel de apel',apply=source)
 r.update(kw);r['folder']=('01. UE' if r['group']=='UE' else '04. Alte organisme')+'/PROGRAM - '+id+' - '+title[:28];cat.append(r);byid[id]=r
append_program('IF26SME','Innovation Fund IF26 SME','https://climate.ec.europa.eu/areas-action/eu-funding-climate-action/innovation-fund/calls-proposals/if26-sme-call_en','Proiecte inovatoare de decarbonizare ale IMM-urilor, CAPEX 2,5–20 mil. EUR','IMM-uri potrivit definiției UE; condiții finale în call fiche',group='UE',source_key='IF26_SME',parent='INNOV',mode='Direct Comisia Europeană / CINEA',instrument='Grant',company='Da',budget='200 mil. EUR – buget PLANIFICAT',status='Anunțat; publicare planificată decembrie 2026',note='Două runde planificate: publicare decembrie 2026 / termen sfârșit martie 2027; publicare aprilie 2027 / termen sfârșit septembrie 2027. Nu a fost publicată o zi-limită exactă. Prezentarea consultării este document de pregătire, nu ghid final.')
append_program('IF26NZT','Innovation Fund IF26 Net-Zero','https://climate.ec.europa.eu/areas-action/eu-funding-climate-action/innovation-fund/calls-proposals/if26-nzt-call_en','Decarbonizare, producție cleantech, materii prime critice și proiecte pilot','Entități cu proiecte eligibile în țările mecanismului; condiții finale în call fiche',group='UE',source_key='IF26_NZT',parent='INNOV',mode='Direct Comisia Europeană / CINEA',instrument='Grant',company='Condiționat',status='Anunțat; deschidere planificată decembrie 2026',note='Închidere planificată aprilie 2027; zi exactă nepublicată. Patru categorii: CAPEX peste 100 mil. EUR pentru large-scale; 2,5–100 mil. EUR pentru small / medium; peste 2,5 mil. EUR pentru cleantech / CRM și pilots. Proiectele eligibile SME / Maritime sunt excluse din NZT, cu excepția pilots. Nu confunda CAPEX cu suma grantului.')
append_program('IF26MAR','Innovation Fund IF26 Maritime','https://climate.ec.europa.eu/areas-action/eu-funding-climate-action/innovation-fund/calls-proposals/if26-maritime-call_en','Tehnologii maritime cu emisii scăzute / zero și infrastructură portuară inovatoare','Entități cu proiecte maritime eligibile, CAPEX minimum 2,5 mil. EUR',group='UE',source_key='IF26_Maritime',parent='INNOV',mode='Direct Comisia Europeană / CINEA',instrument='Grant',company='Condiționat',status='Anunțat; deschidere planificată decembrie 2026',note='Închidere planificată aprilie 2027; zi exactă nepublicată. CAPEX minimum 2,5 mil. EUR reprezintă dimensiunea investiției, nu grantul. Ghidul final / call fiche se verifică la lansare.')
append_program('EURIMAGES','Eurimages – coproducții cinematografice','https://www.coe.int/en/web/eurimages/co-production-funding','Coproducții de lungmetraj: ficțiune, animație și documentare','Companii de producție independente din statele membre Eurimages; România este membră',source_key='Eurimages',deadline='2027-01-13',status='Următorul termen publicat: 13.01.2027; lansarea sesiunii de verificat',instrument='Împrumut condiționat rambursabil / subvenție, conform regulilor',apply='https://eurimages-online.coe.int/WebForms/Default.aspx',note='Ultima sesiune 2026: 08.09.2026, închisă. Calendarul 2027: 13.01, 02.04, 10.09; ora 12:00, ora locală Franța. Proiectul și producătorii trebuie să îndeplinească regulile de coproducție; minimum 50% finanțare confirmată în fiecare țară coproductoare. Sursa calendar: https://www.coe.int/en/web/eurimages/deadlines')
# Additional official Eurimages regulations: save only a validated original PDF.
u='https://eurimages-online.coe.int/Templates/CoprodRegulations_2026_EN.pdf'
try:
 import requests,io
 from pypdf import PdfReader
 resp=requests.get(u,timeout=35,headers={'User-Agent':'Mozilla/5.0'});resp.raise_for_status();raw=resp.content
 assert raw.startswith(b'%PDF-');pages=len(PdfReader(io.BytesIO(raw)).pages);sha=hashlib.sha256(raw).hexdigest()
 folder=OUT/byid['EURIMAGES']['folder']/'1.DOCUMENTE OFICIALE';folder.mkdir(parents=True,exist_ok=True)
 dest=folder/('Eurimages_Regulations_2026_EN_'+sha[:8]+'.pdf');dest.write_bytes(raw)
 manifest.append({'program':'EURIMAGES','url':u,'source':byid['EURIMAGES']['source'],'label':'Eurimages Co-production Support Regulations 2026','archive':False,'ok':True,'bytes':len(raw),'pages':pages,'sha256':sha,'ext':'.pdf','relative_path':dest.relative_to(OUT).as_posix()})
except Exception as e:manifest.append({'program':'EURIMAGES','url':u,'source':byid['EURIMAGES']['source'],'label':'Eurimages Co-production Support Regulations 2026','archive':False,'ok':False,'error':str(e)[:250]})

# Explicitly archive previous work programmes; general grant rules remain current unless replaced.
for d in manifest:
 if d['ok'] and not d['archive'] and re.search(r'/wp-call/(2025|2023-2024|2023-2025)/',d['url']):
  old=OUT/d['relative_path']; dest=OUT/byid[d['program']]['folder']/'0.ARHIVA'/old.name
  dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(old,dest);d['archive']=True;d['relative_path']=dest.relative_to(OUT).as_posix()
  # Original download remains as a preserved source; mark it historical in the index too.
  d['label']='[ISTORIC] '+d['label']
for d in manifest:
 if d['program']=='EUIPO' and d['ok']:d['label']=d['label'].replace(' - URL de verificat','')

groups={'UE':'Uniunea Europeană','SEE':'SEE / Norvegia','CH':'Elveția','ALTE':'Alte organisme europene','ARHIVA':'SEE / Norvegia – arhivă 2014–2021'}
documents=defaultdict(list)
for d in manifest:documents[d['program']].append(d)
def first(m,k,default=''):
 v=m.get(k,[]);return str(v[0]) if isinstance(v,list) and v else str(v) if v else default
def js(m,k):
 try:return json.loads(first(m,k,'{}'))
 except:return {}
def plain(s,n=12000):return BeautifulSoup(s or '','html.parser').get_text(' ',strip=True)[:n]
def date(s):
 try:return datetime.date.fromisoformat(str(s)[:10])
 except:return None
def euro(x):
 try:return f'{float(x):,.0f}'.replace(',','.')+' EUR'
 except:return str(x)
def rel_link(path):return quote(str(path).replace('\\','/'),safe='/:')
topics=[]
for t in portal['results']:
 m=t['metadata'];ident=first(m,'identifier');parent=topic_program(ident);r=byid.get(parent,byid['EU01'])
 dates=sorted(set(x[:10] for x in m.get('deadlineDate',[]) if date(x)))
 nxt=next((x for x in dates if date(x)>=TODAY),'');opening=first(m,'startDate')[:10]
 code=first(m,'status');status={'31094501':'Anunțat / forthcoming','31094502':'Deschis în portal','31094503':'Închis în portal'}.get(code,'Statut de verificat')
 model=first(m,'deadlineModel');conditions=' '.join(m.get('topicConditions',[]));desc=' '.join(m.get('descriptionByte',[]))
 is2=(model=='two-stage' or 'two-stage' in (ident+' '+conditions).lower())
 note='Eligibilitatea României ca stat membru nu garantează eligibilitatea solicitantului. Verifică topicul, anexele, consorțiul și eventualele restricții de securitate / proprietate.'
 if is2 and dates and date(dates[0])<TODAY:
  status='Etapa inițială expirată – etapa următoare poate fi doar pe invitație';note+=' Apel în două etape; termenul viitor nu înseamnă că se pot depune proiecte noi.'
 if opening and date(opening) and date(opening)>TODAY:note+=' Deschidere publicată: '+opening+'.'
 grant='Grant / instrument conform topicului'
 if 'ECHE-CERT' in ident:status='Certificare ECHE – fără grant';grant='Certificare, nu finanțare';note='Această înregistrare apare în filtrul oficial de oportunități, dar nu acordă grant. Carta ECHE este o condiție instituțională pentru Erasmus+.'
 bud=js(m,'budgetOverview');entries=[]
 for vals in bud.get('budgetTopicActionMap',{}).values():
  for a in vals if isinstance(vals,list) else []:
   if ident in str(a.get('action','')):entries.append(a)
 budgets=[];contrib=[];expected=[]
 for a in entries:
  for y,b in (a.get('budgetYearMap') or {}).items():budgets.append(str(y)+': '+euro(b))
  lo=a.get('minContribution');hi=a.get('maxContribution')
  if lo or hi:contrib.append(euro(lo or hi)+((' – '+euro(hi)) if hi and hi!=lo else ''))
  if a.get('expectedGrants') is not None:expected.append(str(a['expectedGrants']))
 budget='; '.join(dict.fromkeys(budgets));fund='; '.join(dict.fromkeys(contrib)) or 'Vezi call fiche / work programme'
 url='https://ec.europa.eu/info/funding-tenders/opportunities/portal/screen/opportunities/topic-details/'+ident
 topics.append({'id':ident,'title':first(m,'title',t.get('summary','')),'parent':parent,'family':r['title'],'call':first(m,'callIdentifier'),'status':status,'status_code':code,'next':nxt,'dates':'; '.join(dates),'opening':opening,'model':model,'instrument':grant,'amount':fund,'budget':budget,'expected':'; '.join(expected),'who':r['who'],'company':'Posibil, verifică topicul' if r['company'] not in ['Nu',''] else 'Neconfirmat pentru firmă; verifică topicul','person':r['person'],'ngo':r['ngo'],'public':r['public'],'apply':url,'folder':r['folder'],'what':plain(desc,10000),'conditions':plain(conditions,14000),'note':note})
topics.sort(key=lambda x:(x['next'] or '9999',x['id']))
assert len(topics)==442 and len({x['id'] for x in topics})==442

# Keep every official topic snapshot in its correctly mapped programme folder.
for t in portal['results']:
 ident=first(t['metadata'],'identifier');r=byid[topic_program(ident)];folder=OUT/r['folder']/'2.APELURI PORTAL';folder.mkdir(parents=True,exist_ok=True)
 fn=re.sub(r'[^A-Za-z0-9_.-]','_',ident)+'.json';target=folder/fn
 target.write_text(json.dumps(t,ensure_ascii=False,indent=2),encoding='utf8')
 # Earlier snapshots are harmless duplicates; a single canonical copy is indexed here.

def eligible(flag):return bool(flag) and flag not in ['Nu','Beneficiar indirect','Conform vechiului ghid']
def route(r):
 mode=r['mode'].lower()
 if r['group']=='ARHIVA':return 'Arhivă'
 if r['instrument'].lower().startswith(('împrumut','credit')) or r['id'] in ['EIB','EIF','EBRD']:return 'Finanțare / asistență prin instituție financiară'
 if 'direct' in mode and 'operator' not in mode and r['group']!='CH':return 'Aplicare la organism european'
 return 'Operator / instituție / finanțare mixtă'
def closed(r):return r['group']=='ARHIVA' or ('închis' in r['status'].lower()) or ('expirat' in r['status'].lower())
def usable(r):return r['group']!='ARHIVA' and not closed(r)

WB=Workbook();WB.remove(WB.active)
navy='123B5D';blue='17658A';teal='087E8B';gold='F6B940';pale='EAF3F8';grey='687783'
table_index=0
def sheet(name,headers,rows,widths=None,links=None):
 global table_index
 ws=WB.create_sheet(name);ws.sheet_view.showGridLines=False
 ws.append([name+' | cercetare oficială la '+DATE]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
 ws.cell(1,1).font=Font(name='Calibri',size=18,bold=True,color='FFFFFF');ws.cell(1,1).fill=PatternFill('solid',fgColor=navy);ws.row_dimensions[1].height=34
 ws.append(['Înapoi la Start • filtrele sunt în antet • condițiile oficiale ale apelului prevalează']);ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers));ws.cell(2,1).hyperlink='#Start!A1';ws.cell(2,1).font=Font(size=10,color=blue,underline='single');ws.row_dimensions[2].height=23
 ws.append(headers)
 for record in rows:
  clean=[]
  for value in record:
   if isinstance(value,str):value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)[:32000]
   if isinstance(value,str) and value.startswith(('=','+','@')):value="'"+value
   clean.append(value)
  ws.append(clean)
 for row in ws.iter_rows(min_row=4):
  ws.row_dimensions[row[0].row].height=62 if name not in ['Documente','Surse','Lipsuri'] else 38
  for cell in row:cell.font=Font(name='Calibri',size=10,color='223648');cell.alignment=Alignment(vertical='top',wrap_text=True)
 for c in ws[3]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=blue);c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[3].height=34
 if rows:
  table_index+=1;tb=Table(displayName='T'+str(table_index),ref=f'A3:{get_column_letter(len(headers))}{ws.max_row}');tb.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tb)
 for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=(widths or {}).get(h,27)
 ws.freeze_panes='C4';ws.auto_filter.ref=f'A3:{get_column_letter(len(headers))}{ws.max_row}'
 if links:
  for j,kind in links.items():
   for rowno in range(4,ws.max_row+1):
    c=ws.cell(rowno,j);v=c.value
    if v:
     c.hyperlink=rel_link(v) if kind=='local' else v;c.font=Font(name='Calibri',size=10,color=blue,underline='single');c.alignment=Alignment(vertical='top',wrap_text=True)
 ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:3'
 return ws

start=WB.create_sheet('Start');start.sheet_view.showGridLines=False
start.merge_cells('A1:H2');start['A1']='FINANȚĂRI EUROPENE PENTRU ROMÂNIA';start['A1'].fill=PatternFill('solid',fgColor=navy);start['A1'].font=Font(size=23,bold=True,color='FFFFFF');start['A1'].alignment=Alignment(vertical='center');start.row_dimensions[1].height=30;start.row_dimensions[2].height=20
success=[d for d in manifest if d['ok']];failed=[d for d in manifest if not d['ok']]
stats={'programe_scheme':len(cat),'topicuri_portal':len(topics),'documente_originale_distincte':len({d['sha256'] for d in success}),'url_documente_reusite':len({d['url'] for d in success}),'url_documente_nereusite':len({d['url'] for d in failed}),'copii_documente_indexate':len(success),'programe_fara_document_descarcat':sum(not any(d['ok'] for d in documents[r['id']]) for r in cat)}
intro=[
 ('Data verificării',DATE+' • situație la data cercetării; termenele și corrigendumurile se pot modifica.'),
 ('Ce conține',f"{len(cat)} programe, subprograme și scheme (inclusiv arhivă); {len(topics)} topicuri oficiale UE cu cel puțin un termen începând de la 01.10.2026."),
 ('Documente salvate',f"{stats['documente_originale_distincte']} documente originale distincte după conținut; {stats['url_documente_reusite']} URL-uri descărcate; {stats['copii_documente_indexate']} copii indexate în folderele programelor."),
 ('Sensul inventarului','Oportunități și mecanisme de finanțare. Nu este baza istorică a beneficiarilor români care au primit bani. Un program poate avea mai multe componente și topicuri; numerele nu reprezintă apeluri independente.'),
 ('Acces direct','Coloana MOD ACCES arată dacă se aplică la organismul european sau prin operator, instituție gazdă ori bancă. SEE/Norvegia și Elveția sunt incluse separat chiar când gestionarea este în România.'),
 ('Eligibilitatea României','Programele UE sunt candidate pentru entități românești ca stat membru. Un topic poate limita solicitanții, țările, proprietatea sau consorțiul. Filtrul de firmă / ONG / persoană este orientativ, nu aprobare de eligibilitate.'),
 ('Statut și termene','PROGRAM ACTIV nu înseamnă APEL DESCHIS. Topicurile viitoare, etapele pe invitație și certificarea fără grant sunt marcate. Ora exactă de închidere se verifică în portal; orele ISO 00:00 nu au fost interpretate drept ore-limită.'),
 ('Bugete','Bugetul apelului / programului este separat de suma pe proiect. Nu aduna bugetele: programe-umbrelă, componente și ani se suprapun. EUR / CHF / lei și grant / credit / capital sunt păstrate distinct.'),
 ('Ghiduri și versiuni','Ghiduri / call fiche / work programmes / reguli / anexe originale unde au fost accesibile. Rezumatele FISA FINANTARII.txt și copiile paginilor web nu sunt ghiduri oficiale. Versiunile vechi sunt indicate ca arhivă.'),
 ('Lipsuri',f"{stats['url_documente_nereusite']} URL-uri de documente nu au putut fi salvate. Foaia Lipsuri arată și programele fără document descărcat; linkurile oficiale sunt păstrate."),
 ('Cum folosești','Începe cu Firme, Persoane sau ONG, apoi verifică Registru și Apeluri UE. Termene apropiate acoperă 01.10–31.12.2026; înregistrările oficial închise / certificările sunt excluse din această selecție.'),
 ('Limitele acoperirii','Cercetare largă în surse oficiale, fără garanție de exhaustivitate. Portalul a fost interogat pentru oportunități tip 1, limba EN, termen ≥01.10.2026; 442 rezultate unice, nu totalul tuturor apelurilor europene. Apelurile exclusiv externe portalului și finanțările în cascadă trebuie urmărite la operatori.'),
 ('Organizare','Fiecare finanțare are 1.DOCUMENTE OFICIALE, 0.ARHIVA, fișă și link de sursă. Topicurile UE au copii ale datelor oficiale în 2.APELURI PORTAL. Structura urmează modelul folderelor românești existente.')]
for no,(label,value) in enumerate(intro,4):
 start.cell(no,1,label);start.merge_cells(start_row=no,start_column=1,end_row=no,end_column=2);start.cell(no,1).font=Font(bold=True,color=navy,size=11);start.cell(no,1).alignment=Alignment(wrap_text=True,vertical='top')
 start.cell(no,3,value);start.merge_cells(start_row=no,start_column=3,end_row=no,end_column=8);start.cell(no,3).font=Font(size=11,color='223648');start.cell(no,3).alignment=Alignment(wrap_text=True,vertical='top');start.row_dimensions[no].height=47
 if no%2==0:
  for c in start[no]:c.fill=PatternFill('solid',fgColor=pale)
for col in 'ABCDEFGH':start.column_dimensions[col].width=18
start.freeze_panes='A4'

headers=['ID','FINANȚARE / PROGRAM','FAMILIE','STATUT 01.10.2026','TERMEN PUBLICAT','FIRME','ONG','PERSOANE FIZICE','INSTITUȚII PUBLICE','MOD ACCES','INSTRUMENT','CE FINANȚEAZĂ','SOLICITANȚI / RO','SUMĂ PE PROIECT','BUGET PROGRAM / APEL','RATĂ / COFINANȚARE','PARTENERIAT','SURSĂ OFICIALĂ','APLICARE / PORTAL','FOLDER LOCAL','DOCUMENTE SALVATE','OBSERVAȚII','PROGRAM-PĂRINTE']
def regrow(r):return [r['id'],r['title'],groups[r['group']],r['status'],r['deadline'],r['company'],r['ngo'],r['person'],r['public'],r['mode'],r['instrument'],r['what'],r['who'],r['amount'],r['budget'],r['rate'],r['partnership'],r['source'],r['apply'],r['folder'],len([d for d in documents[r['id']] if d['ok']]),r['note'],r['parent']]
widths={h:25 for h in headers};widths.update({'ID':13,'FINANȚARE / PROGRAM':43,'FAMILIE':24,'STATUT 01.10.2026':38,'TERMEN PUBLICAT':18,'CE FINANȚEAZĂ':43,'SOLICITANȚI / RO':45,'OBSERVAȚII':75,'SURSĂ OFICIALĂ':36,'APLICARE / PORTAL':36,'FOLDER LOCAL':43,'DOCUMENTE SALVATE':14})
rlinks={18:'web',19:'web',20:'local'}
sheet('Registru',headers,[regrow(r) for r in cat],widths,rlinks)
for title,flag in [('Firme','company'),('Persoane','person'),('ONG','ngo'),('Instituții','public')]:sheet(title,headers,[regrow(r) for r in cat if usable(r) and eligible(r[flag])],widths,rlinks)
th=['COD TOPIC','TITLU OFICIAL','PROGRAM','ID PROGRAM','COD APEL','STATUT / ETAPĂ','URMĂTORUL TERMEN','TOATE TERMENELE','DESCHIDERE PUBLICATĂ','INSTRUMENT','CONTRIBUȚIE INDICATIVĂ / PROIECT','BUGET TOPIC / ACȚIUNE PE AN','NUMĂR INDICATIV GRANTURI','PROFIL SOLICITANT ORIENTATIV','FIRME – DE VERIFICAT','PERSOANE','LINK OFICIAL APEL','FOLDER PROGRAM','OBIECT / REZULTATE AȘTEPTATE','CONDIȚII OFICIALE – EXTRAS','OBSERVAȚII','STATUT COD PORTAL']
trows=[[t['id'],t['title'],t['family'],t['parent'],t['call'],t['status'],t['next'],t['dates'],t['opening'],t['instrument'],t['amount'],t['budget'],t['expected'],t['who'],t['company'],t['person'],t['apply'],t['folder'],t['what'],t['conditions'],t['note'],t['status_code']] for t in topics]
sheet('Apeluri UE',th,trows,{'COD TOPIC':44,'TITLU OFICIAL':65,'PROGRAM':38,'COD APEL':42,'STATUT / ETAPĂ':38,'CONDIȚII OFICIALE – EXTRAS':85,'OBIECT / REZULTATE AȘTEPTATE':80,'OBSERVAȚII':70,'LINK OFICIAL APEL':45,'FOLDER PROGRAM':45},{17:'web',18:'local'})
near=[]
for t in topics:
 d=date(t['next'])
 if d and TODAY<=d<=datetime.date(2026,12,31) and not any(s in t['status'].lower() for s in ['închis','inițială expirată','certificare']):near.append([t['next'],t['id'],t['title'],'UE – topic',t['status'],t['who'],t['amount'],t['apply'],t['folder'],t['note']])
for r in cat:
 d=date(r['deadline'])
 if d and TODAY<=d<=datetime.date(2026,12,31) and usable(r):near.append([r['deadline'],r['id'],r['title'],'Program / schemă – vezi apelul',r['status'],r['who'],r['amount'],r['apply'],r['folder'],r['note']])
near.sort(key=lambda x:(x[0],x[1]))
sheet('Termene apropiate',['TERMEN','COD','FINANȚARE','TIP ÎNREGISTRARE','STATUT','SOLICITANȚI','SUMĂ PE PROIECT','LINK APLICARE','FOLDER LOCAL','OBSERVAȚII'],near,{'FINANȚARE':62,'OBSERVAȚII':80,'SOLICITANȚI':48},{8:'web',9:'local'})
sheet('Prin operatori',headers,[regrow(r) for r in cat if r['group']!='ARHIVA' and route(r)!='Aplicare la organism european'],widths,rlinks)
sheet('Arhivă și închise',headers,[regrow(r) for r in cat if closed(r)],widths,rlinks)

dh=['ID PROGRAM','PROGRAM','DOCUMENT OFICIAL','TIP','VERSIUNE / ARHIVĂ','DESCĂRCARE','FIȘIER LOCAL','URL DOCUMENT OFICIAL','PAGINĂ SURSĂ','DIMENSIUNE MiB','PAGINI PDF','SHA-256','DETALIU / EROARE']
drows=[]
for d in manifest:drows.append([d['program'],byid[d['program']]['title'],d['label'],d.get('ext','').lstrip('.').upper(),'ISTORIC / ARHIVĂ' if d['archive'] else 'Document la data verificării','SALVAT' if d['ok'] else 'NESALVAT',d.get('relative_path',''),d['url'],d.get('source',''),round(d.get('bytes',0)/1048576,3) if d['ok'] else None,d.get('pages'),d.get('sha256',''),d.get('error','') if not d['ok'] else 'Original verificat PDF / Office / ZIP; clasificarea documentului se verifică în conținut.'])
sheet('Documente',dh,drows,{'DOCUMENT OFICIAL':65,'FIȘIER LOCAL':65,'URL DOCUMENT OFICIAL':50,'PAGINĂ SURSĂ':45,'SHA-256':30,'DETALIU / EROARE':70},{7:'local',8:'web',9:'web'})
gaps=[]
for r in cat:
 ds=documents[r['id']]
 if not any(d['ok'] for d in ds):gaps.append([r['id'],r['title'],'Niciun document descărcat','Nu a fost identificat sau accesat un ghid / document original descărcabil. Pagina oficială și portalul rămân sursele de verificare.',r['source'],r['folder']])
for d in failed:gaps.append([d['program'],byid[d['program']]['title'],'URL document nesalvat',d.get('error',''),d['url'],byid[d['program']]['folder']])
sheet('Lipsuri',['ID','PROGRAM','SITUAȚIE','MOTIV / URMĂTORUL PAS','LINK OFICIAL','FOLDER LOCAL'],gaps,{'PROGRAM':45,'MOTIV / URMĂTORUL PAS':95,'LINK OFICIAL':65,'FOLDER LOCAL':55},{5:'web',6:'local'})
sources=[]
for p in sorted(CACHE.glob('*.json')):
 s=json.loads(p.read_text(encoding='utf8'));key=s.get('key',p.stem)
 sources.append([key,s.get('title',key),s.get('url_final',s['url']),'Pagină descărcată' if s['ok'] else ('Verificată în browser; descărcare blocată' if key in ['EYF_2027','Eurimages','Swiss','Swiss_Projects','EUIPO'] else 'Acces automat nereușit'),DATE,len(s.get('documents',[])),s.get('error','')])
sources.extend([['PORTAL_API','Funding & Tenders – 442 rezultate unice','https://api.tech.ec.europa.eu/search-api/prod/rest/search','Export oficial JSON complet pentru filtrul folosit',DATE,442,'Interogare tip=1, limba=en, deadlineDate ≥2026-10-01. API key SEDIA publică.'],['EURIMAGES_CAL','Eurimages – calendar 2026 / 2027','https://www.coe.int/en/web/eurimages/deadlines','Verificat în browser',DATE,0,'13.01.2027 următorul termen, 12:00 ora locală Franța.']])
sheet('Surse',['CHEIE','SURSA / TITLU','URL OFICIAL','ACCES / VERIFICARE','DATA','LINKURI DOCUMENTE GĂSITE','DETALIU'],sources,{'SURSA / TITLU':60,'URL OFICIAL':65,'DETALIU':90},{3:'web'})
exclusions=[
 ['PNRR / Fonduri de coeziune / PAC / AFIR / APIA','Gestionare în România','Nu sunt apeluri directe ale Comisiei către o firmă în sensul acestui registru. Proiectele românești existente rămân sursa pentru aceste mecanisme.','https://commission.europa.eu/funding-and-tenders/find-funding/funding-management-mode_en'],
 ['Fondul pentru Modernizare','Gestionare prin statul beneficiar / autoritățile din România','Fond european relevant pentru România, dar finanțarea solicitantului este prin scheme naționale. Nu este confundat cu Innovation Fund direct.','https://climate.ec.europa.eu/eu-action/eu-funding-climate-action/modernisation-fund_en'],
 ['TSI / Fiscalis / Customs','Beneficiari instituționali desemnați','Nu sunt granturi generale pentru SRL sau persoane fizice. Serviciile / achizițiile publice nu reprezintă grant direct pentru ofertant.','https://commission.europa.eu/funding-and-tenders/find-funding/eu-funding-programmes_en'],
 ['Achiziții publice europene / contracte de consultanță','Contracte comerciale, nu granturi','Au reguli și proceduri distincte. Nu sunt inventariate toate licitațiile sau posturile de muncă europene.','https://ec.europa.eu/info/funding-tenders/opportunities/portal/'],
 ['Eurimages – pilot seriale 2026','România: coproducător posibil; țara solicitantului este restrânsă','România figurează între statele Eurimages, dar nu în lista țărilor participante la pilotul 2026. Nu presupune solicitant român direct. Apelul 2026 este închis (14.04).','https://www.coe.int/en/web/programme-for-series-co-productions/co-production-support-2026'],
 ['Finanțare în cascadă / FSTP','Prin proiecte și operatori','Posibilă pentru firme / persoane românești, dar nu există un singur registru complet. Este menționată când apare în topic; eligibilitatea și apelul final se verifică la proiectul care acordă subgrantul.','https://ec.europa.eu/info/funding-tenders/opportunities/portal/']]
sheet('Delimitare',['MECANISM','MOD ACCES','EXPLICAȚIE','SURSĂ OFICIALĂ'],exclusions,{'MECANISM':44,'MOD ACCES':48,'EXPLICAȚIE':105,'SURSĂ OFICIALĂ':65},{4:'web'})
for no,name in enumerate(WB.sheetnames[1:],19):
 start.cell(no,1,'DESCHIDE →');start.cell(no,1).font=Font(color=teal,bold=True);start.merge_cells(start_row=no,start_column=3,end_row=no,end_column=8);c=start.cell(no,3,name);c.hyperlink="#'"+name+"'!A1";c.font=Font(color=blue,underline='single',bold=True,size=12);start.row_dimensions[no].height=25
for ws in WB:
 ws.sheet_properties.tabColor=navy if ws.title=='Start' else teal if ws.title in ['Firme','Persoane','ONG','Termene apropiate'] else blue if ws.title in ['Registru','Apeluri UE'] else grey
WB.properties.creator='Cercetare surse oficiale – Codex';WB.properties.title='Finanțări europene pentru România – 01.10.2026';WB.properties.description='Programe, apeluri, documente originale și surse. Eligibilitatea se verifică în ghid.'
WB.save(OUT/NAME)

index=OUT/'00. Index si surse';index.mkdir(exist_ok=True)
(BASE/'catalogue_final.json').write_text(json.dumps(cat,ensure_ascii=False,indent=2),encoding='utf8')
(BASE/'manifest_final.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
(index/'Registru_programe.json').write_text(json.dumps(cat,ensure_ascii=False,indent=2),encoding='utf8')
(index/'Inventar_documente.json').write_text(json.dumps([{k:v for k,v in d.items() if k!='file'} for d in manifest],ensure_ascii=False,indent=2),encoding='utf8')
shutil.copy2(BASE/'portal_all.json',index/'Funding_Tenders_export_oficial_442.json')
(index/'Surse_verificate.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf8')
readme='FINANȚĂRI EUROPENE PENTRU ROMÂNIA – '+DATE+'\n\n'+'\n\n'.join(a.upper()+'\n'+b for a,b in intro)+'\n\nFIȘIER PRINCIPAL\n'+NAME+'\n\nLINKURI OFICIALE PRINCIPALE\nhttps://commission.europa.eu/funding-and-tenders/find-funding/eu-funding-programmes_en\nhttps://ec.europa.eu/info/funding-tenders/opportunities/portal/\nhttps://eeagrants.org/en/fmo/news/renewed-cooperation-romania\nhttps://eeagrants.org/en/eea-civil-society-fund-romania/search\nhttps://www.swiss-contribution.ro/\nhttps://www.schweiz-rumaenien.eda.admin.ch/en/second-swiss-contribution-projects\n\nFOARTE APROAPE\nCSF: 7 apeluri ONG, termen 08.10.2026.\nEYF: 01.10.2026, 23:59 CET, conform anunțului.\nERC Starting 2027: 14.10.2026.\nCOST: 28.10.2026.\nEIC Accelerator: 04.11.2026, cu condiții pentru etapa scurtă.\nInnowwide: 01.12.2026.\n'
(OUT/'CITESTE-MA.txt').write_text(readme,encoding='utf-8-sig')
for r in cat:
 folder=OUT/r['folder'];(folder/'1.DOCUMENTE OFICIALE').mkdir(parents=True,exist_ok=True);(folder/'0.ARHIVA').mkdir(exist_ok=True)
 ds=documents[r['id']];ok=[d for d in ds if d['ok']];bad=[d for d in ds if not d['ok']]
 pairs=[('ID',r['id']),('PROGRAM',r['title']),('DATA VERIFICĂRII',DATE),('STATUT',r['status']),('TERMEN',r['deadline'] or 'Nu a fost stabilit un termen unic de program'),('CE FINANȚEAZĂ',r['what']),('SOLICITANȚI',r['who']),('FIRME',r['company']),('ONG',r['ngo']),('PERSOANE FIZICE',r['person']),('INSTITUȚII PUBLICE',r['public']),('MOD DE ACCES',r['mode']),('INSTRUMENT',r['instrument']),('SUMĂ PE PROIECT',r['amount']),('BUGET PROGRAM / APEL',r['budget'] or 'Se verifică la nivel de apel'),('RATĂ / COFINANȚARE',r['rate']),('PARTENERIAT',r['partnership']),('SURSĂ OFICIALĂ',r['source']),('APLICARE',r['apply']),('OBSERVAȚII',r['note']),('DOCUMENTE',f'{len(ok)} documente indexate salvate; {len(bad)} linkuri nesalvate')]
 text='FIȘĂ DE ORIENTARE – NU ESTE GHID OFICIAL\nCondițiile ghidului / topicului prevalează.\n\n'+'\n\n'.join(a+'\n'+str(b) for a,b in pairs)
 if not ok:text+='\n\nNu a fost identificat / accesat un document original descărcabil. Consultă linkul oficial; nu există un PDF de ghid inventat în acest folder.'
 if bad:text+='\n\nDOCUMENTE NESALVATE\n'+'\n'.join(d['label']+'\n'+d['url']+'\n'+d.get('error','')+'\n' for d in bad)
 (folder/'FISA FINANTARII.txt').write_text(text,encoding='utf-8-sig');(folder/'SURSA OFICIALA.url').write_text('[InternetShortcut]\nURL='+r['source']+'\n',encoding='utf8')
 p=CACHE/(r['source_key']+'.json')
 if p.exists():
  s=json.loads(p.read_text(encoding='utf8'))
  if s.get('ok'):
   (folder/'1.DOCUMENTE OFICIALE'/'Pagina oficiala - copie text.txt').write_text('COPIE TEXT A PAGINII OFICIALE, NU GHID PDF\nURL: '+s.get('url_final',s['url'])+'\nData descărcării: '+DATE+'\n\n'+s.get('text',''),encoding='utf-8-sig')

htmlrows=''.join('<tr><td>'+html.escape(r['id'])+'</td><td><a href="'+rel_link(r['folder'])+'/FISA%20FINANTARII.txt">'+html.escape(r['title'])+'</a></td><td>'+html.escape(groups[r['group']])+'</td><td>'+html.escape(r['status'])+'</td><td>'+html.escape(r['deadline'])+'</td><td><a href="'+html.escape(r['source'],quote=True)+'">Sursa oficială</a></td></tr>' for r in cat)
(OUT/'INDEX.html').write_text('<!doctype html><html lang="ro"><meta charset="utf-8"><title>Finanțări RO – 01.10.2026</title><style>body{font:15px system-ui;margin:32px;color:#123b5d}h1{font-size:28px}table{border-collapse:collapse;width:100%}td,th{padding:10px;border-bottom:1px solid #dbe4ed;text-align:left}thead{background:#123b5d;color:white}a{color:#087e8b}tr:nth-child(even){background:#eef5f8}input{padding:12px;width:65%;margin:16px 0}</style><h1>Finanțări europene pentru România</h1><p>Verificare: 01.10.2026. '+str(len(cat))+' programe / scheme; 442 topicuri UE. Ghidurile și apelurile oficiale prevalează.</p><p><a href="'+NAME+'">Deschide registrul Excel</a> · <a href="CITESTE-MA.txt">Cum folosești inventarul</a></p><input placeholder="Caută un program, domeniu sau statut" oninput="let q=this.value.toLowerCase();document.querySelectorAll(\'tbody tr\').forEach(r=>r.hidden=!r.textContent.toLowerCase().includes(q))"><table><thead><tr><th>ID</th><th>Finanțare</th><th>Familie</th><th>Statut</th><th>Termen</th><th>Link</th></tr></thead><tbody>'+htmlrows+'</tbody></table></html>',encoding='utf8')

# Verify the saved deliverable and every indexed original document.
check=load_workbook(OUT/NAME,read_only=False);assert check['Registru'].max_row==len(cat)+3;assert check['Apeluri UE'].max_row==445;assert len({r['id'] for r in cat})==len(cat)
assert all((OUT/r['folder']/'FISA FINANTARII.txt').is_file() for r in cat)
errors=[]
for d in success:
 p=OUT/d['relative_path']
 if not p.is_file() or p.stat().st_size!=d['bytes'] or hashlib.sha256(p.read_bytes()).hexdigest()!=d['sha256']:errors.append(d['relative_path'])
assert not errors,errors
for id in ['CSF1','CSF2','CSF3','CSF4','CSF5','CSF6','CSF7']:
 assert any(d['ok'] and d.get('ext')=='.pdf' for d in documents[id]),id
 assert any(d['ok'] and 'ghidul-solicitant' in d['url'] for d in documents[id]),id+' Romanian guide'
 assert any(d['ok'] and d.get('ext')=='.xlsx' for d in documents[id]),id
for ws in check:
 for row in ws:
  for c in row:
   if c.hyperlink and not c.hyperlink.target.startswith(('https:','#')):
    from urllib.parse import unquote
    assert (OUT/unquote(c.hyperlink.target)).exists(),c.hyperlink.target
with zipfile.ZipFile(OUT/NAME) as z:assert z.testzip() is None
allfiles=[p for p in OUT.rglob('*') if p.is_file()]
stats.update(files=len(allfiles),size_MiB=round(sum(p.stat().st_size for p in allfiles)/1048576,2),excel=NAME,sheets=check.sheetnames,near_deadlines=len(near),max_network_path=max(len(str(NETWORK/p.relative_to(OUT))) for p in allfiles),verification='Excel reopened, 442 unique topics, all indexed files hashed, all local hyperlinks exist, seven CSF guide and budget sets checked')
(BASE/'final_stats.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf8');(index/'Verificare_finala.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(stats,ensure_ascii=False,indent=2),flush=True)
