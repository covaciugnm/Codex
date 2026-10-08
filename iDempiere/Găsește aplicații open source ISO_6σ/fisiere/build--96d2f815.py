from pathlib import Path
import json,csv,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
root=Path(__file__).resolve().parent.parent
out=root/'livrare';out.mkdir(exist_ok=True)
sources=json.loads((root/'work/source-checks.json').read_text(encoding='utf-8'))
txt=(root/'work/report.md').read_text(encoding='utf-8')+'\n'+(root/'work/annex.md').read_text(encoding='utf-8')
for s in sources:
    status=s['status']
    txt+=f"\n### {s['id']} {s['titlu']}\n\nUtilizare: {s['utilizare']}. Verificare: {s.get('verificat','2026-10-08')}; acces: {status}.\n\n[{s['titlu']}]({s['url']})\n"
(out/'Analiza_QMS_ISO_SixSigma_EVA.md').write_text(txt,encoding='utf-8')
with (out/'Registru_sur­se.csv'.replace('\u00ad','')).open('w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=['id','grup','titlu','url','utilizare','status','verificat','sha256','final_url'],extrasaction='ignore');w.writeheader();w.writerows(sources)
d=Document();sec=d.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name,size in [('Normal',10),('Title',23),('Heading 1',16),('Heading 2',12)]:
    st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='EVA  •  Managementul calității  •  Analiză pentru conducere și IT';header.style='Caption'
p=sec.footer.paragraphs[0];p.alignment=2
p.add_run('8 octombrie 2026  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'215778');prop.append(color);r.append(prop)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
lines=txt.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',x) for x in cells):rows.append(cells)
            i+=1
        t=d.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid';t.autofit=False
        widths=([6.9]+[2.1]*5) if len(rows[0])==6 else ([4.3,6.1,7] if len(rows[0])==3 else [17.4/len(rows[0])]*len(rows[0]))
        for rowidx,row in enumerate(rows):
            cells=t.add_row().cells
            trpr=cells[0]._tc.getparent().get_or_add_trPr()
            cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if rowidx==0:
                rep=OxmlElement('w:tblHeader');trpr.append(rep)
            for j,value in enumerate(row):
                cells[j].width=Cm(widths[j]);p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                r=p.add_run(value);r.font.size=Pt(8.5);r.bold=rowidx==0
                if rowidx==0:p.paragraph_format.keep_with_next=True
                if rowidx==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEFF1');cells[j]._tc.get_or_add_tcPr().append(sh)
        d.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if line.startswith('# '):d.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        p=d.add_paragraph(line[3:],'Heading 1')
        if line.startswith('## 17 '):p.paragraph_format.page_break_before=True
    elif line.startswith('### '):d.add_paragraph(line[4:],'Heading 2')
    else:
        p=d.add_paragraph()
        m=re.fullmatch(r'\[(.+)\]\((.+)\)',line)
        if m:link(p,m[1]+' — deschide sursa',m[2])
        else:p.add_run(line)
    i+=1
d.core_properties.title='Analiza soluțiilor de management al calității pentru Eva Accounting și iDempiere'
d.core_properties.subject='Comparație funcțională și proiect de integrare QMS ISO Six Sigma'
d.save(out/'Analiza_QMS_ISO_SixSigma_EVA.docx')
(out/'README.md').write_text('# Analiza QMS pentru EVA\n\nRaport din 8 octombrie 2026 pentru conducere și IT.\n\n- [Raport Word](Analiza_QMS_ISO_SixSigma_EVA.docx)\n- [Raport navigabil](Analiza_QMS_ISO_SixSigma_EVA.md)\n- [Registru de surse](Registru_surse.csv)\n\nBază de analiză: commit `1c039193e0c1626cb03fb8ca37f9980a1703fd0c`. Este o analiză documentară și de cod, fără teste în producție. Arhitectura QMS nativă este propusă, nu implementată.\n',encoding='utf-8')
print('Sources',len(sources),'words',len(txt.split()),'files',[p.name for p in out.iterdir()])
