from pathlib import Path
import json,csv,re,html,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(r'\\192.168.100.169\Comun\00.Roboti\iDempiere\Documentatie\BI_iDempiere')
LOCAL=Path(__file__).parent
for n in ['00_Raport','01_iDempiere','02_Matrice','04_Arhitectura_AI','05_Pilot_si_operare','08_Administrare']: (ROOT/n).mkdir(exist_ok=True)
groups=[
('Acces și modelare',[
('PostgreSQL','C C C C C N E','Driver și cont read-only; configurația Eva rămâne de făcut.'),
('Conectare la surse SQL multiple','N N N N C L H','Surse multiple nu înseamnă join federat automat.'),
('Join între baze eterogene','D D C L D D G','Necesită warehouse/federare; Redash QRDS este limitat.'),
('Fișiere CSV pentru analiză','C C C C C C U','Importul ERP nu echivalează cu dataset BI self-service.'),
('ETL complet și orchestrare','D L L D C D G','dbt și orchestrator extern la Lightdash; funcțiile noi Metabase depind de ediție.'),
('CDC și ștergeri incrementale','D D D D D D G','Infrastructură separată pentru toate produsele.'),
('Vederi și seturi virtuale SQL','N N C N C C E','Semantică ce trebuie proiectată.'),
('Metrici reutilizabile','N C N L N C L','Nu sunt metrici RO preinstalate.'),
('Dimensiuni și agregări','N N N N N N E','Tipul agregării depinde de grain.'),
('Semantică versionată în Git','C L L C N C G','Export/import Metabase avansat este comercial; Lightdash orientat către cod.'),
('Glosar business guvernat','D L P D C C G','Nu confundăm descrieri de câmp cu glosar guvernat.'),
('Istoric SCD și snapshots','D D D D C D G','Materializarea istoriei aparține stratului de date.'),
('Cache rezultate','C C C N C L U','Scope și invalidare de verificat.'),
('Interogări asincrone','C C C N C C U','Configurație worker/limite.'),
('Replica read-only','C C C C C C A','iDempiere documentează suport; nu verificat în runtime Eva.'),
('KPI financiar RO gata definit','D D D D D C E','Eva are formule fixe; cele externe necesită catalog.'),
]),
('Explorare și vizualizare',[
('Editor SQL','N N C N C C H','În ERP: dezvoltator/dicționar, nu SQL Lab pentru business.'),
('Editor vizual fără SQL','N N N L N L G','Superset Explore operează peste dataset pregătit.'),
('Filtre și parametri','N N N N N N E','Filtrele nu sunt politici de securitate.'),
('Filtre încrucișate','C C L L C L G','Depinde de tipurile de grafic și ediție.'),
('Drill-down analitic','N N C L N C L','Metabase N în aplicație; în guest embedded este limitat.'),
('Link spre document ERP','C C C C C N E','Necesită verificare acces în sistemul țintă.'),
('Tabele pivot','N N P N N C U','Knowage: widget pivot avansat EE; OLAP CE este separat.'),
('OLAP multidimensional dedicat','D L N L L L G','Report Cube iDempiere este sumar contabil, nu server MDX.'),
('Grafice bară linie arie pie','N N N N N N L','Graficele generice cer date și mapări.'),
('Vizualizări avansate și geo','N N L C C L G','Tipurile exacte se verifică pe versiunea aleasă.'),
('Comparații perioade și cumulări','N C C C N C E','Limbajul metricii trebuie să precizeze calendarul.'),
('Layout dashboard editabil','N N N N N N G','Eva managerial inspectat este formular tabular fix.'),
('KPI țintă și praguri','C N N C C N L','Datele buget/țintă nu apar automat.'),
('Preferințe individuale','C N C C C N A','Prezența frameworkului nu confirmă configurarea Eva.'),
('Utilizare pe ecran mic','C C C C C C U','Test vizual pe dashboardul real; fără garanție generală.'),
('Localizare și vocabular companie','C C C C C C E','Separăm traducerea interfeței de semantică.'),
]),
('Raportare și distribuție',[
('Raport tabular','N N N N N N E','Raportare uzuală.'),
('Raport paginat pixel-perfect','D L N D D N E','Jasper/BIRT în Knowage; Jasper/print engine ERP.'),
('Export CSV','N N N N N N E','Control acces identic vizualizării.'),
('Export XLSX','C N C C C C U','Formatul exact și motorul se verifică.'),
('Captură sau document PDF','C C N C C N E','PDF dashboard nu este raport contabil paginat.'),
('Programare distribuție','C C P C C C U','Knowage scheduler batch este EE în matricea curentă.'),
('Alertă bazată pe date','C N C N C N G','Threshold fără notificare nu este alertă completă.'),
('Distribuție e-mail','C C C C C C U','SMTP și destinatari de configurat; automatizarea poate fi limitată de ediție.'),
('Webhook sau integrare externă','C C C C C D G','Depinde de API/canal; nu se presupune bidirecțional.'),
('Colecții foldere catalog','N N N L N N L','Organizarea rapoartelor.'),
('Dashboard public','C C C C C C G','Nu recomandăm pentru datele financiare.'),
('Dossier Word/PDF din dashboarduri','D D P D D D G','Funcție distinctă de exportul unui grafic.'),
('Forecast statistic','D D C D D D G','Necesită model, istoric și validare; nu inclus doar prin grafic temporal.'),
('Buget versus realizat','C C C C C C L','Necesită buget cu versiune și aceeași granularitate.'),
('Writeback tranzacțional ERP','D D D D D N E','BI nu scrie direct în tabele contabile.'),
('Audit decizie și proveniență KPI','D D C D C C G','Necesită registru comun chiar dacă există loguri de aplicație.'),
]),
('Securitate integrare și AI',[
('Roluri și permisiuni obiecte','N N N N N N E','Modele diferite de autorizare.'),
('RLS per utilizator în aplicație','C P C D C N L','Redash: izolare externă pe surse; Eva: verificare org/schema obligatorie.'),
('Protecție coloane sensibile','C P C D C C U','Vederi DB restrânse când produsul nu oferă control suficient.'),
('Administrare multi-tenant completă','D L P D L N E','Nu confundăm filtrele de rând cu tenancy complet.'),
('SSO configurabil','C L C C C C U','Metabase JWT/SAML/OIDC și SSO embedded sunt comerciale; unele loginuri de bază sunt gratuite.'),
('Embedding securizat de consum','C C C L P D G','Necesită adaptor Eva; Redash nu echivalent guest RLS.'),
('Self-service în interfața embedded','D P C D P D G','Superset SDK guest nu oferă editor complet.'),
('API pentru configurare','N N C N N C H','Acoperirea endpointurilor și ediția se validează.'),
('Audit acces și export','C P P L L N U','Loguri tehnice diferite de audit dashboard dedicat.'),
('Configurare BI din profil companie','D D D D D D G','Niciun flux gratuit complet demonstrat pentru Eva.'),
('AI local generare dashboard','L L U D P D G','Superset MCP și Metabase AI creează conținut; profilul complet Eva cere dezvoltare. Lightdash self-hosted AI este EE.'),
('NLQ cu AI nativ gratuit confirmat','L C U D P D G','Metabase permite AI cu provider propriu fără plan plătit. Superset cere client AI extern prin MCP. S51-S55.'),
('Conector model local propriu','L C U D D D H','Metabase documentează vLLM; Superset prin clientul MCP. Eva gateway istoric, nu configurator BI activ. S53 S55.'),
('MCP nativ self-hosted gratuit','C C U D P D G','MCP documentat pentru Superset 6.1 și Metabase. Modelul/clientul AI și securizarea se configurează. S52 S55.'),
('Promovare și rollback configurație','C L P C C C G','Snapshot și API nu echivalează cu proces validat.'),
('Backup și restaurare','C C C C C C U','Se testează DB metadate și config, nu doar containere.'),
])]
sources=['S01 S40 S46 S17 S18; manualele locale ale produselor','S01 S06 S10 S17 S18; manualele de vizualizare','S10 S12-S14 S18 S39; manualele de raportare','S06-S09 S12-S17 S20 S34 S38 S41-S42 S50-S55; audit Eva']
matrix=[]
for gi,(cat,rows) in enumerate(groups):
 for i,(name,vals,note) in enumerate(rows):matrix.append({'id':f'F{len(matrix)+1:03}','categorie':cat,'functie':name,**dict(zip(['Superset','Metabase_OSS','Knowage_CE','Redash','Lightdash_OSS','iDempiere','Eva'],vals.split())),'nota':note,'surse':sources[gi]})
with (ROOT/'02_Matrice/matrice_64_functii.csv').open('w',encoding='utf-8-sig',newline='') as f:w=csv.DictWriter(f,fieldnames=matrix[0].keys());w.writeheader();w.writerows(matrix)
weights=[20,20,20,15,10,10,5]
scoredata={'Superset':[5,5,4,5,4,2,3],'Metabase OSS':[4,3,3,5,4,4,5],'Knowage CE':[4,3,3,3,3,2,3],'Redash':[3,2,2,3,3,4,2],'Lightdash OSS':[4,2,3,4,3,3,4]}
with (ROOT/'02_Matrice/scoruri.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['produs','analiza_20','integrare_20','securitate_20','AI_15','rapoarte_10','operare_10','business_5','scor_100']);w.writerows([[p,*v,sum(a*b for a,b in zip(v,weights))/5] for p,v in scoredata.items()])
text=(LOCAL/'report_content.txt').read_text(encoding='utf-8')
for gi,(cat,rows) in enumerate(groups):
 md='Legenda: N nativ gratuit; C configurare; L limitat; P comercial; D dezvoltare; U neconfirmat. ERP = iDempiere; EVA: E cod observat, A moștenit de verificat, G gol, H dormant.\n\n'
 md+='|Funcție|SUP|MET|KNO|RED|LIG|ERP|EVA|\n|---|---|---|---|---|---|---|---|\n'
 md+='\n'.join('|'+name+'|'+'|'.join(vals.split())+'|' for name,vals,note in rows)
 md+='\n\nDetalii și limite pe funcție: Comparatie_BI_iDempiere_Eva.xlsx și matrice_64_functii.csv.\nSurse: '+sources[gi]+'.'
 text=text.replace('{{MATRIX'+str(gi+1)+'}}',md)
score='|Produs|Analiză|Integrare|Securitate|AI|Raport|Operare|Business|Scor|\n|---|---|---|---|---|---|---|---|---|\n'
score+='\n'.join('|'+p+'|'+'|'.join(str(x) for x in v)+'|'+str(round(sum(a*b for a,b in zip(v,weights))/5,1))+'|' for p,v in scoredata.items())
text=text.replace('{{SCORES}}',score)
src=json.loads((ROOT/'06_Surse_oficiale/surse.json').read_text(encoding='utf-8'))
for n in range(0,len(src),17):
 text+='\n@@PAGE Registrul surselor '+str(n//17+1)+'\n'
 text+='Surse consultate la 8 octombrie 2026. URL-urile sunt active în versiunea digitală. Fișierele locale și hashurile sunt în 06_Surse_oficiale/surse.json. HTTP 200 nu dovedește singur o pagină utilă; redirecționările scurte Superset S02/S03/S05 sunt înlocuite documentar prin S38/S39/S40 și manualele din repository.\n\n'
 for r in src[n:n+17]:
  title=r['title'].split('_',1)[1].replace('_',' ')
  state='salvat local' if r.get('local') else 'descărcare indisponibilă; referință online'
  text+=f"{r['id']} • {title} • {state}\n{r['url']}\n\n"
pages=[p.strip().split('\n',1) for p in text.split('@@PAGE ')[1:]]
md='\n\n'.join('# '+title+'\n\n'+body for title,body in pages)
(ROOT/'00_Raport/Raport_BI_iDempiere_Eva_2026-10-08.md').write_text(md,encoding='utf-8')
(ROOT/'08_Administrare/report_pages.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
doc=Document();st=doc.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.08
for sn in ['Title','Heading 1','Heading 2','Heading 3','Subtitle']:
 st=doc.styles[sn];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(22 if sn=='Title' else 17 if sn=='Heading 1' else 12)
def setsec(sec,land):
 sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
 sec.page_width=Cm(29.7 if land else 21);sec.page_height=Cm(21 if land else 29.7)
 sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
setsec(doc.sections[0],False)
footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Eva Accounting • BI • ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def addlink(p,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');props=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'174F79');props.append(c);r.append(props);t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)
def table(lines,land):
 rows=[[x.strip() for x in l.strip('|').split('|')] for l in lines if not re.match(r'^\|[- :|]+\|$',l)]
 t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 width=26.3 if land else 17.6
 widths=([9.5]+[(width-9.5)/7]*7) if land and len(rows[0])==8 else ([4.0]+[(width-4)/(len(rows[0])-1)]*(len(rows[0])-1))
 if len(rows[0])==9:widths=[3.5]+[(width-3.5)/8]*8
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  pr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  if ri==0:rep=OxmlElement('w:tblHeader');pr.append(rep)
  for ci,txt in enumerate(row):
   cell=cells[ci];cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173B55' if ri==0 else 'F1F5F8' if ri%2==0 else 'FFFFFF');tcPr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for e in ['top','left','bottom','right']:
    el=OxmlElement('w:'+e);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcPr.append(borders);m=OxmlElement('w:tcMar')
   for e in ['top','bottom','left','right']:
    el=OxmlElement('w:'+e);el.set(qn('w:w'),'40' if land and e in ['top','bottom'] else '75');el.set(qn('w:type'),'dxa');m.append(el)
   tcPr.append(m)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
   if (land and ci>0) or len(rows[0])==9 and ci>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(txt);r.font.size=Pt(9 if len(rows[0])<8 else 8.5);r.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
prevland=False
for ix,(title,body) in enumerate(pages):
 land=title.startswith('Matrice de ')
 if ix:
  if land!=prevland:setsec(doc.add_section(WD_SECTION_START.NEW_PAGE),land)
  else:doc.add_page_break()
 prevland=land
 doc.add_paragraph(title,'Title' if ix==0 else 'Heading 1')
 lines=body.splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   ll=[]
   while i<len(lines) and lines[i].startswith('|'):ll.append(lines[i]);i+=1
   table(ll,land);continue
  p=doc.add_paragraph()
  if line.startswith('http'):addlink(p,line);p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1
  else:
   r=p.add_run(line)
   if title.startswith('Registrul surselor'):r.font.size=Pt(9)
  i+=1
for element in [doc.styles.element,doc.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='Business Intelligence pentru Eva Accounting și iDempiere';doc.core_properties.author='Codex';doc.core_properties.subject='Evaluare tehnică și managerială pentru integrare BI și configurare AI locală'
dest=ROOT/'00_Raport/Raport_BI_iDempiere_Eva_2026-10-08.docx';doc.save(dest)
# Offline report with print-ready chapters and complete links, no external dependencies.
def tohtml(md):
 out=[];intable=False
 for line in md.splitlines():
  if line.startswith('|'):
   if re.match(r'^\|[- :|]+\|$',line):continue
   tag='td' if intable else 'th'
   if not intable:out.append('<table>');intable=True
   out.append('<tr>'+''.join('<'+tag+'>'+html.escape(c.strip())+'</'+tag+'>' for c in line.strip('|').split('|'))+'</tr>');continue
  if intable:out.append('</table>');intable=False
  if line.startswith('# '):out.append('<h1>'+html.escape(line[2:])+'</h1>')
  elif line.startswith('http'):out.append('<p><a href="'+html.escape(line,quote=True)+'">'+html.escape(line)+'</a></p>')
  elif line:out.append('<p>'+html.escape(line)+'</p>')
 if intable:out.append('</table>')
 return '\n'.join(out)
content=tohtml(md)
htmlpage='<!doctype html><html lang="ro"><meta charset="utf-8"><title>BI pentru Eva și iDempiere</title><style>body{max-width:1100px;margin:40px auto;font:16px/1.55 Arial;color:#172c3c;padding:20px}h1{margin-top:50px;color:#000}table{border-collapse:collapse;width:100%;font-size:14px;margin:22px 0}td,th{border:1px solid #cbd5df;padding:9px}th{background:#173b55;color:white}tr:nth-child(even){background:#f1f5f8}a{overflow-wrap:anywhere}@media print{h1{break-before:page}body{font-size:11pt}}</style><body>'+content+'</body></html>'
(ROOT/'00_Raport/Raport_BI_iDempiere_Eva_2026-10-08.html').write_text(htmlpage,encoding='utf-8')
print('Report authored',len(pages),'planned pages;',len(matrix),'features;',len(md.split()),'words')
