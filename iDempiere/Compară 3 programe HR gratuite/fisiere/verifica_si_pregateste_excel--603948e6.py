from pathlib import Path
import json,hashlib,shutil,zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

BASE=Path(__file__).parent
NAME='Analiza_iDempiere_HR_SSM_SU_2026-10-08'
ROOT=BASE/NAME
data=json.loads((BASE/'date_comparativ_excel.json').read_text(encoding='utf8'))
file=BASE/'outputs/01a11ac9-7ce0-7481-87a6-870f72ae1107/Comparativ_iDempiere_HR_SSM_SU_2026-10-08.xlsx'
# Artifact Tool nu calculează HYPERLINK și nu expune setterul nativ în API-ul documentat.
# Se adaugă doar relațiile hyperlink OOXML, fără rescrierea datelor, stilurilor sau regulilor.
MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('',MAIN)
ET.register_namespace('r',REL)
sheetpath='xl/worksheets/sheet2.xml'
relpath='xl/worksheets/_rels/sheet2.xml.rels'
with zipfile.ZipFile(file) as z:
    content={n:z.read(n) for n in z.namelist()}
sheetxml=ET.fromstring(content[sheetpath])
relations=ET.fromstring(content[relpath]) if relpath in content else ET.Element('{'+PKG+'}Relationships')
old=sheetxml.find('{'+MAIN+'}hyperlinks')
if old is not None:sheetxml.remove(old)
for el in list(relations):
    if el.attrib.get('Id','').startswith('rIdDocLink'):relations.remove(el)
hyperlinks=ET.Element('{'+MAIN+'}hyperlinks')
for i,p in enumerate(data['products'],7):
    for col,ref in [('I',p['refs'][0]),('J',p['refs'][min(1,len(p['refs'])-1)])]:
        rid='rIdDocLink'+col+str(i)
        ET.SubElement(hyperlinks,'{'+MAIN+'}hyperlink',{'ref':col+str(i),'{'+REL+'}id':rid})
        ET.SubElement(relations,'{'+PKG+'}Relationship',{'Id':rid,'Type':REL+'/hyperlink','Target':data['sources'][ref]['url'],'TargetMode':'External'})
late={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
idx=next((i for i,x in enumerate(sheetxml) if x.tag.split('}')[-1] in late),len(sheetxml))
sheetxml.insert(idx,hyperlinks)
content[sheetpath]=ET.tostring(sheetxml,encoding='utf-8',xml_declaration=True)
content[relpath]=ET.tostring(relations,encoding='utf-8',xml_declaration=True)
tmp=file.with_suffix('.native.tmp')
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in content.items():z.writestr(n,b)
tmp.replace(file)
wb=load_workbook(file,data_only=False)
wc=load_workbook(file,data_only=True)
assert wb.sheetnames==['Comparativ','Produse și versiuni','Dovezi și explicații','Legendă și utilizare']
s=wb['Comparativ']
assert s.freeze_panes=='D7',str(s.freeze_panes)
assert s.max_row==125 and s.max_column==27,(s.max_row,s.max_column)
assert s.tables['ComparativFunctii'].ref=='A6:AA125'
assert s.tables['ComparativFunctii'].autoFilter.ref=='A6:AA125'
assert s.sheet_view.showGridLines is False
assert len(s.data_validations.dataValidation)==1
assert len(s.conditional_formatting)==1
rules=s.conditional_formatting[next(iter(s.conditional_formatting))]
assert len(rules)==5
assert all(x.type=='expression' and x.formula[0].startswith('D7=') for x in rules)
count=0
for i,r in enumerate(data['rows'],7):
    assert [s.cell(i,j).value for j in range(1,4)]==[r['id'],r['chapter'],r['feature']]
    for j,p in enumerate(data['products'],4):
        assert s.cell(i,j).value==data['labels'][r['cells'][p['id']]['status']],(i,j)
        count+=1
e=wb['Dovezi și explicații']
assert e.max_row==len(data['evidence'])+6
for i,r in enumerate(data['evidence'],7):
    assert [e.cell(i,j).value or '' for j in range(1,10)]==[v or '' for v in r],i
errors=[]
for sheet in wc:
    for line in sheet:
        for cell in line:
            if cell.data_type=='e':errors.append([sheet.title,cell.coordinate,cell.value])
assert not errors,errors
for i,p in enumerate(data['products'],7):
    for j,ref in [(9,p['refs'][0]),(10,p['refs'][min(1,len(p['refs'])-1)])]:
        assert wb['Produse și versiuni'].cell(i,j).hyperlink.target==data['sources'][ref]['url']
dest=ROOT/'09_Comparativ_Excel'/file.name
dest.parent.mkdir(exist_ok=True)
shutil.copy2(file,dest)
rel=dest.relative_to(ROOT).as_posix()
readme=ROOT/'README.md'
text=readme.read_text(encoding='utf8')
if '## Comparație Excel' not in text:
    section='## Comparație Excel\n\n[Deschide tabelul comparativ Excel]('+rel+'). Conține **119 funcții și criterii în 12 capitole**, cu coloane pentru **24 de aplicații, module și dependențe**. Marcaje colorate: Da, Parțial, Nu, Neconfirmat și N/A. Include versiuni, filtre, antete fixe și 1.230 de evaluări cu explicații și surse. Fișierul are patru foi și se poate folosi independent de bibliotecă.\n\n'
    text=text.replace('## Conținut\n',section+'## Conținut\n',1)
    text=text.replace('| 08_Registre_si_verificari | Proveniență, starea descărcărilor și verificarea integrității |','| 08_Registre_si_verificari | Proveniență, starea descărcărilor și verificarea integrității |\n| 09_Comparativ_Excel | Matrice comparativă XLSX, versiuni și dovezi |')
    readme.write_text(text,encoding='utf8')
index=ROOT/'INDEX.html'
text=index.read_text(encoding='utf8')
if rel not in text:
    text=text.replace('<body>','<body><p><a href="'+rel+'"><strong>Excel comparativ: 119 funcții, 24 de produse / module</strong></a></p>',1)
    index.write_text(text,encoding='utf8')
verification={'data':'2026-10-08','fisier':rel,'foi':wb.sheetnames,'functii':len(data['rows']),'capitole':len(data['chapters']),'produse_module':len(data['products']),'celule_de_statut_verificate':count,'evaluari_cu_explicatii':len(data['evidence']),'panouri_fixe':'D7','filtre':'Tabel Excel A6:AA125','formatare_conditionata':'5 reguli relative, culori verificate în randare','linkuri_documentatie':48,'erori_formule':errors,'verificare_vizuala':'Toate cele 4 foi inspectate; antetele modulelor inspectate suplimentar.','motor_de_verificare':'Artifact Tool pentru calcul/randare; citire independentă XLSX cu openpyxl, fără rescriere. Nu s-a executat Excel desktop.','sha256':hashlib.sha256(dest.read_bytes()).hexdigest()}
(ROOT/'08_Registre_si_verificari/verificare_comparativ_excel.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
manifest=[]
for p in sorted(ROOT.rglob('*')):
    if p.is_file() and p.name!='SHA256SUMS.txt':manifest.append(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(ROOT).as_posix())
(ROOT/'SHA256SUMS.txt').write_text('\n'.join(manifest)+'\n',encoding='utf8')
archive=BASE/(NAME+'.zip')
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in sorted(ROOT.rglob('*')):
        if p.is_file():z.write(p,NAME+'/'+p.relative_to(ROOT).as_posix())
with zipfile.ZipFile(archive) as z:assert z.testzip() is None
print(json.dumps({'verificare':verification,'fisiere_pachet':len(manifest)+1,'arhiva_octeti':archive.stat().st_size},ensure_ascii=False))
