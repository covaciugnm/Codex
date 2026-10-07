from pathlib import Path
import fitz, hashlib, json
from docx import Document
from PIL import Image, ImageDraw

work=Path(__file__).resolve().parent
out=work.parent
pdf=next(out.glob('01 *.pdf'))
docx=next(out.glob('01 *.docx'))
offer=next(out.glob('03 *.pdf'))
doc=fitz.open(pdf); original=fitz.open(offer)
texts=[p.get_text() for p in doc]
matches=[]
for i in range(len(doc)-2):
    if all(doc[i+j].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).samples == original[j].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).samples for j in range(3)):
        matches.append([i+1,i+2,i+3])
word=Document(docx)
wordtext='\n'.join(p.text for p in word.paragraphs)
report={'pdf_pages':len(doc),'word_paragraphs':len(word.paragraphs),'word_inline_images':len(word.inline_shapes),'original_offer_pages_exact_render_matches':matches,'original_offer_hash':hashlib.sha256(offer.read_bytes()).hexdigest(),'pdf_hash':hashlib.sha256(pdf.read_bytes()).hexdigest(),'docx_hash':hashlib.sha256(docx.read_bytes()).hexdigest(),'source_hash':hashlib.sha256((work/'CONTRACT_DE.md').read_bytes()).hexdigest(),'pdf_pages_char_counts':[len(t) for t in texts],'word_keychecks':{k:k in wordtext for k in ['8.250,00','9.900,00','31.12.2026','spätestens bei Auftragsvergabe','zwei Arbeitstagen','625,00','200,00','135,00','Anlage 3']}}
(work/'AUDIT_ARTEFACTE.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
for start in range(0,len(doc),6):
    sheet=Image.new('RGB',(1050,1515),'#cccccc'); d=ImageDraw.Draw(sheet)
    for j in range(start,min(start+6,len(doc))):
        pix=doc[j].get_pixmap(matrix=fitz.Matrix(.57,.57),alpha=False)
        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        x=((j-start)%3)*350+5;y=((j-start)//3)*755+25
        sheet.paste(im,(x,y));d.text((x,y-18),f'PDF page {j+1}',fill='black')
    sheet.save(work/f'AUDIT_contact_{start+1:02d}.png')
print(json.dumps(report,ensure_ascii=False,indent=2))
