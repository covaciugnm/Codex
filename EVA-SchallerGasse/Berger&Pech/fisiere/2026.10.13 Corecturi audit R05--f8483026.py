from pathlib import Path
import json,hashlib,importlib.util
from docx import Document
ROOT=Path(r'D:\00. Downloads\Apartamente Viena\00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
WORK=ROOT/'00.Proiect/09. Vertraege/2026.10.13 Reorganizare juridica si audit'
VIS=Path(r'C:\Users\User\.codex\visualizations\2026\10\01\01a0f6eb-c105-7760-973e-241d6f6be8c1')
spec=importlib.util.spec_from_file_location('builder',VIS/'2026.10.13 Construire contracte juridice.py');builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def writej(p,v):
    assert not p.exists(),str(p)
    p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
oldmanifest=json.loads((WORK/'2026.10.13 Export contracte R04.json').read_text(encoding='utf-8-sig'))
records=[]
for name,domain in [('BERGER','bmberger.at'),('PECH','zt-pech.at')]:
    src=WORK/f'2026.10.13 Werkvertrag {name} R04.json';data=json.loads(src.read_text(encoding='utf-8-sig'))
    data['version']=data['version'].replace('R04','R05');data['preamble'][-1]=data['preamble'][-1].replace('R04','R05')
    ss={x['number']:x for x in data['sections']};aa={x['number']:x for x in data['annexes']}
    if name=='PECH':
        ss[5]['paragraphs'][2]+=' Die gesetzliche Prüfpflicht wird durch die Terminzahl nicht begrenzt; die Vergütung zusätzlicher, nicht bereits abgegoltener Ortstermine richtet sich nach § 7.'
        old='Eigene Fehlerkorrekturen und bereits im PI-Paket geschuldete Beschauen sind keine solche Zusatzleistung.'
        new='Eigene Fehlerkorrekturen und bereits durch einen enthaltenen oder vergüteten Termin abgegoltene Beschauen sind keine Zusatzleistung. Eine gemeinsame PI-/Statik-Wahrnehmung zählt als ein Ortstermin; bei einem zusätzlichen erforderlichen, noch nicht abgegoltenen Ortstermin gilt die Vergütungsregel dieses Absatzes, ohne die gesetzliche Prüfpflicht einzuschränken.'
        assert old in ss[7]['paragraphs'][3]
        ss[7]['paragraphs'][3]=ss[7]['paragraphs'][3].replace(old,new)
        ss[11]['paragraphs'][0]+=' Zwölf Ortstermine sind im Grundhonorar enthalten. Die gesonderte Vergütung zusätzlicher, nicht bereits abgegoltener Ortstermine nach § 7 bleibt vorbehalten; sie umfasst deren dort genannte Vorbereitung, Durchführung und übliche Dokumentation, auch bei Wahrnehmung gesetzlicher PI-Prüfpflichten. Bereits enthaltene beziehungsweise bezahlte Leistungen werden nicht doppelt verrechnet.'
        aa[4]['paragraphs']=[p for p in aa[4]['paragraphs'] if not p.startswith('Für n wirksam zusätzlich beauftragte und ausgeführte Abnahmen')]
        aa[4]['paragraphs'].append('Die Vergütung zusätzlicher, noch nicht abgegoltener physischer Ortstermine folgt § 7, auch bei gesetzlich erforderlichen PI-Beschauen. EUR 14.000,00 / zwölf ist kein Zusatzterminpreis. Die Option von bis zu sechs Terminen zu jeweils EUR 450,00 netto wird nur bei wirksamer Einzelbeauftragung und tatsächlicher Ausführung vergütet; maßgeblich sind die gesetzlichen Umsatzsteuerregeln je Rechnung.')
        for row in aa[6]['tables'][1]['rows']:
            if row[0]=='PI und Bauwerksbuch':row[1]+=' Im Grundhonorar sind zwölf Ortstermine enthalten; zusätzliche noch nicht abgegoltene physische Termine werden nach § 7 gesondert vergütet. Übliche Vorbereitung und Dokumentation gehören zum jeweiligen Terminpreis; derselbe Termin wird nicht mehrfach als PI und Statik abgerechnet.'
    js=WORK/f'2026.10.13 Werkvertrag {name} R05.json';writej(js,data)
    outdir=ROOT/f'00.Proiect/09. Vertraege/{domain}/2026.10.13 Juristischer Werkvertrag'
    docx=outdir/f'2026.10.13 Werkvertrag {name} R05.docx';assert not docx.exists()
    builder.build(js,docx);d=Document(docx)
    for p in list(d.paragraphs):
        if p.text.startswith('Anlage 5 '):
            prev=p._p.getprevious()
            if prev is not None and prev.tag.endswith('}p') and not ''.join(prev.xpath('.//w:t/text()')):prev.getparent().remove(prev)
    for table in d.tables:
        if table.rows[0].cells[0].text=='Kostenfall':
            for ri,row in enumerate(table.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:p.paragraph_format.keep_with_next=ri<len(table.rows)-1
    for sec in d.sections:
        for p in sec.footer.paragraphs:
            for run in p.runs:
                if 'Entwurf 2026.10.13' in run.text:run.text=run.text.replace('Entwurf 2026.10.13','Entwurf 2026.10.13 R05')
    d.save(docx)
    full=''.join(x.text or '' for e in Document(docx).element.body for x in e.xpath('.//w:t'))
    for c in data['sections']+data['annexes']:
        for p in c.get('paragraphs',[]):assert p in full,p[:80]
        for tb in c.get('tables',[]):
            for r in [tb['headers']]+tb['rows']:
                for v in r:assert str(v) in full,str(v)
    rec=next(x.copy() for x in oldmanifest['documente'] if x['prestator']==name)
    rec.update(docx=docx.relative_to(ROOT).as_posix(),json=js.relative_to(ROOT).as_posix(),sha256_docx=sha(docx),sha256_json=sha(js))
    records.append(rec)
writej(WORK/'2026.10.13 Export contracte R05.json',{'data':'2026.10.13','documente':records})
print(json.dumps(records,ensure_ascii=True))
