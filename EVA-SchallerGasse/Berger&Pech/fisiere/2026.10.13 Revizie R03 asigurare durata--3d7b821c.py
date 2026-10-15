from pathlib import Path
import json, hashlib, importlib.util
from docx import Document

ROOT=Path(r'D:\00. Downloads\Apartamente Viena\00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
WORK=ROOT/'00.Proiect/09. Vertraege/2026.10.13 Reorganizare juridica si audit'
VIS=Path(r'C:\Users\User\.codex\visualizations\2026\10\01\01a0f6eb-c105-7760-973e-241d6f6be8c1')
spec=importlib.util.spec_from_file_location('builder',VIS/'2026.10.13 Construire contracte juridice.py')
builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)
def writej(p,obj):
    assert not p.exists(),str(p)
    p.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
records=[];changes=[]
for name,domain in [('BERGER','bmberger.at'),('PECH','zt-pech.at')]:
    src=WORK/f'2026.10.13 Werkvertrag {name} R02.json'
    data=json.loads(src.read_text(encoding='utf-8-sig'))
    old_data=json.loads(src.read_text(encoding='utf-8-sig'))
    data['version']='Vertragsentwurf 2026.10.13 R03 - 16 Monate Bauphase und Versicherungsnachweise'
    data['preamble'].append('Fassung R03, 2026.10.13: sechzehn Monate Bauphase ab tatsächlichem B0 sind im unveränderten Paketpreis vorgeschlagen. Dauer, erweiterter Leistungsumfang und Versicherungsregelung sind gemeinsam ausdrücklich anzunehmen; dieser Entwurf belegt keine Zustimmung des AN.')
    sections={x['number']:x for x in data['sections']}
    annexes={x['number']:x for x in data['annexes']}
    s9=sections[9]['paragraphs']
    common='Die vereinbarte Bauphase umfasst sechzehn Kalendermonate ab B0 zum unveränderten Gesamtpauschalhonorar nach § 11. B0 ist der in Anlage 3 dokumentierte tatsächliche Beginn der baulichen Ausführung am Objekt; Vertragsunterzeichnung, Planungsbeginn T0, PI-Bestellung oder bloße administrative Vorbereitung sind kein B0. Vor Unterzeichnung ist der voraussichtliche B0 zu vereinbaren; der tatsächliche Beginn und das daraus errechnete Ende des sechzehnmonatigen Bauphasenrahmens werden bei Baubeginn schriftlich festgehalten. Die vor B0 erforderliche Planung ab T0 verbraucht diesen Zeitraum nicht. Die acht Betreuungstermine nach § 7 werden bedarfsgerecht über die Bauphase verteilt; gesetzliche PI-Prüfpflichten und vollständige Abschlussleistungen bleiben gesondert geschuldet.' if name=='BERGER' else 'Die vereinbarte Bauphase umfasst sechzehn Kalendermonate ab B0 zum unveränderten Gesamtpauschalhonorar nach § 11. B0 ist der in Anlage 3 dokumentierte tatsächliche Beginn der baulichen Ausführung am Objekt; Vertragsunterzeichnung, Planungsbeginn T0, PI-Bestellung oder bloße administrative Vorbereitung sind kein B0. Vor Unterzeichnung ist der voraussichtliche B0 zu vereinbaren; der tatsächliche Beginn und das daraus errechnete Ende des sechzehnmonatigen Bauphasenrahmens werden bei Baubeginn schriftlich festgehalten. Die vor B0 erforderliche Planung ab T0 verbraucht diesen Zeitraum nicht. Zwölf Abnahmen nach § 7 werden bedarfsgerecht über die Bauphase verteilt; zwölf Abnahmen sind keine zwölf Monate. Gesetzliche PI-Prüfpflichten und vollständige Abschlussleistungen bleiben gesondert geschuldet.'
    timing=' D1 ist bis drei, D2 bis sechs, D3 bis acht, D4 bis zehn und D5 bis zwölf Kalenderwochen ab T0 zu liefern, sofern in Anlage 3 keine beiderseits unterschriebenen abweichenden relativen Fristen festgelegt werden.' if name=='BERGER' else ' Vereinbart werden D1 bis drei Kalenderwochen ab T0 und ein innerhalb von fünf Werktagen ab T0 beiderseits bestätigter Lieferplan für D2-D5 vor den jeweiligen Bestellungen und Ausführungsarbeiten; die ursprünglichen Angebotsfristen waren mit der AG abzustimmen.'
    s9[1]=common+' Der bestätigte phasenweise Ablauf innerhalb dieses Rahmens begründet für sich keine Behinderungs- oder Zusatzvergütung. Gesetzlich frühere Fristen, insbesondere für das Bauwerksbuch des bestehenden Gebäudes, gehen dem Bauphasenrahmen vor; die rechtzeitige Erstellung beziehungsweise Registrierung und anschließende Fortschreibung für die neuen Eingriffe sind im Grundumfang enthalten.'+timing
    s9.append('Wird eine Überschreitung des sechzehnmonatigen Bauphasenrahmens erkennbar, legt der AN unverzüglich, bei vorhersehbarer Überschreitung spätestens vier Wochen vor dessen Ende, einen aktualisierten Ablauf- und Kostenplan vor; später erkennbare Ursachen sind unverzüglich nach Bekanntwerden anzuzeigen. Auszuweisen sind Ursache, betroffene Leistungen, Abgrenzung zum bereits geschuldeten Paket, zusätzliche Termine beziehungsweise Rollen und Stunden, ursprünglicher anwendbarer Tarif, Höchstbetrag und Terminwirkung. Der bloße Ablauf des sechzehnten Monats erzeugt weder einen monatlichen Zuschlag noch einen neuen Pauschalauftrag und beendet insbesondere PI-, Bauwerksbuch- und übliche Schlussleistungen nicht automatisch. Echte Zusatzleistungen werden nach § 11 vereinbart; nachgewiesene gesetzliche Ansprüche wegen Behinderung oder Zeitverlust aus der Sphäre der AG richten sich nach §§ 9 und 15 sowie § 1168 ABGB. Fehlende oder verspätete Kostenabstimmung ist weder stillschweigende Kostenfreigabe noch automatischer Verlust gesetzlicher Ansprüche. Gesetzlich oder vertraglich berechtigte Unterbrechungs- und Beendigungsrechte, insbesondere nach § 1170b ABGB, bleiben unberührt.')
    s11=sections[11]['paragraphs']
    s11[0]+=' Im Gesamtpreis enthalten sind die vereinbarte Planung ab T0, die übernommenen Leistungen während der gesamten sechzehnmonatigen Bauphase ab B0 sowie die zugehörigen üblichen Abschluss- und Dokumentationsleistungen, auch bei ihrer Fertigstellung nach Ende dieses Zeitrahmens. Eine vereinbarte Bauphasenlänge ist keine monatliche Vergütung; die in §§ 9, 11 und 15 geregelten echten Änderungen und gesetzlichen Behinderungsfolgen bleiben getrennt.'
    a3=annexes[3]
    a3['paragraphs'][1]+=' Die vorgesehene Zahl wird nicht automatisch erhöht, weil die Bauphase sechzehn Monate dauert; erforderliche echte Mehrtermine sind nach § 7 und § 11 zu behandeln.'
    rows=a3['tables'][0]['rows']
    idx=next(i for i,r in enumerate(rows) if r[0]=='Bauphase')
    rows[idx]=['Bauphase','Sechzehn Kalendermonate ab tatsächlichem B0 im Paket; Planung ab T0 davor und übliche Schlussleistungen danach bleiben enthalten','________________']
    rows[idx:idx]=[['B0 geplant','Vor Unterzeichnung zu vereinbarender voraussichtlicher Beginn der baulichen Ausführung','________________'],['B0 tatsächlich','Tatsächlicher Baubeginn; separat von T0 dokumentieren','________________'],['Ende des Bauphasenrahmens','Tatsächlicher B0 + sechzehn Kalendermonate; kein automatisches Vertragsende','________________']]
    a4=annexes[4]
    a4['paragraphs'][0]+=' Der Paketpreis umfasst den Bauphasenrahmen von sechzehn Monaten nach § 9; acht Betreuungstermine sind keine monatlichen Raten.' if name=='BERGER' else ' Der Paketpreis umfasst den Bauphasenrahmen von sechzehn Monaten nach § 9; zwölf Abnahmen sind keine monatlichen Raten.'
    if name=='BERGER':
        a4['paragraphs'].append('Zusatzkosten: Für einen echten zusätzlichen Betreuungstermin über acht hinaus enthält das Originalangebot keinen festen Terminpreis; Preis oder ausdrücklich freigegebene Regie mit Stunden- und Kostenobergrenze sind vor Beauftragung nach § 11 festzulegen. EUR 2.000,00 geteilt durch acht ist kein vereinbarter Zusatzterminpreis. Bereits geschuldete PI-Beschauen, normale Abschlussleistungen und eigene Fehlerkorrekturen sind kein kostenpflichtiger Zusatz allein wegen Zeitablaufs. Es besteht kein angebotener Monatszuschlag; gesetzliche Behinderungsansprüche sind nach §§ 9 und 15 gesondert nachzuweisen.')
        sections[14]['paragraphs'][3]='Als vertraglicher Ausgangswert gilt eine Deckungssumme von EUR 1.000.000,00 je Versicherungsfall für die beauftragten beruflichen Tätigkeiten, sofern nicht anhand aktueller Versicherungsnachweise eine andere Deckungssumme in Anlage 6 ausdrücklich vereinbart wird. Zwingende gesetzliche Versicherungspflichten bleiben unberührt. Die maßgeblichen Jahreshöchstleistungen, gemeinsamen Versicherungssummen, Sublimits, Selbstbehalte und wesentlichen Ausschlüsse sind anhand aktueller Nachweise in Anlage 6 offenzulegen; eine zusätzliche Verpflichtung zur zweifachen Jahresdeckung wird nicht begründet. Eine historische Polizze oder eine bloße Erklärung des AN bestätigt nicht die aktuelle Deckung der einzelnen übernommenen Funktionen.'
        a6rows=annexes[6]['tables'][0]['rows']
        r=next(r for r in a6rows if r[0]=='Deckung je Fall und Jahresmaximierung')
        r[1]='Ausgangswert EUR 1.000.000,00 je Fall; abschließend vereinbarte Summe: __________; tatsächliche Jahresmaximierung und Sublimits laut aktuellem Nachweis: __________'
        annexes[6]['paragraphs'].append('Versicherungsquellen: E-Mail des AN vom 2026.07.22 erklärt «1mio»; die übermittelte Allianz-Urkunde A559847312 vom 2016.08.09 nennt getrennt EUR 3.500.000,00 allgemeine Betriebshaftpflicht und EUR 1.000.000,00 Pflichtversicherung Baumeister nach Bedingung 6436. Diese historischen Angaben ersetzen weder aktuelle Bedingungen einschließlich 6436 noch eine Bestätigung der tatsächlich übernommenen Planung-, PI-, §-128- und Bauwerksbuchfunktionen. Gemeinsame Deckungen, Tätigkeitsgrenzen und akzeptierte Sublimits sind vor Unterzeichnung verbindlich festzuhalten; eine einheitliche Summe für jede Schadensart wird nicht ohne Nachweis unterstellt.')
    else:
        a4['paragraphs'].append('Zusatzkosten: Eine echte zusätzliche Abnahme über zwölf hinaus kostet nach § 7 EUR 450,00 netto beziehungsweise bei 20 % Umsatzsteuer EUR 540,00 brutto je Termin, einschließlich der dort vereinbarten Terminleistungen. Sechzehn Monate bedeuten nicht sechzehn Abnahmen. Nur bei vier tatsächlich erforderlichen und wirksam zusätzlich beauftragten Abnahmen entstehen rechnerisch EUR 1.800,00 netto beziehungsweise EUR 2.160,00 brutto zusätzlich; dieser Betrag ist nicht Bestandteil des vereinbarten Grundhonorars und kein automatisch beauftragter Zusatz. Bereits geschuldete PI-Leistungen, übliche Schlussleistungen und eigene Fehlerkorrekturen werden nicht nochmals als Abnahme oder Regie verrechnet. Ein Monatszuschlag wird nicht vereinbart; echte Änderungen und gesetzliche Behinderungsansprüche richten sich nach §§ 9, 11 und 15.')
    deviation=next(r for r in annexes[6]['tables'][1]['rows'] if r[0]=='Termine')
    deviation[1]='T0 für Planung und tatsächlicher B0 für die sechzehnmonatige Bauphase sind getrennt; Grundhonorar umfasst diesen Rahmen und übliche Schlussleistungen. Relative Liefer-/Buchfristen nach §§ 6 und 9 sowie Anlage 3; keine bereits begonnenen Termine. Die ursprüngliche Dauer von zirka zwölf Monaten wird ersetzt.' if name=='BERGER' else 'T0 für Planung und tatsächlicher B0 für die sechzehnmonatige Bauphase sind getrennt; Grundhonorar umfasst diesen Rahmen und übliche Schlussleistungen. Zwölf angebotene Abnahmen bleiben zwölf Termine. Relative Liefer-/Buchfristen nach §§ 6 und 9 sowie Anlage 3; keine bereits begonnenen Termine. Im Originalangebot war keine Monatsdauer zugesagt.'
    alltext=json.dumps(data,ensure_ascii=False)
    for forbidden in ['rund zwölf Monate Bauphase','bis fünfzehn Monate','zwölfmonatige Bauphase','Kalkulation zwölf Monate','EUR 2.000.000,00 je Fall','mindestens zweifach']:
        assert forbidden not in alltext,(name,forbidden)
    assert data['financial_check']==old_data['financial_check']
    outdir=ROOT/f'00.Proiect/09. Vertraege/{domain}/2026.10.13 Juristischer Werkvertrag'
    newj=WORK/f'2026.10.13 Werkvertrag {name} R03.json'
    writej(newj,data)
    docx=outdir/f'2026.10.13 Werkvertrag {name} R03.docx'
    assert not docx.exists()
    builder.build(newj,docx)
    d=Document(docx)
    # Preserve R02 layout correction: no empty body paragraph before the forced Annex 5 page break.
    for p in list(d.paragraphs):
        if p.text.startswith('Anlage 5 '):
            prev=p._p.getprevious()
            if prev is not None and prev.tag.endswith('}p') and not ''.join(prev.xpath('.//w:t/text()')):
                prev.getparent().remove(prev)
    for section in d.sections:
        for p in section.footer.paragraphs:
            for r in p.runs:
                if 'Entwurf 2026.10.13' in r.text:r.text=r.text.replace('Entwurf 2026.10.13','Entwurf 2026.10.13 R03')
    d.save(docx)
    # Semantic fidelity and unchanged fee tables against the authored JSON.
    txt=''.join(x.text or '' for el in d.element.body for x in el.xpath('.//w:t'))
    for container in data['sections']+data['annexes']:
        for p in container.get('paragraphs',[]):assert p in txt,p[:80]
        for table in container.get('tables',[]):
            for row in [table['headers']]+table['rows']:
                for c in row:assert str(c) in txt,str(c)
    records.append({'prestator':name,'docx':docx.relative_to(ROOT).as_posix(),'json':newj.relative_to(ROOT).as_posix(),'sha256_json':sha(newj),'sha256_docx':sha(docx),'net_EUR':45000 if name=='BERGER' else 69300})
    changes.append({'prestator':name,'sursa_R02_sha256':sha(src),'R03_sha256':sha(newj),'sectiuni_modificate': [9,11,14] if name=='BERGER' else [9,11],'anexe_modificate':[3,4,6],'preturi_originale_pastrate':True,'Bauphase_luni':16,'B0_separat_de_T0':True,'statut':'NETRIMIS NESEMNAT; extinderea este propunere AG, neacceptată'})
writej(WORK/'2026.10.13 Export contracte R03.json',{'data':'2026.10.13','documente':records})
writej(WORK/'2026.10.13 Modificari contracte R03.json',{'data':'2026.10.13','revizii':changes})
print(json.dumps(records,ensure_ascii=False))
