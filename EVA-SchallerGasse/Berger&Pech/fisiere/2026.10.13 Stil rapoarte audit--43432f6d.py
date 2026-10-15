from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
def field(p,key):
    for kind,val in [('begin',None),('instr',key),('end',None)]:
        r=p.add_run()
        e=OxmlElement('w:instrText' if kind=='instr' else 'w:fldChar')
        if kind=='instr':e.text=' '+val+' ';e.set(qn('xml:space'),'preserve')
        else:e.set(qn('w:fldCharType'),kind)
        r._r.append(e)
def newdoc(title,subject):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21.59);s.page_height=Cm(27.94)
    s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.9);s.left_margin=Cm(2.0);s.right_margin=Cm(2.0)
    s.header_distance=Cm(.7);s.footer_distance=Cm(.7)
    st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.06;st.paragraph_format.space_after=Pt(5);st.paragraph_format.widow_control=True
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ro-RO');st._element.get_or_add_rPr().append(lang)
    for n,z in [('Title',17),('Heading 1',13),('Heading 2',11.5)]:
        st=d.styles[n];st.font.name='Arial';st.font.size=Pt(z);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
        if st._element.pPr is not None:
            for e in list(st._element.pPr):
                if e.tag==qn('w:pBdr'):st._element.pPr.remove(e)
    d.core_properties.title=title;d.core_properties.author='A&C Wohnart — analiză cu auditori AI independenți';d.core_properties.subject=subject
    d.add_paragraph(title,'Title')
    d.add_paragraph('Schallergasse 35, 1120 Wien  |  2026.10.13  |  Contracte R05 înghețate pentru analiză')
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('2026.10.13  |  Audit separat  |  ');field(footer,'PAGE');footer.add_run(' / ');field(footer,'NUMPAGES')
    for r in footer.runs:r.font.name='Arial';r.font.size=Pt(8.5);r.font.color.rgb=RGBColor(0,0,0)
    return d
def para(d,t,bold_label=None):
    p=d.add_paragraph()
    if bold_label:
        p.add_run(bold_label+' ').bold=True;p.add_run(str(t))
    else:p.add_run(str(t))
    p.paragraph_format.widow_control=True
    return p
def table(d,headers,rows,widths):
    tb=d.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    fac=17.59/sum(widths);widths=[w*fac for w in widths]
    for i,w in enumerate(widths):tb.columns[i].width=Cm(w)
    for i,x in enumerate(headers):tb.rows[0].cells[i].text=str(x)
    for vals in rows:
        assert len(vals)==len(headers)
        for c,v in zip(tb.add_row().cells,vals):c.text=str(v)
    for ri,row in enumerate(tb.rows):
        rp=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');rp.append(cs)
        if ri==0:
            h=OxmlElement('w:tblHeader');h.set(qn('w:val'),'true');rp.append(h)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            cp.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EBEF' if ri==0 else ('F7F8FA' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            m=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
            cp.append(m)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.02;p.paragraph_format.space_before=Pt(1.5);p.paragraph_format.space_after=Pt(1.5)
                p.paragraph_format.keep_with_next=ri==0;p.paragraph_format.widow_control=True
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.name='Arial';r.font.size=Pt(9.5);r.font.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(2)
    return tb
def source(d,title,url):
    p=para(d,title+' — ')
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),d.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);rr.append(pr)
    tx=OxmlElement('w:t');tx.text=url;rr.append(tx);h.append(rr);p._p.append(h)
    for r in p.runs:r.font.size=Pt(9)
    return p
def save(d,path):
    assert not path.exists(),str(path)
    d.save(path)
    # No content corrections are made to the frozen contracts by this module.
    return path
