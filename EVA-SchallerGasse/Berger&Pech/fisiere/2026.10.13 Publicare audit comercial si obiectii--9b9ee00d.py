from pathlib import Path
import json,hashlib,importlib.util
from decimal import Decimal
from docx import Document
ROOT=Path(r'D:\00. Downloads\Apartamente Viena\00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
WORK=ROOT/'00.Proiect/09. Vertraege/2026.10.13 Reorganizare juridica si audit'
VIS=Path(r'C:\Users\User\.codex\visualizations\2026\10\01\01a0f6eb-c105-7760-973e-241d6f6be8c1')
sp=importlib.util.spec_from_file_location('sty',VIS/'2026.10.13 Stil rapoarte audit.py');sty=importlib.util.module_from_spec(sp);sp.loader.exec_module(sty)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
cp=WORK/'2026.10.13 Audit comercial comparativ R05 final.json';assert sha(cp)=='0ab9f9510bd6e8a920a3bdecee3d18fe27221abde30f250d16b8c4835e927b8b'
c=json.loads(cp.read_text(encoding='utf-8-sig'));j=json.loads((WORK/'2026.10.13 Checklist juridic comparativ R05.json').read_text(encoding='utf-8-sig'))
frozen=json.loads((WORK/'2026.10.13 Surse inghetate audit R05.json').read_text(encoding='utf-8-sig'))
for x in frozen['surse']:assert sha(ROOT/x['cale'])==x['sha256']
labels={'conform ofertei':'Conform ofertei','AG nou':'Condiție AG nouă','conflict':'Conflict / delimitare','requires proof':'Necesită probă','AG nou / conflict':'Condiție AG nouă / conflict'}
def status(x):return labels.get(str(x),str(x))
def money(x):return f'{Decimal(str(x)):,.2f}'.replace(',','X').replace('.',',').replace('X','.')
def src(s):
    if isinstance(s,list):return '; '.join(src(x) for x in s)
    if isinstance(s,dict):return str(s.get('id',''))+(' — '+str(s['pagina']) if s.get('pagina') else '')
    return str(s or '')
def p(d,t):return sty.para(d,t)
def h(d,t):d.add_paragraph(t,'Heading 1')
def findings(d,items):
    for x in items:
        d.add_paragraph(str(x['id'])+' — '+x['tema']+' — '+x.get('prestator',''),'Heading 2')
        for key,label in [('severitate','Prioritate'),('referinta','Clauză'),('constatare','Constatare'),('efect','Efect'),('text_propus_doar_observatie','Formulare propusă, neimplementată'),('recomandare','Recomandare neimplementată')]:
            if x.get(key):sty.para(d,x[key],label+':')
        if x.get('sursa'):sty.para(d,src(x['sursa']),'Sursă:')
d=sty.newdoc('Audit comercial comparativ și checklist','BERGER / PECH — ofertă → contract și contract → sursă; numai sesizări R05')
p(d,'Auditor: rol AI comercial independent. Verificarea juridică a celuilalt auditor este în raport separat. Contractele R05 au rămas înghețate; toate formulările propuse în acest raport sunt observații, fără introducere în contracte după instrucțiunea utilizatorului.')
h(d,'1. Concluzii și documente acoperite')
for line in c['concluzii']:p(d,line)
p(d,'Acoperire: două oferte comerciale complete distincte, câte patru pagini, 12 aliasuri cu hash identic; 14 mesaje arhivate (7 pentru fiecare partener, 5 primite și 9 trimise AG), fără corp marcat trunchiat în selecție. Registrul surselor conține '+str(len(c['documente_verificate']))+' documente. Comparația are 60 de teme, 44 grupuri inverse și 401 unități de alineat/rând trasabile în JSON/TXT. Reperul TOMS servește comparației de scop, fără a transfera acceptarea sau prețurile lui asupra BERGER/PECH.')
sty.table(d,['Control numeric','BERGER','PECH'],[
['Poziții originale nete (EUR)',', '.join(money(x) for x in c['control_numeric']['BERGER']['pozitii_net']),', '.join(money(x) for x in c['control_numeric']['PECH']['pozitii_net'])],
['Total net al pozițiilor / tranșelor','45.000,00 / 45.000,00','69.300,00 / 69.300,00'],
['TVA / brut, agregat','9.000,00 / 54.000,00','13.860,00 / 83.160,00'],
['Durată în R05','18 luni ab B0; T0 pentru planificare separat','18 luni ab B0; T0 pentru planificare separat'],
['Original → R05, vizite','8 → 12 incluse; PI distinct','12 → 12 incluse'],
['Până la 6 opțiuni','Formula AG nouă 2.000/12 EUR net','450 EUR net per termen, oferit'],
['Toate cele 6, net / brut','1.000,00 / 1.200,00','2.700,00 / 3.240,00'],
['Bază + toate 6, net / brut','46.000,00 / 55.200,00','72.000,00 / 86.400,00']],[4.2,6.69,6.7])
p(d,'Sumele brute exemplifică 20% TVA calculată pe total. Rotunjirea fiscală a facturilor individuale rămâne aplicabilă. Cele șase opțiuni nu sunt comandate anticipat și nu sunt un plafon universal pentru schimbări reale, servicii externe sau creanțe legale distincte.')
sty.table(d,['n vizite suplimentare','BERGER cumul net','BERGER total net','PECH cumul net','PECH total net'],[[str(x['n']),money(x['berger_cumul_net']),money(x['berger_total_net']),money(x['pech_cumul_net']),money(x['pech_total_net'])] for x in c['optiuni_0_6']],[3.19,3.6,3.6,3.6,3.6])
p(d,'BERGER: rotunjire cumulativă la cent, conform formulei din contract; șase afișaje de 166,67 nu se multiplică pentru a factura 1.000,02. Suma exactă a celor șase este 1.000 EUR net. PECH: 6 × 450 = 2.700 EUR net.')
h(d,'2. Cele cinci observații comerciale rămase')
findings(d,c['observatii'])
h(d,'3. Checklist comparativ: fiecare aspect din ofertă și cererea conexă')
p(d,'În fiecare celulă sunt separate condiția inițială, textul contractual și calificarea. Un preț numeric păstrat nu certifică identitatea scopului. Sursa indicată prin cod și pagină este identificată cu calea și hashul în registrul final; citatele originale complete rămân în JSON.')
for start in range(0,len(c['checklist']),15):
    subset=c['checklist'][start:start+15];d.add_paragraph('Temele '+subset[0]['id']+'–'+subset[-1]['id'],'Heading 2')
    rows=[]
    for x in subset:
        cells=[]
        for provider in ['berger','pech']:
            y=x[provider]
            cells.append('OFERTĂ: '+y.get('oferta','')+'\nCONTRACT: '+y.get('contract','')+'\nSTATUT: '+status(y.get('statut',''))+'\nSURSA: '+src(y.get('sursa','')))
        rows.append([x['id']+' — '+x['tema'],cells[0],cells[1],x.get('severitate','')+'\nRecomandare neimplementată: '+x.get('corectie','')])
    sty.table(d,['Tema','BERGER — ofertă / contract','PECH — ofertă / contract','Prioritate / recomandare'],rows,[2.6,5.6,5.6,3.79])
h(d,'4. Contract → ofertă: toate capitolele, anexele și componentele')
p(d,'Fiecare grup de mai jos include toate componentele mapate. Lista indică referința, statutul și codul sursei pentru fiecare unitate; textul german integral al celor 401 unități este păstrat în JSON/TXT. Observația de grup se citește împreună cu cele 60 de teme și auditul juridic pe alineate.')
unit_count=0
for provider in ['BERGER','PECH']:
    d.add_paragraph(provider,'Heading 2');rows=[]
    for x in c['mapare_inversa']:
        if x['prestator']!=provider:continue
        pieces=[]
        for unit in x['componente']:
            unit_count+=1
            pieces.append(unit['referinta']+' — '+status(unit['statut'])+' ['+src(unit.get('sursa',''))+']')
        rows.append([x['referinta']+'\n'+x['tema'],'\n'.join(pieces),x.get('observatie','')])
    sty.table(d,['Capitol / anexă','Componente și proveniență','Observație comercială'],rows,[3.6,8.2,5.79])
assert unit_count==401,unit_count
h(d,'5. Reperul de scop TOMS')
p(d,'Cele 47 de repere sunt lista de comparație din arhivă. Prețurile și confirmările TOMS nu sunt condiții acceptate de BERGER ori PECH. Primele patru sunt istoricul negocierii; BauKG și managementul general REALiG rămân excluse în R05.')
sty.table(d,['Reper','Referință R05','Calificare'],[[x['id']+' — '+x['tema'],x['referinta_R05'],x['statut']] for x in c['reper_TOMS']],[9.4,3.2,4.99])
h(d,'6. Istoricul ofertelor și registrul surselor')
for provider in ['BERGER','PECH']:
    x=c['istoric_oferte'][provider]
    p(d,provider+': '+x['data']+'; '+x['referinta']+'; expirare calculată '+x['expirare_calculata']+'; ultimă comunicare externă arhivată '+x['ultima_comunicare_externa']+'. O durată contractuală de 18 luni nu prelungește valabilitatea ofertei.')
for source_index,x in enumerate(c['documente_verificate'],1):
    d.add_paragraph(x.get('id','DOC'+str(source_index))+' — '+x.get('data','')+' — '+x.get('tip',''),'Heading 2')
    p(d,x['cale']);p(d,'SHA-256: '+x.get('hash',''))
    if x.get('verificare'):p(d,x['verificare'])
h(d,'7. Limite și condiții de interpretare')
for line in c['limite']:p(d,line)
p(d,'QA vizual separat: toate cele 45 de pagini ale contractelor înghețate au fost inspectate; nota QA R05 este în același dosar. Acest control de aspect nu rezolvă acceptarea, autorizarea persoanelor, dovada asigurării sau observațiile de fond.')
comout=WORK/'2026.10.13 Audit comercial comparativ R05.docx';sty.save(d,comout)
body=''.join(t.text or '' for e in Document(comout).element.body for t in e.xpath('.//w:t'))
for x in c['checklist']:
    assert x['id'] in body
    for provider in ['berger','pech']:
        for field in ['oferta','contract']:assert x[provider][field] in body,(x['id'],provider,field)
for g in c['mapare_inversa']:
    for unit in g['componente']:assert unit['referinta'] in body,unit['referinta']
o=sty.newdoc('Obiecții posibile ale ofertanților','BERGER / PECH — observații anticipate de doi auditori AI independenți')
p(o,'Acestea sunt obiecții anticipate de auditori, pentru evaluarea negocierii. Fiecare are o bază documentară și o poziție recomandată. Sunt separate de răspunsurile efectiv primite. Contractele R05 sunt înghețate; nicio recomandare de mai jos nu a fost introdusă în ele după cererea utilizatorului.')
p(o,'Pachet: 18 luni ab B0, 12 vizite incluse și cel mult 6 opțiuni plătite. BERGER: baza 45.000 EUR net și formula AG nouă 2.000/12; PECH: baza 69.300 EUR net și supliment 450 EUR net. Menținerea prețului nu dovedește acceptarea extinderii.')
total=0
for role,items in [('Auditorul comercial — 18 obiecții',c['obiectii']),('Auditorul juridic — 12 obiecții',j['obiectii'])]:
    h(o,role)
    for i,x in enumerate(items,1):
        total+=1;o.add_paragraph(str(i)+'. '+x['prestator']+' — '+x['tema'],'Heading 2')
        for key,label in [('obiectie_probabila','Obiecție posibilă'),('evaluare','Evaluare echilibrată'),('pozitie_recomandata','Poziție recomandată')]:sty.para(o,x[key],label+':')
        if x.get('sursa'):sty.para(o,src(x['sursa']),'Sursă:')
assert total==30
h(o,'Documente de referință')
p(o,'Auditul comercial comparativ R05 și auditul juridic în ambele sensuri din același dosar conțin checklisturile complete, sursele cu hash, textele normative, jurisprudența și limitele verificării. O obiecție comercială nu reprezintă automat nulitate; o clauză acceptată nu dovedește automat abilitatea ori asigurarea necesară pentru exercitarea funcției.')
p(o,'Punctele deschise centrale: delimitarea PECH Pos.04/05; Materialprüfung versus teste externe; versiunea cererii tehnice și limitele reziduale; acceptarea scopului și calendarului; documentele profesionale și asigurarea actuală; valabilitatea calculat expirată a ofertei BERGER.')
obsout=WORK/'2026.10.13 Obiectii posibile BERGER PECH R05.docx';sty.save(o,obsout)
for x in frozen['surse']:assert sha(ROOT/x['cale'])==x['sha256']
assert sha(cp)=='0ab9f9510bd6e8a920a3bdecee3d18fe27221abde30f250d16b8c4835e927b8b'
manifest={'data':'2026.10.13','documente':[{'prestator':'Audit comercial','docx':comout.relative_to(ROOT).as_posix()},{'prestator':'Obiecții posibile','docx':obsout.relative_to(ROOT).as_posix()}]}
mf=WORK/'2026.10.13 Export rapoarte comercial si obiectii.json';assert not mf.exists();mf.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'commercial':str(comout),'commercial_sha256':sha(comout),'objections':str(obsout),'objections_sha256':sha(obsout),'teme':60,'grupuri_inverse':44,'unitati':unit_count,'obiectii':total},ensure_ascii=True))
