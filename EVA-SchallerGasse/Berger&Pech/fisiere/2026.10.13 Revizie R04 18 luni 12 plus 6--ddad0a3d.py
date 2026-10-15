from pathlib import Path
import json,hashlib,importlib.util
from docx import Document
ROOT=Path(r'D:\00. Downloads\Apartamente Viena\00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
WORK=ROOT/'00.Proiect/09. Vertraege/2026.10.13 Reorganizare juridica si audit'
VIS=Path(r'C:\Users\User\.codex\visualizations\2026\10\01\01a0f6eb-c105-7760-973e-241d6f6be8c1')
spec=importlib.util.spec_from_file_location('builder',VIS/'2026.10.13 Construire contracte juridice.py');builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def writej(p,v):
    assert not p.exists(),str(p)
    p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def horizon(v):
    if isinstance(v,str):return v.replace('sechzehn','achtzehn').replace('Sechzehn','Achtzehn').replace('16 Monate','18 Monate')
    if isinstance(v,list):return [horizon(x) for x in v]
    if isinstance(v,dict):return {k:horizon(x) for k,x in v.items()}
    return v
records=[]
for name,domain in [('BERGER','bmberger.at'),('PECH','zt-pech.at')]:
    src=WORK/f'2026.10.13 Werkvertrag {name} R03.json'
    old=json.loads(src.read_text(encoding='utf-8-sig'))
    data=horizon(old)
    data['version']='Vertragsentwurf 2026.10.13 R04 - 18 Monate Bauphase; 12 Termine und 6 Optionen'
    data['preamble'][-1]='Fassung R04, 2026.10.13: achtzehn Monate Bauphase ab tatsächlichem B0 und zwölf enthaltene Ortstermine sind im unveränderten Paketpreis vorgeschlagen; bis zu sechs weitere Termine sind nur bei wirksamer Einzelbeauftragung zusätzlich zu vergüten. Der erweiterte Umfang einschließlich Dauer, Preisregelung und Versicherungsnachweisen ist gemeinsam ausdrücklich anzunehmen; dieser Entwurf belegt keine Zustimmung des AN.'
    ss={x['number']:x for x in data['sections']};aa={x['number']:x for x in data['annexes']}
    if name=='BERGER':
        ss[7]['paragraphs'][0]=ss[7]['paragraphs'][0].replace('acht geplante Ortstermine','zwölf geplante Ortstermine').replace('der weitere Termin wird nach Bedarf festgelegt','weitere enthaltene Termine werden für Bestandswände/Anschlüsse, Holz-/Stahldetails, tragwerksrelevanten Ausbau und Abschluss nach Anlage 3 festgelegt')
        ss[7]['paragraphs'][3]='Die AG kann zusätzlich zu den zwölf enthaltenen Betreuungsterminen bis zu sechs weitere echte Betreuungstermine nach Bedarf einzeln beauftragen. Der Preis ergibt sich je Termin aus EUR 2.000,00 / zwölf, als gerundeter Informationswert rund EUR 166,67 netto; die verbindliche kumulative Berechnung und Rundung regelt Anlage 4. Enthalten sind Vorbereitung, An- und Abreise, Terminwahrnehmung, Protokoll und übliche unmittelbare Nachbearbeitung. Anlass, Bauteil, Termin und Beauftragung durch eine befugte Person nach § 16 sind vor Durchführung nach § 11 zu dokumentieren. Die Option ist keine Vorausbestellung; nur tatsächlich beauftragte und ausgeführte Zusatztermine werden vergütet. Eigene Fehlerkorrekturen und bereits im PI-Paket geschuldete Beschauen sind keine solche Zusatzleistung; dieselbe Leistung darf nicht zusätzlich nach Regie oder als Nebenkosten verrechnet werden. Diese zwölf enthaltenen Termine und der abgeleitete Optionspreis sind neue Vertragsregelungen, keine im Originalangebot bereits zugesagten Bedingungen.'
        ss[9]['paragraphs'][1]=ss[9]['paragraphs'][1].replace('Die acht Betreuungstermine','Die zwölf Betreuungstermine')
        aa[3]['paragraphs'][1]=aa[3]['paragraphs'][1].replace('acht Betreuungsbesprechungen','zwölf Betreuungsbesprechungen')
        aa[3]['tables'][1]['rows']=[['1','Startbesprechung'],['2','Fundamentplatte'],['3','Aufzugsschacht'],['4','Deckenertüchtigung und Kernbohrungen'],['5','Bestandswände und Auswechslungen'],['6','Dachstuhl'],['7','Holz-/Stahl- und Bestandsanschlüsse'],['8','Balkon-/Anschlussphase'],['9','Rohbauabschluss'],['10','Tragwerksrelevante Ausbauphase'],['11','Abgestimmter Bedarfs-/Koordinierungstermin'],['12','Abschluss und Übergabekoordinierung']]
        aa[4]['paragraphs'][0]=aa[4]['paragraphs'][0].replace('acht Betreuungstermine','zwölf Betreuungstermine')
        for tb in aa[4]['tables']:
            for row in tb['rows']:
                for i,c in enumerate(row):
                    if isinstance(c,str):row[i]=c.replace('Acht dokumentierte Betreuungstermine','Zwölf dokumentierte Betreuungstermine')
        original_count='Das Originalangebot nennt maximal acht Betreuungstermine für EUR 2.000,00 netto. Zwölf enthaltene Termine bei unverändertem Grundhonorar und der Preis der bis zu sechs Optionen werden mit diesem Vertrag ausdrücklich neu vereinbart; das Original enthält keinen festen Zusatzterminpreis.'
        price_rule='Für n wirksam zusätzlich beauftragte und ausgeführte Betreuungstermine, mit n von null bis sechs, beträgt das kumulative Netto-Optionsentgelt n × EUR 2.000,00 / zwölf. Die Berechnung erfolgt mit dem ungerundeten Verhältnis; erst der kumulative Rechnungsstand wird kaufmännisch auf Cent gerundet. Bereits für diese Optionen verrechnete Nettoentgelte werden abgezogen. Der gerundete Anzeigewert von EUR 166,67 ist nicht als sechsfach gerundeter Festpreis zu multiplizieren. Bei sechs Terminen sind insgesamt höchstens EUR 1.000,00 netto zu vergüten, bei 20 % Umsatzsteuer auf diesen Gesamtbetrag EUR 1.200,00 brutto; die gesetzlichen Regeln zur Umsatzsteuer je Rechnung bleiben maßgeblich.'
        cap=['Bis zu 6 Zusatztermine','n × EUR 2.000,00 / 12; rund EUR 166,67 je Termin','1.000,00','1.200,00']
        total=['Grundhonorar + alle 6 Optionen','Nur wenn alle sechs wirksam beauftragt und ausgeführt','46.000,00','55.200,00']
        visits_deviation='Zwölf statt ursprünglich maximal acht enthaltene Betreuungstermine; bis zu sechs einzeln beauftragbare echte Zusatztermine zum abgeleiteten Preis EUR 2.000,00 / zwölf je Termin nach Anlage 4. Diese Erweiterung und der Tarif sind neue AG-Vertragsvorschläge. Erforderliche im Paket enthaltene PI-Beschauen werden gesondert erbracht und nicht nochmals als Zusatztermin verrechnet.'
    else:
        ss[7]['paragraphs'][3]='Die AG kann zusätzlich zu den zwölf enthaltenen Abnahmen bis zu sechs weitere echte Abnahmen nach Bedarf einzeln zu EUR 450,00 netto je Termin, bei 20 % Umsatzsteuer EUR 540,00 brutto, beauftragen. Enthalten sind Vorbereitung, An- und Abreise, Terminwahrnehmung, Protokoll und übliche unmittelbare Nachbearbeitung. Anlass, Bauteil, Termin und Beauftragung durch eine befugte Person nach § 16 sind vor Durchführung nach § 11 zu dokumentieren. Die Option ist keine Vorausbestellung; nur tatsächlich beauftragte und ausgeführte Zusatztermine werden vergütet. Eigene Fehlerkorrekturen und bereits im PI-Paket geschuldete Beschauen sind keine solche Zusatzleistung. Dieselbe Leistung darf nicht zusätzlich nach Regie oder als Nebenkosten verrechnet werden. Bei vollständiger Beauftragung und Ausführung aller sechs Optionen beträgt der Zuschlag EUR 2.700,00 netto beziehungsweise EUR 3.240,00 brutto bei 20 % Umsatzsteuer.'
        original_count='Das Originalangebot umfasst zwölf Abnahmen in Position 05 und EUR 450,00 netto je weiterer Abnahme. Diese Grundanzahl und der angebotene Zusatzterminpreis bleiben erhalten; die Option von bis zu sechs Terminen, der achtzehnmonatige Bauphasenrahmen und der erweiterte Leistungsumfang werden ausdrücklich vereinbart.'
        price_rule='Für n wirksam zusätzlich beauftragte und ausgeführte Abnahmen, mit n von null bis sechs, beträgt das Netto-Optionsentgelt n × EUR 450,00. Höchstens EUR 2.700,00 netto beziehungsweise bei 20 % Umsatzsteuer auf diesen Gesamtbetrag EUR 3.240,00 brutto entstehen aus diesen sechs Optionen. Die gesetzlichen Regeln zur Umsatzsteuer je Rechnung bleiben maßgeblich. EUR 14.000,00 / zwölf ist kein vereinbarter Zusatzterminpreis; der ausdrücklich angebotene Preis von EUR 450,00 ist maßgeblich.'
        cap=['Bis zu 6 Zusatzabnahmen','450,00 netto je tatsächlich ausgeführtem Zusatztermin','2.700,00','3.240,00']
        total=['Grundhonorar + alle 6 Optionen','Nur wenn alle sechs wirksam beauftragt und ausgeführt','72.000,00','86.400,00']
        visits_deviation='Zwölf enthaltene Abnahmen; bis zu sechs einzeln beauftragbare echte Zusatzabnahmen zu EUR 450,00 netto je Termin. Grundanzahl und Zusatzpreis stammen aus dem Originalangebot; die Option, der Gesamtumfang und der Bauphasenrahmen sind ausdrücklich anzunehmen. Gemeinsame Terminwahrnehmung und bereits geschuldete PI-Leistungen werden nicht doppelt verrechnet.'
    ss[7]['paragraphs'].append('Die Option von bis zu sechs Zusatzterminen beschränkt keine gesetzlich oder sicherheitsbedingt notwendige Prüfung. Weiterer echter Zusatzbedarf über diese Option hinaus wird rechtzeitig mit Ursache, Leistungsabgrenzung, Preis und Terminwirkung nach § 11 gesondert vereinbart; aus dem Ablauf von achtzehn Monaten oder dem Verbrauch von achtzehn geplanten Terminen folgt weder eine Prüfung als erledigt noch ein automatisches Vertragsende. Unaufschiebbare Gefahrenabwehr und gesetzlich begründete Vergütungsfolgen bleiben nach §§ 9, 11 und 15 unberührt.')
    aa[3]['paragraphs'].append('Bis zu sechs optionale echte Zusatztermine werden nur einzeln nach § 7 und § 11 beauftragt und zusätzlich zum Grundhonorar bezahlt. Sie sind in der laufenden Terminliste mit Auftrag, Leistung, Befund und Preis getrennt von den zwölf enthaltenen Terminen und den gesonderten PI-Pflichten zu erfassen. Keine Monatszählung löst eine automatische Zusatzbestellung aus.')
    aa[4]['paragraphs']=[p for p in aa[4]['paragraphs'] if not p.startswith('Zusatzkosten:')]
    aa[4]['paragraphs'] += [original_count,price_rule,'Die nachfolgende Optionssumme ist ein bedingtes Budget und erhöht das Grundhonorar nicht ohne wirksame Beauftragung und tatsächliche Ausführung. Keine Kosten entstehen für ungenutzte Optionen. Echte andere Leistungsänderungen, Untersuchungs-/Fremdleistungen und gesetzlich begründete Behinderungsansprüche sind hiervon getrennt nach §§ 8, 9, 11 und 15 zu beurteilen; der Optionshöchstbetrag ist kein pauschaler Verzicht auf diese anderen Ansprüche.']
    aa[4]['tables'].append({'headers':['Kostenfall','Preisgrundlage','Netto EUR','Brutto EUR bei 20 %'],'widths_cm':[3.3,7.1,3.1,3.1],'rows':[cap,total]})
    for row in aa[6]['tables'][1]['rows']:
        if row[0]=='Ortstermine':row[1]=visits_deviation
    assert data['financial_check']==old['financial_check']
    text=json.dumps(data,ensure_ascii=False)
    assert 'sechzehn' not in text.lower()
    for bad in ['acht geplante Ortstermine','acht enthaltenen Termine','acht Betreuungsbesprechungen','Acht dokumentierte Betreuungstermine','Die acht Betreuungstermine']:
        assert bad not in text,(name,bad)
    pathj=WORK/f'2026.10.13 Werkvertrag {name} R04.json'
    writej(pathj,data)
    outdir=ROOT/f'00.Proiect/09. Vertraege/{domain}/2026.10.13 Juristischer Werkvertrag'
    docx=outdir/f'2026.10.13 Werkvertrag {name} R04.docx'
    assert not docx.exists()
    builder.build(pathj,docx)
    d=Document(docx)
    for p in list(d.paragraphs):
        if p.text.startswith('Anlage 5 '):
            prev=p._p.getprevious()
            if prev is not None and prev.tag.endswith('}p') and not ''.join(prev.xpath('.//w:t/text()')):prev.getparent().remove(prev)
    for sec in d.sections:
        for p in sec.footer.paragraphs:
            for run in p.runs:
                if 'Entwurf 2026.10.13' in run.text:run.text=run.text.replace('Entwurf 2026.10.13','Entwurf 2026.10.13 R04')
    d.save(docx)
    check=Document(docx);body=''.join(x.text or '' for e in check.element.body for x in e.xpath('.//w:t'))
    for c in data['sections']+data['annexes']:
        for p in c.get('paragraphs',[]):assert p in body,p[:80]
        for tb in c.get('tables',[]):
            for r in [tb['headers']]+tb['rows']:
                for cell in r:assert str(cell) in body,str(cell)
    records.append({'prestator':name,'docx':docx.relative_to(ROOT).as_posix(),'json':pathj.relative_to(ROOT).as_posix(),'sha256_json':sha(pathj),'sha256_docx':sha(docx),'net_EUR':45000 if name=='BERGER' else 69300,'TVA_EUR':9000 if name=='BERGER' else 13860,'brut_EUR':54000 if name=='BERGER' else 83160,'luni_Bauphase':18,'vizite_incluse':12,'vizite_optionale_maxim':6,'optiuni_max_net_EUR':1000 if name=='BERGER' else 2700,'total_cu_toate_optiunile_net_EUR':46000 if name=='BERGER' else 72000,'total_cu_toate_optiunile_brut_EUR':55200 if name=='BERGER' else 86400,'statut':'NETRIMIS NESEMNAT; propunere AG neacceptată'})
writej(WORK/'2026.10.13 Export contracte R04.json',{'data':'2026.10.13','documente':records})
print(json.dumps(records,ensure_ascii=False))
