from pathlib import Path
from docx import Document
import json,hashlib
ROOT=Path(r'D:\00. Downloads\Apartamente Viena\00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
SOURCE=ROOT/'04. Firme + Executie/2026.10.12 Propuneri contracte Berger si Pech'
out=[]
def blocks(d):
 return [''.join(x.text or '' for x in el.xpath('.//w:t')) for el in d.element.body]
for name,domain,start in [('BERGER','bmberger.at',15),('PECH','zt-pech.at',12)]:
 original=SOURCE/f'2026.10.12 Propunere contract {name}.docx'
 d=Document(original)
 orig_blocks=blocks(d)
 assert orig_blocks[start]==f'Werkvertrag {name}'
 assert all(x.tag.endswith('}p') for x in list(d.element.body)[:start])
 expected=orig_blocks[start:]
 for el in list(d.element.body)[:start]: d.element.body.remove(el)
 d.paragraphs[0].paragraph_format.page_break_before=False
 dest=ROOT/'00.Proiect/09. Vertraege'/domain/'2026.10.13 Propunere contract'
 target=dest/f'2026.10.13 Werkvertrag {name} DE.docx'
 assert not target.exists(),target
 d.save(target)
 check=Document(target)
 assert blocks(check)==expected
 assert len(check.tables)==len(Document(original).tables)
 out.append(dict(prestator=name,sursa=original.relative_to(ROOT).as_posix(),docx=target.relative_to(ROOT).as_posix(),paragrafe_intro_eliminate=start,tabele_pastrate=len(check.tables),blocuri_germane_pastrate=len(expected),sha256=hashlib.sha256(target.read_bytes()).hexdigest(),modificare='Numai eliminarea introducerii și sintezei române; început fără salt de pagină inițial. Textul și tabelele contractuale germane păstrate exact.',statut='NETRIMIS și NESEMNAT; acceptare neconfirmată'))
manifest=ROOT/'00.Proiect/09. Vertraege/2026.10.13 Versiuni contracte numai germana.json'
manifest.write_text(json.dumps(dict(data='2026.10.13',documente=out),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2))

