from pathlib import Path
import importlib.util,json,re,hashlib
from docx import Document
ROOT=Path(r'D:\00. Downloads\Apartamente Viena\00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
WORK=ROOT/'00.Proiect/09. Vertraege/2026.10.13 Reorganizare juridica si audit'
VIS=Path(r'C:\Users\User\.codex\visualizations\2026\10\01\01a0f6eb-c105-7760-973e-241d6f6be8c1')
sp=importlib.util.spec_from_file_location('sty',VIS/'2026.10.13 Stil rapoarte audit.py');sty=importlib.util.module_from_spec(sp);sp.loader.exec_module(sty)
data=json.loads((WORK/'2026.10.13 Checklist juridic comparativ R05.json').read_text(encoding='utf-8-sig'))
frozen=json.loads((WORK/'2026.10.13 Surse inghetate audit R05.json').read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
for x in frozen['surse']:assert sha(ROOT/x['cale'])==x['sha256']
assert len(data['checklist'])==97 and len(data['oferte_spre_contract'])==80
def no_url(s):return re.sub(r'\s*RIS:\s*https?://\S+','',str(s)).strip()
def severity(s):return str(s).split(' — ')[0]
def h(d,t):d.add_paragraph(t,'Heading 1')
def p(d,t):return sty.para(d,t)
d=sty.newdoc('Audit juridic și checklist în ambele sensuri','BERGER / PECH — sesizări distincte; fără modificarea contractelor R05')
p(d,'Auditor: rol AI juridic independent. Raportul comercial al celui de-al doilea auditor este separat. Prezenta analiză fixează versiunile R05 existente; recomandările sunt numai sesizări și nu au fost introduse în contracte după instrucțiunea utilizatorului de înghețare.')
p(d,data['concluzie'])
h(d,'1. Sinteză pentru decizie')
p(d,'Prețurile de bază sunt 45.000 EUR net pentru BERGER și 69.300 EUR net pentru PECH. Cele 18 luni de execuție, 12 vizite incluse și opțiunea de până la 6 suplimentare sunt condiții ale proiectelor R05. Extinderile față de sursele inițiale și probele încă lipsă nu reprezintă acceptarea prestatorilor.')
sty.table(d,['Element','BERGER','PECH'],[
['Preț de bază net / TVA / brut','45.000 / 9.000 / 54.000 EUR','69.300 / 13.860 / 83.160 EUR'],
['Bauphase în R05','18 luni de la B0 efectiv; T0 separat','18 luni de la B0 efectiv; T0 separat'],
['Vizite în oferta originală','Maximum 8 de asistență; PI distinct','12 Abnahmen'],
['Vizite în R05','12 incluse + până la 6 opționale','12 incluse + până la 6 opționale'],
['Preț suplimentar în R05','2.000 / 12 EUR net per vizită; ~166,67 afișat. Tarif nou propus de AG.','450 EUR net / 540 EUR cu TVA per vizită, tarif oferit.'],
['Toate 6 opțiunile','Maximum 1.000 EUR net / 1.200 brut pe total','2.700 EUR net / 3.240 brut pe total'],
['Bază + toate cele 6 opțiuni','46.000 EUR net / 55.200 brut pe total','72.000 EUR net / 86.400 brut pe total']],[4.1,6.75,6.75])
p(d,'TVA de 20% este ilustrată pe total; regulile fiscale și rotunjirile fiecărei facturi rămân aplicabile. Totalul opțiunilor presupune comanda distinctă și executarea tuturor celor șase. Nu este plafon al tuturor celorlalte schimbări sau creanțe legale posibile. Nicio opțiune neutilizată nu este facturată automat.')
h(d,'2. Sesizări juridice deschise și starea constatărilor')
for item in data['constatari']:
    d.add_paragraph(item['id']+' — '+item.get('provider','')+' — '+item.get('severitate',''),'Heading 2')
    for key,label in [('constatare','Constatare'),('efect','Efect'),('recomandare','Recomandare neimplementată'),('R05','Stare în textul înghețat')]:
        if item.get(key):sty.para(d,item[key],label+':')
h(d,'3. Ofertă și cerere AG → contract: fiecare condiție identificată')
p(d,'Această direcție urmărește atât ofertele primite, cât și cererea tehnică transmisă de AG, la care PECH face trimitere. Cererea AG este tratată ca sursă transmisă, fără a presupune acceptarea fără rezerve a tuturor termenilor prin referința generică. Calificarea contractuală este distinctă de verificarea sumei. S0 arată că nu a fost identificat un defect material în acel punct, menținând necesitatea acordului și a probelor.')
for provider in ['BERGER','PECH','AG']:
    rows=[x for x in data['oferte_spre_contract'] if x.get('prestator')==provider]
    if not rows:continue
    d.add_paragraph(provider+' — '+str(len(rows))+' condiții','Heading 2')
    tab=[]
    for x in rows:
        details=x.get('berger','') if provider=='BERGER' else x.get('pech','') if provider=='PECH' else ('BERGER: '+x.get('berger','')+'\nPECH: '+x.get('pech',''))
        tab.append([x['id']+' — '+x['tema']+'\n'+x.get('sursa',''),x.get('clauza','')+'\n'+details,severity(x.get('severitate',''))+'\n'+x.get('verdict','')])
    sty.table(d,['Aspect și sursă','Contract / diferență constatată','Calificare'],tab,[4.8,8.5,4.29])
others=[x for x in data['oferte_spre_contract'] if x.get('prestator') not in ['BERGER','PECH','AG']]
if others:
    d.add_paragraph('Cerințe AG și condiții comune — '+str(len(others))+' condiții','Heading 2')
    sty.table(d,['Aspect și sursă','Contract / diferență constatată','Calificare'],[[x['id']+' — '+x['tema']+'\n'+x.get('sursa',''),x.get('clauza','')+'\nBERGER: '+x.get('berger','')+'\nPECH: '+x.get('pech',''),severity(x.get('severitate',''))+'\n'+x.get('verdict','')] for x in others],[4.8,8.5,4.29])
h(d,'4. Contract → ofertă și lege: fiecare alineat și anexă')
p(d,'Cele 97 rânduri acoperă individual cele 91 alineate ale celor 16 capitole și toate cele 6 anexe. Textul german integral și datele de trasabilitate sunt păstrate în fișierul JSON al checklistului; în tabel apar referința exactă și evaluarea pentru fiecare prestator. Normele sunt identificate în prima coloană, cu legăturile primare în secțiunea de surse. Recomandările generice de confirmare și regula fără modificări se aplică tuturor rândurilor S0.')
for section in range(1,17):
    rows=[x for x in data['checklist'] if x['clauza'].startswith('§ '+str(section)+' ')]
    if not rows:continue
    d.add_paragraph('§ '+str(section)+' — '+rows[0]['tema'],'Heading 2')
    tab=[]
    for x in rows:
        obs=severity(x['severitate'])+'\n'+x['verdict']+'\nSursă: '+no_url(x['sursa'])
        if not str(x.get('severitate','')).startswith('S0'):obs+='\nRecomandare: '+x.get('corectie','')
        tab.append([x['id']+'\n'+x['clauza']+'\n'+x.get('lege',''),x.get('berger',''),x.get('pech',''),obs])
    sty.table(d,['Clauză / temei','BERGER','PECH','Calificare / sursă'],tab,[3.5,4.6,4.6,4.89])
annexrows=[x for x in data['checklist'] if not x['clauza'].startswith('§ ')]
d.add_paragraph('Anexele 1–6','Heading 2')
sty.table(d,['Anexă / temei','BERGER','PECH','Calificare / sursă'],[[x['id']+'\n'+x['clauza']+'\n'+x.get('lege',''),x.get('berger',''),x.get('pech',''),severity(x['severitate'])+'\n'+x['verdict']+'\nSursă: '+no_url(x['sursa'])+('\nRecomandare: '+x.get('corectie','') if not str(x.get('severitate','')).startswith('S0') else '')] for x in annexrows],[3.5,4.6,4.6,4.89])
h(d,'5. Surse primare și jurisprudență verificată')
p(d,'Sursele au fost recitite de auditor la 2026.10.13. Textele normative, codul profesional/standardele etice și hotărârile au funcții distincte. Cele șase hotărâri sunt comparatori aplicați, nu hotărâri asupra acestor contracte. Cercetarea este selectivă și nu certifică un rezultat judiciar.')
for x in data['surse']['norme']:sty.source(d,x['id']+' ('+x.get('data_verificarii','2026.10.13')+')',x['url'])
case_lines=[line for line in (WORK/'2026.10.13 Audit juridic bidirectional R05.txt').read_text(encoding='utf-8-sig').splitlines() if line.startswith('OGH ')]
assert len(case_lines)==6
for x,line in zip(data['surse']['jurisprudenta'],case_lines):
    assert x['numar'].replace(' ','') in line.replace(' ','')
    p(d,line)
    sty.source(d,'Sursa primară OGH '+x['numar']+' — '+x['data'],x['url'])
h(d,'6. Documente originale și versiunile înghețate')
for key,x in data['originale'].items():
    d.add_paragraph(key,'Heading 2');p(d,str(x.get('path','')));p(d,'SHA-256: '+x.get('sha256',''))
    if x.get('status'):p(d,x['status'])
for x in data['documente']:
    d.add_paragraph(x['provider']+' R05 — '+str(x['pages'])+' pagini PDF','Heading 2')
    for ext,path in x['files'].items():p(d,ext.upper()+': '+path+'\nSHA-256: '+x['sha256'][ext])
    p(d,'Integritate: '+str(x['paragraph_count'])+' paragrafe, '+str(x['tables_count'])+' tabele și '+str(x['cell_count'])+' celule verificate în Word/PDF; fără lipsuri de text identificate.')
h(d,'7. Limitele concrete ale verificării')
for line in data['limite']:p(d,line)
p(d,'Nota finală este neatribuită în această verificare. Sunt raportate punctele deschise, cu prioritatea lor, fără a transforma verificarea într-o acceptare comercială, probă a asigurării sau confirmare a datelor profesionale.')
target=WORK/'2026.10.13 Audit juridic bidirectional R05 final.docx'
sty.save(d,target)
alltext=''.join(t.text or '' for el in Document(target).element.body for t in el.xpath('.//w:t'))
for row in data['checklist']+data['oferte_spre_contract']:
    assert row['id'] in alltext,row['id']
    for key in ['berger','pech']:
        if row.get(key) and not row[key].startswith('Nu se aplică'):assert row[key] in alltext,(row['id'],key)
for x in frozen['surse']:assert sha(ROOT/x['cale'])==x['sha256']
print(json.dumps({'output':str(target),'sha256':sha(target),'checklist_legal':97,'oferta_cerere_spre_contract':80},ensure_ascii=True))
