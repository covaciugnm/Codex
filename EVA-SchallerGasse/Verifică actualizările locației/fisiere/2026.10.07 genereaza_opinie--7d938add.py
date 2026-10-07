from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'08. Corespondenta'/'2026.10.07 Analiza juridica Donau auditata'
SRC=OUT/'2026.10.07 Analiza redactor.txt'
DEST=OUT/'2026.10.07 Opinie juridica Donau Schallergasse 35.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.73); sec.bottom_margin=Inches(.73)
sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[style_name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(11)
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.line_spacing=1.08
    fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
    if fonts is not None:
        for attr in list(fonts.attrib):
            if attr.endswith('Theme'):del fonts.attrib[attr]
    for unwanted in ['w:spacing','w:kern']:
        for el in list(st.element.get_or_add_rPr().findall(qn(unwanted))):st.element.get_or_add_rPr().remove(el)
for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(14); doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12); doc.styles['Heading 2'].font.bold=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
for sn in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[sn].paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.widow_control=True
header=sec.header.paragraphs[0]
header.text='A&C Wohnart Immobilien GmbH   |   DONAU 2044001194'
header.runs[0].font.size=Pt(9)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('2026.10.07  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
for r in foot.runs:r.font.size=Pt(9)

def add_text(p, text):
    urls=re.findall(r'https?://\S+',text)
    if urls:
        for url in urls:
            before,_,text=text.partition(url)
            add_text(p,before)
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
            run=OxmlElement('w:r');rpr=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'17365D');rpr.append(color)
            run.append(rpr);t=OxmlElement('w:t');t.text='RIS sursă oficială';run.append(t);link.append(run);p._p.append(link)
        if text:add_text(p,text)
        return
    # Preserve citations, and support light Markdown emphasis without exposing syntax.
    for chunk in re.split(r'(\*\*.*?\*\*)',text):
        if chunk.startswith('**') and chunk.endswith('**'):p.add_run(chunk[2:-2]).bold=True
        else:p.add_run(chunk)

def table(lines):
    rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.match(r'^:?-+:?$',v.replace(' ','')) for v in r)]
    if not rows:return
    columns=max(map(len,rows)); t=doc.add_table(rows=0,cols=columns)
    t.autofit=False
    widths={2:[2.0,4.9],3:[1.45,2.4,3.05],4:[1.4,1.9,1.8,1.8]}.get(columns,[6.9/columns]*columns)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');margins.append(el)
    props.append(margins)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(c,txt) in enumerate(zip(cells,row)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.02
            add_text(p,txt)
            for run in p.runs:run.font.size=Pt(9.5);run.bold=(i==0)
            if i==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EDF2');c._tc.get_or_add_tcPr().append(fill)
        if i==0:
            repeat=OxmlElement('w:tblHeader');t.rows[i]._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');t.rows[i]._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=SRC.read_text(encoding='utf-8-sig').splitlines()
start=next((i for i,line in enumerate(lines) if line.startswith('1. ')),0)
lines=lines[start:]
doc.add_paragraph('Opinie juridică privind asigurarea clădirii Schallergasse 35',style='Title')
doc.add_paragraph('Rezilierea poliței DONAU și drepturile A&C Wohnart Immobilien GmbH',style='Subtitle')
doc.add_paragraph('2026.10.07   |   Polița 2044001194   |   Analiză pentru decizie')
doc.add_paragraph('Redactare și audit realizate de doi agenți AI separați, pe documentele dosarului și surse juridice austriece primare. Documentul nu este semnat și nu este atribuit unui avocat. Concluziile condiționate de probe lipsă sunt indicate în text.')
i=0
appendix=False
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
        table(block);continue
    # Skip writer title metadata if duplicated by document front matter.
    if i<8 and (line.startswith('2026.10.07') or line.upper().startswith(('OPINIE JURIDIC','PUNCT DE VEDERE JURIDIC'))):continue
    if line.startswith('ANEXĂ'):
        appendix=True
        p=doc.add_paragraph('Anexă cu registrul probelor și surselor',style='Heading 1')
        continue
    if line.startswith('### ') or re.match(r'^\d+\.\d+\s',line):
        txt=re.sub(r'[^\w\s]',' ',line.lstrip('# '));txt=re.sub(r'\s+',' ',txt).strip()
        p=doc.add_paragraph(txt,style='Heading 2')
    elif line.startswith('## ') or re.match(r'^\d+[.)]\s+[^.]{4,120}$',line) or line in ['SINTEZĂ PENTRU DECIZIE','SURSE','SURSE PRIMARE','CONCLUZIE']:
        txt=re.sub(r'[^\w\s]',' ',line.lstrip('# '));txt=re.sub(r'\s+',' ',txt).strip()
        p=doc.add_paragraph(txt,style='Heading 1')
    else:
        p=doc.add_paragraph();add_text(p,line)
        if appendix:
            for r in p.runs:r.font.size=Pt(9.5)
            p.paragraph_format.space_after=Pt(5)
    if line.startswith('https://'):
        for r in p.runs:r.font.size=Pt(8.5)
        p.paragraph_format.space_after=Pt(3)

doc.core_properties.title='Opinie juridică privind asigurarea DONAU pentru Schallergasse 35'
doc.core_properties.subject='Reziliere, prime restante, acoperire, clauza 1000K și regres contractual'
doc.core_properties.author='Analiză documentară asistată de AI'
for el in list(doc.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.save(DEST)
print(json.dumps({'docx':str(DEST),'words':len(SRC.read_text(encoding='utf-8-sig').split()),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
