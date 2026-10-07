from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
import json,hashlib
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from docx import Document
R=Path(__file__).resolve().parent.parent
D='2026.10.05';P=R/'08. Corespondenta'/f'{D} Actualizare comunicari';B=R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail'
evidence=json.loads((P/f'{D} Dovezi Eva-Mail.json').read_text(encoding='utf-8'));emails=evidence['emails']
states=json.loads((P/f'{D} Statusuri verificate.json').read_text(encoding='utf-8'))
reg=json.loads((B/'2026.10.02 Registru comunicatii.json').read_text(encoding='utf-8'))
attachpath=B/'2026.10.02 Registru atasamente.json'
if not attachpath.exists():attachpath=B/'2026.10.01 Registru atasamente.json'
atts=json.loads(attachpath.read_text(encoding='utf-8'))
oldids={e['id'] for e in reg['comunicatii']};oldatts={a['id'] for a in atts}
newrows=[];newatts=[];audit=[]
index={e['cale']:e for e in json.loads((R/'folder map/inventar.json').read_text(encoding='utf-8'))}
def rel(p):return p.relative_to(R).as_posix()
def txt(p,t):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(t,encoding='utf-8-sig')
def js(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def check(p):
 e=index.get(rel(p))
 if e:
  audit.append({'cale':rel(p),'hash_identic_index':hashlib.sha256(p.read_bytes()).hexdigest()==e['sha256'],'dimensiune_identic_index':p.stat().st_size==e['octeti']})
def partner(e):
 if 'schauer' in e['from_address']:return 'SCHAUERLEUTE'
 if 'schmitt' in e['from_address']:return 'Schmitt + Sohn'
 if 'bau-werte' in e['from_address']:return 'BAU-WERTE'
 if 'cerha' in e['from_address']:return 'CERHA HEMPEL'
 return 'TOMS'
check(B/'2026.10.02 Registru comunicatii.json')
for e in emails:
 n=partner(e);stamp=datetime.fromisoformat(e['received_at'].replace('Z','+00:00'));local=stamp.astimezone(ZoneInfo('Europe/Bucharest'))
 paths=[];m=EmailMessage(policy=SMTP)
 for k,v in [('From',e['from_address']),('To',', '.join(e['to'])),('Subject',e['subject'].replace('\r','').replace('\n',' ')),('X-Eva-Mail-ID',e['id']),('X-Eva-Status','SENT' if e['is_sent'] else 'RECEIVED'),('X-Archive-Export','EVA API text; reconstructed EML')]:m[k]=v
 if e['cc']:m['Cc']=', '.join(e['cc'])
 m['Date']=format_datetime(stamp);m.set_content(e['text'])
 for a in e['attachments']:
  ap=P/f'{D} Atasamente'/(D+' '+a['id'][:8]+' '+a['name']);raw=ap.read_bytes();assert len(raw)==a['size'];sha=hashlib.sha256(raw).hexdigest();paths.append(rel(ap))
  ct=(a.get('content_type') or 'application/octet-stream').split('/',1);m.add_attachment(raw,maintype=ct[0],subtype=ct[1],filename=a['name'])
  if a['id'] not in oldatts:
   at={'id':a['id'],'name':a['name'],'size':len(raw),'email_id':e['id'],'date':e['received_at'],'cale':rel(ap),'sha256':sha,'stare':'salvat original integral','data':local.strftime('%Y.%m.%d'),'partener':n,'nume_original':a['name'],'octeti':len(raw)}
   atts.append(at);newatts.append(at);oldatts.add(a['id'])
 stem=D+' '+('TRIMIS' if e['is_sent'] else 'PRIMIT')+' '+n+' '+e['id'][:8]
 ep=lambda ext:P/(stem+ext)
 head=f"{D} | Export verificare | {'TRIMIS' if e['is_sent'] else 'PRIMIT'}\nData mesaj: {local.strftime('%Y.%m.%d %H:%M:%S')} Romania\nData tehnica: {e['received_at']}\nDe la: {e['from_name']} <{e['from_address']}>\nCatre: {', '.join(e['to'])}\nCC: {', '.join(e['cc']) or '—'}\nSubiect: {e['subject']}\nID Eva-Mail: {e['id']}\nSursa: Eva-Mail API; EML reconstruit cu atasamente originale, nu export MIME integral.\n"
 txt(ep('.txt'),head+'\nAtasamente:\n'+'\n'.join(paths)+'\n\n--- TEXT ORIGINAL EVA NEMODIFICAT ---\n\n'+e['text']);js(ep('.json'),e);ep('.eml').write_bytes(m.as_bytes())
 if e['id'] not in oldids:
  row={'data':local.strftime('%Y.%m.%d'),'ora_utc':e['received_at'][11:16],'partener':n,'domeniu':next(s['domeniu'] for s in reg['statusuri'] if s['partener']==n),'sens':'TRIMIS' if e['is_sent'] else 'PRIMIT','subiect':e['subject'].replace('\r','').replace('\n',' '),'expeditor':e['from_address'],'destinatari':'; '.join(e['to']),'cc':'; '.join(e['cc']),'id':e['id'],'cale':rel(ep('.txt')),'atasamente':len(paths),'trunchiat':e['truncated']}
  reg['comunicatii'].append(row);newrows.append(row);oldids.add(e['id'])
js(B/'Surse'/f'{D} Completare verificare comunicari.json',emails)
for n,info in states.items():
 st=next(s for s in reg['statusuri'] if s['partener']==n)
 st.update(status=info['status'],urmatorul_pas=info['urmatorul_pas'],sursa_status=info['surse']+[rel(P/f'{D} Dovezi Eva-Mail.json')],nivel='Sinteza verificata')
 if info.get('ultimul_primit'):st['ultimul_primit']=info['ultimul_primit']
 if info.get('subiect'):st['subiect_ultimul_primit']=info['subiect']
 st['ultima_comunicare']=info.get('ultima_comunicare',st['ultima_comunicare'])
 st['mesaje']+=sum(row['partener']==n for row in newrows)
 prior=sorted((B/'Parteneri').glob('* Log discutii - '+n+'.txt'))[-1];check(prior)
 note=f"{D} | Verificare comunicari\nStatus scurt: {st['status']}\nUltimul raspuns: {st['ultimul_primit']} | {st['subiect_ultimul_primit']}\nUrmatorul pas: {st['urmatorul_pas']}\nSurse: {'; '.join(st['sursa_status'])}\n\n"
 old=prior.read_text(encoding='utf-8-sig');target=B/'Parteneri'/f'{D} Log discutii - {n}.txt'
 txt(target,note+'ISTORIC — STATUSURILE VECHI POT FI DEPASITE\n\n'+old)
 doc=Document();doc.add_paragraph(note);doc.add_paragraph('ISTORIC');doc.add_paragraph(old);doc.save(target.with_suffix('.docx'))
cover=reg['acoperire'];cover['mesaje_salvate']=len(reg['comunicatii']);cover['mesaje_suplimentare_fata_de_lista']=cover['mesaje_salvate']-cover['mesaje_candidate']
cover['atasamente_salvate']+=len(newatts);cover['atasamente_identificate']+=len(newatts)
docs=sum(a['name'].lower().endswith(('.pdf','.docx')) for a in newatts)
for k in ['atasamente_documentare','atasamente_documentare_salvate']:cover[k]+=docs
cover['verificare_incrementala']={'data':D,'mesaje_noi':len(newrows),'originale_noi':len(newatts),'originale_exportate':32,'perioada':'2026.10.01–2026.10.05','office_sincronizat':'2026.10.05 11:31:09 Romania','nota':'Verificate mesaje office si cautari globale Schallergasse. Cosmin@ig.ro sincronizat doar pana la 2026.08.12. Nu este identificat raspuns relevant nou din 2026.10.03–05 in cautarile disponibile.'}
reg['data_actualizarii']=D
reg.setdefault('actualizari_punctuale',[]).append({'data':D,'subiect':'Verificare generala si automatizare zilnica 08:00 Europe/Bucharest','sursa':rel(P/f'{D} Dovezi Eva-Mail.json')})
reg['monitorizare']={'id':'verificare-zilnic-schallergasse-35','status':'ACTIVE','ora':'08:00','fus_orar':'Europe/Bucharest','afisare':'in acest chat','autorizare':'cerere explicita utilizator'}
js(B/f'{D} Registru comunicatii.json',reg);js(B/f'{D} Registru atasamente.json',atts)
summary=f"{D} | ACTUALIZARE SCHALLERGASSE 35\n\n"
for n,info in states.items():
 st=next(s for s in reg['statusuri'] if s['partener']==n)
 summary+=f"{n}\nStatus scurt: {st['status']}\nUltimul raspuns: {st['ultimul_primit']}\nUrmatorul pas: {st['urmatorul_pas']}\nSurse: {'; '.join(st['sursa_status'])}\n\n"
summary+="Verificare zilnica ACTIVE la 08:00 Europe/Bucharest, in acest chat. ID: verificare-zilnic-schallergasse-35. Sincronizare office 2026.10.05 11:31:09 Romania; cosmin@ig.ro pana la 2026.08.12. Nu sunt identificate noi raspunsuri relevante din 2026.10.03–05. Nicio comunicare externa trimisa in aceasta verificare. Incetarea DONAU ramane neconfirmata; raspuns cerut 2026.10.07, termen Commerz 2026.10.12 distinct. Jurnalele anterioare sunt pastrate.\n\n"
txt(P/f'{D} Rezumat comunicari noi.txt',summary)
txt(R/f'{D} Status proiect.txt',summary+'ISTORIC\n\n'+(R/'2026.10.02 Status proiect.txt').read_text(encoding='utf-8-sig'))
check(R/'2026.10.02 Log progres proiect.xlsx');w=load_workbook(R/'2026.10.02 Log progres proiect.xlsx')
ws=w['Status pe scurt'];cols=[c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
 if row[0].value in states:
  st=next(s for s in reg['statusuri'] if s['partener']==row[0].value)
  for i,k in enumerate(cols,1):
   v=st.get(k);ws.cell(row[0].row,i,'; '.join(v) if isinstance(v,list) else v)
for sheet,added in [('Comunicatii',newrows),('Atasamente',newatts)]:
 ws=w[sheet];cols=[c.value for c in ws[1]]
 for a in added:ws.append([a.get(k) for k in cols])
ws=w['Acoperire']
for row in ws.iter_rows(min_row=2):
 if row[0].value in cover:
  v=cover[row[0].value];row[1].value=str(v) if isinstance(v,dict) else v
sh=w.create_sheet('Verificare 2026.10.05');sh.append(['Partener','Status','Ultimul raspuns','Urmatorul pas','Surse'])
for n in states:
 st=next(s for s in reg['statusuri'] if s['partener']==n);sh.append([n,st['status'],st['ultimul_primit'],st['urmatorul_pas'],'; '.join(st['sursa_status'])])
for c in ['A','B','C','D','E']:sh.column_dimensions[c].width=25 if c in ['A','C'] else 90
for row in sh:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
sh.freeze_panes='A2';w.save(R/f'{D} Log progres proiect.xlsx')
doc=Document(R/'2026.10.02 Log progres proiect.docx');doc.paragraphs[0].insert_paragraph_before(summary);doc.save(R/f'{D} Log progres proiect.docx')
offers=sorted((R/'04. Firme + Executie').glob('* Ultimele raspunsuri ofertanti.xlsx'))[-1];check(offers);ow=load_workbook(offers)
for ws in ow:
 for row in ws.iter_rows(min_row=2):
  if row[1].value in states:
   n=row[1].value;row[0].value=D;row[2].value=states[n]['status'];row[3].value=states[n]['urmatorul_pas'];row[4].value='; '.join(states[n]['surse'])
ow.save(R/'04. Firme + Executie'/f'{D} Ultimele raspunsuri ofertanti.xlsx')
txt(R/'04. Firme + Executie/01. Verificator (Pruefingenieur)/TOMS/2026.10.02 Contract v4.1 - data actualizata'/f'{D} Status trimitere verificat.txt',states['TOMS']['status']+'\nSursa: 64c91100-89c3-4206-8e7e-988e575385b8\n')
read=R/'folder map/README.md'
content=read.read_text(encoding='utf-8')
head=f"\n\n## {D} — Status curent si verificare zilnica\n\nPuncte curente: ../{D} Status proiect.txt; ../{D} Log progres proiect.xlsx; ../08. Corespondenta/2026.09.30 Arhiva Eva-Mail/{D} Registru comunicatii.json si {D} Registru atasamente.json. Dovezi si originale: ../08. Corespondenta/{D} Actualizare comunicari/.\n\nSCHAUERLEUTE propune 2026.10.09 09:30 Viena, de confirmat; BAU-WERTE transmite oferta Bauwerksbuch; Schmitt + Sohn oferta lift; CERHA transmite scrisoare Sturm. TOMS v4.1 este TRIMIS la 2026.10.02, neacceptat/nesemnat. Paragrafele vechi privind DRAFT sunt istoricul. Jurnalele curente ale partenerilor au prefix {D}. Oferte: ../04. Firme + Executie/{D} Ultimele raspunsuri ofertanti.xlsx.\n\nLa cererea expresa a utilizatorului, verificare zilnica ACTIVE la 08:00 Europe/Bucharest in acest chat, ID verificare-zilnic-schallergasse-35. Sincronizare office 2026.10.05 11:31:09 Romania; cosmin@ig.ro ramane la 2026.08.12. Nu sunt identificate comunicari relevante noi din 2026.10.03–05. Exportul ramane API text, nu MIME integral.\n"
read.write_text(content.replace('# Harta arhivei Schallergasse 35','# Harta arhivei Schallergasse 35'+head,1),encoding='utf-8')
js(P/f'{D} Audit actualizare.json',{'data':D,'mesaje_noi':len(newrows),'originale_noi':len(newatts),'documentare_noi':docs,'audit_index':audit,'monitorizare':reg['monitorizare']})
assert len(newrows)==5
print('Actualizat:',len(reg['comunicatii']),'mesaje;',len(newatts),'originale noi; 6 parteneri; verificare zilnica activa.')

