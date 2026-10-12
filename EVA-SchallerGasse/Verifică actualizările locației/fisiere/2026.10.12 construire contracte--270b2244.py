from pathlib import Path
import json, re, sys
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DEST=Path(__file__).resolve().parent

def text(x):
    if isinstance(x,(int,float)):
        return f'{x:,.2f}'.replace(',','~').replace('.',',').replace('~','.')
    s=str(x or '')
    s=re.sub(r'\bEUR\s*([0-9]{1,3}(?:,[0-9]{3})+\.[0-9]{2})',lambda m:'EUR '+f'{float(m.group(1).replace(",","")):,.2f}'.replace(',','~').replace('.',',').replace('~','.'),s)
    # Typographic cleanup only. Preserve URLs, email addresses and document IDs.
    saved=[]
    def hold(m):
        saved.append(m.group());return f'ZZPROTECTED{len(saved)-1}ZZ'
    s=re.sub(r'https?://\S+|[\w.+%-]+@[\w.-]+\.[A-Za-z]{2,}|[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}',hold,s,flags=re.I)
    s=s.replace('\u00a0',' ').replace('—','-').replace('–','-')
    s=re.sub(r'([,;:])(?=[A-Za-zÄÖÜäöüĂÂÎȘȚăâîșț])',r'\1 ',s)
    s=s.replace('Anschlüsse/Anker/Schweißnähte','Anschlüsse / Anker / Schweißnähte')
    s=s.replace('Balkone/Stiegen/Resttragwerk','Balkone / Stiegen / Resttragwerk')
    s=re.sub(r'(§{1,2})(?=\d)',r'\1 ',s)
    s=re.sub(r'\bEUR(?=\d)', 'EUR ',s)
    s=re.sub(r'(?<=\d)(?=(?:EUR|netto|brutto|net|brut|vizite|visite|Wochen|Monate|Tage|luni|pagini|Abnahmen|Betreuungstermine|TVA)\b)',' ',s)
    s=s.replace('EURnet','EUR net').replace('EURbrut','EUR brut').replace('EURnetto','EUR netto').replace('EURbrutto','EUR brutto')
    s=re.sub(r'\b(T0)\+(\d+)',r'\1 + \2',s)
    s=s.replace('Bestand+Neu','Bestand und Neu').replace('existent+nou','existent și nou')
    for i,v in enumerate(saved):s=s.replace(f'ZZPROTECTED{i}ZZ',v)
    return s

def heading(s):
    s=text(s).replace('§','')
    s=re.sub(r'[-—–:;,()/]',' ',s)
    return re.sub(r'\s+',' ',s).strip()

def borders(cell):
    pr=cell._tc.get_or_add_tcPr()
    node=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');node.append(e)
    pr.append(node)
    margins=OxmlElement('w:tcMar')
    for edge,value in [('top','90'),('left','100'),('bottom','90'),('right','100')]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)

def table(doc, headers, rows, money=False):
    headers=['AG Anteil am Paket' if 'Zahlungsanteil' in str(h) else h for h in headers]
    n=len(headers);tb=doc.add_table(rows=1,cols=n)
    tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    presets={2:[6,11.2],3:[3.2,8.4,5.6],4:[3.3,3.4,7.0,3.5],5:[5.5,2.0,2.0,2.2,5.5]}
    widths=presets.get(n,[17.2/n]*n)
    has_currency=any('EUR' in str(h) for h in headers)
    if n==3 and has_currency:widths=[2.4,11.6,3.2]
    elif n==3 and str(headers[0]).startswith('Nr'):widths=[1.2,10.4,5.6]
    elif n==3 and str(headers[0]).startswith('Bereich'):widths=[4.8,6.8,5.6]
    elif n==3 and str(headers[0]).startswith('Nachweis'):widths=[8.0,5.2,4.0]
    elif n==3 and str(headers[0]).startswith('Punkt'):widths=[5.3,7.9,4.0]
    elif n==3 and str(headers[0]).startswith('Meilenstein'):widths=[4.8,8.4,4.0]
    elif n==3 and str(headers[0]).startswith('Thema'):widths=[3.8,6.3,7.1]
    if n==4 and str(headers[0]).startswith('Meilenstein'):widths=[3.4,4.4,6.6,2.8]
    for c,w in zip(tb.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):tb.rows[0].cells[i].text=text(h)
    for row in rows:
        cs=tb.add_row().cells
        for i in range(n):
            val=text(row[i] if i<len(row) else '')
            cs[i].text='______________' if re.fullmatch(r'_{15,}',val) else val
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;borders(cell)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EAF0' if ri==0 else ('F6F7F9' if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.03
                p.paragraph_format.keep_with_next=ri==0
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT
                if ri>0 and money and 0<ci<4:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                if ri>0 and has_currency and ('EUR' in str(headers[ci])):p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:run.font.name='Arial';run.font.size=Pt(10);run.font.color.rgb=RGBColor(0,0,0);run.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return tb

def add_field(p,name):
    r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    r=p.add_run();node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=' '+name+' ';r._r.append(node)
    r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)

def build(provider):
    src=DEST/f'2026.10.12 Contract {provider} continut.json'
    data=json.loads(src.read_text(encoding='utf-8'))
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    sec.header_distance=Cm(0.7);sec.footer_distance=Cm(0.8)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.1
    normal.paragraph_format.widow_control=True
    for nm,sz in [('Title',18),('Heading 1',13),('Heading 2',12),('Subtitle',11)]:
        st=doc.styles[nm];st.font.name='Arial';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0)
        st.font.bold=nm!='Subtitle';st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(12 if nm.startswith('Heading') else 0)
        st.paragraph_format.keep_with_next=True
        if st._element.pPr is not None:
            for child in list(st._element.pPr):
                if child.tag==qn('w:pBdr'):st._element.pPr.remove(child)
    doc.core_properties.author='A&C Wohnart Immobilien GmbH'
    doc.core_properties.title=f'Vertragsvorschlag {provider} Schallergasse 35'
    doc.core_properties.subject='Tragwerksplanung Prüfingenieur und Bauwerksbuch'
    doc.core_properties.keywords='AG Entwurf 2026.10.12 Nicht unterzeichnet Nicht versendet'
    doc.core_properties.comments=''
    settings=doc.settings._element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'de-AT');normal._element.get_or_add_rPr().append(lang)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run(f'{provider}  |  AG Entwurf 2026.10.12  |  Seite ');add_field(f,'PAGE');f.add_run(' von ');add_field(f,'NUMPAGES')
    for r in f.runs:r.font.name='Arial';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph(f'Propunere de contract {provider}',style='Title')
    doc.add_paragraph('Schallergasse 35 1120 Wien',style='Subtitle')
    doc.add_paragraph('2026.10.12  |  Sinteză în română pentru lectură')
    ros=[a for a in data.get('annexes',[]) if 'RO' in a.get('title','') or 'Sinteză' in a.get('title','')]
    paras=ros[0].get('paragraphs',[]) if ros else data.get('summary_ro',[])
    doc.add_paragraph('Servicii prețuri termene și modificări propuse',style='Heading 1')
    for s in paras:doc.add_paragraph(text(s))
    doc.add_paragraph('Textul german care urmează este propunerea contractuală. Sinteza în română explică modificările și nu are efect contractual autonom.')
    german_title=doc.add_paragraph(f'Werkvertrag {provider}',style='Title')
    german_title.paragraph_format.page_break_before=True
    doc.add_paragraph('Schallergasse 35 1120 Wien',style='Subtitle')
    for i,s in enumerate(data['sections']):
        if i>0:doc.add_paragraph(heading(s['title']),style='Heading 1')
        for p in s.get('paragraphs',[]):
            shown=text(p)
            if 'Unterschrift' in p and '_' in p:
                if p.startswith('Für die AG') or p.startswith('Für den AN'):
                    shown=shown.replace('Name/Funktion ', '\nName/Funktion ').replace('Unterschrift ', '\nUnterschrift ')
                else:
                    shown=shown.replace(' Für die AG:', '\nFür die AG:').replace(' Für den AN:', '\nFür den AN:').replace(' Tatsächlicher T0', '\nTatsächlicher T0')
            out=doc.add_paragraph(shown)
            if 'Unterschrift' in p and '_' in p:out.paragraph_format.keep_together=True
        if s.get('table'):table(doc,s['table']['headers'],s['table']['rows'])
        if 'Festhonorar' in s['title'] or ('Honorar' in s['title'] and 'Abrechnung' in s['title']):
            table(doc,data.get('price_headers',['Position','Netto EUR','USt EUR','Brutto EUR','Hinweis']),data['price_rows'],money=True)
    for a in data.get('annexes',[]):
        if a in ros:continue
        if not a.get('paragraphs') and not a.get('table') and not a.get('source'):continue
        doc.add_paragraph(heading(a['title']),style='Heading 1')
        for p in a.get('paragraphs',[]):doc.add_paragraph(text(p))
        if a.get('table'):table(doc,a['table']['headers'],a['table']['rows'])
        if a.get('source') and not a.get('paragraphs'):
            doc.add_paragraph('Das unverändert beigefügte Originalangebot 2026-470/AN01 vom 2026.09.17 umfasst vier Seiten und ist als Datei 2026.10.12 Oferta originala PECH 2026.09.17.pdf beigefügt. Für abweichende Bedingungen gilt die ausdrücklich vereinbarte Rangordnung dieses Vertragsangebots.')
    # Outstanding signature and evidence fields are already in the operative
    # parties section and Anlage 4; avoid a duplicate trailing page of metadata.
    outfile=DEST/f'2026.10.12 Propunere contract {provider}.docx'
    doc.save(outfile)
    print(str(outfile))

if __name__=='__main__':
    for provider in sys.argv[1:] or ['BERGER','PECH']:build(provider)
