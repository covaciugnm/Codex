from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
import json,hashlib,shutil
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from docx import Document
import fitz
R=Path(__file__).resolve().parent.parent
D='2026.10.01'
P=R/'08. Corespondenta'/f'{D} Actualizare comunicari'
B=R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail'
data=json.loads((P/f'{D} Dovezi Eva-Mail.json').read_text(encoding='utf-8'))
emails=data['emails']
idx={x['cale']:x for x in json.loads((R/'folder map/inventar.json').read_text(encoding='utf-8'))}
audit=[]
backupdir=P/(D+' Istoric inainte de verificare '+datetime.now().strftime('%H-%M-%S'))
def rel(p):return p.relative_to(R).as_posix()
def js(p,v):
 p.parent.mkdir(parents=True,exist_ok=True)
 p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def txt(p,s):
 p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8-sig')
def check(p):
 e=idx.get(rel(p))
 if e:
  st=p.stat();sha=hashlib.sha256(p.read_bytes()).hexdigest()
  audit.append({'cale':rel(p),'metadate_neschimbate':st.st_size==e['octeti'] and st.st_mtime_ns==e.get('modificat_ns'),'hash_identic_index':sha==e['sha256'],'sha256':sha})
def backup(p):
 if p.exists():
  check(p);q=backupdir/p.relative_to(R);q.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,q)
reg=json.loads((B/'2026.09.30 Registru comunicatii.json').read_text(encoding='utf-8'))
check(B/'2026.09.30 Registru comunicatii.json')
assert not ({e['id'] for e in reg['comunicatii']} & {e['id'] for e in emails})
partners=['Donau Versicherung','Paknehad','MA6 - Taxe','MA6 - Taxe','TOMS','TOMS']
pair={e['id']:p for e,p in zip(emails,partners)}
states={
'Donau Versicherung':[
'2026.10.01: cererea de incetare a politei 2044001194 este TRIMISA la 10:20 Romania / 09:20 Viena. Catre DONAU si Loschy; CC efectiv: Maritczak si Stefanie Gruber. Capra NU apare in CC al mesajului trimis. Incetarea si data ei nu sunt confirmate. Raspuns cerut pana la 2026.10.07. Creanta anterioara 5.326,15 EUR include Q4; termen Commerz 2026.10.12, distinct si nesuspendat.',
'Asteptarea confirmarii DONAU privind primirea, data incetarii, acoperirea si decontul final; clarificarea creantei inainte de 2026.10.12. Implicarea Capra necesita o comunicare distincta, neefectuata aici.',
['23eaa675-18ad-45b1-b535-25f7583087b1','1fa19552-d1fd-459e-92b3-57ed5319d18c']],
'Maritczak - Commerz Inkasso':[
'2026.10.01: Maritczak si Stefanie Gruber sunt in CC al cererii de incetare DONAU trimise. Nu este identificat raspuns nou dupa Commerz 2026.09.30. Suma solicitata 5.326,15 EUR include Q4; termen 2026.10.12. Cererea de incetare nu suspenda termenul si nu confirma soldul acceptat.',
'Clarificarea principalului, alocarii, accesoriilor si acoperirii inainte de 2026.10.12; asteptarea confirmarii DONAU.',
['23eaa675-18ad-45b1-b535-25f7583087b1','1fa19552-d1fd-459e-92b3-57ed5319d18c']],
'Paknehad':[
'2026.09.30: raspuns PRIMIT la cererea Bauwerksbuch si documentarea starii inainte de lucrari. Cere planuri, statica, autorizatii si rapoarte prin download-link, data inceperii lucrarilor si daca BauKG a fost atribuit, inclusiv pretul atribuirii. Disponibilitate de ofertare; NU este oferta cu pret si nu confirma atribuirea BauKG.',
'Pregatirea documentelor si raspunsului privind data estimata de inceput si stadiul real BauKG. Nicio trimitere efectuata in aceasta verificare.',
['8827e952-0877-4a73-879a-2ce077005838']],
'TOMS':[
'2026.09.30: propunerea de contract v3 si cererea pentru oferta separata de detaliere sunt TRIMISE, cu 5 anexe. Total propus 37.900 EUR net: 34.500 statica/verificare + 3.400 Bauwerksbuch. Emailul separat pentru acces la portal este TRIMIS. Nu este identificat raspuns ulterior; contractul nu este acceptat sau semnat prin aceste mesaje. Oferta prelungita anterior pana la 2026.12.31.',
'Asteptarea contractului semnat/observatiilor, ofertei de detaliere, confirmarii accesului si a doua propuneri de vizita pana la 2026.10.08. Calendarul D1-D5 ramane propus si conditionat de semnarea efectiva.',
['560aebe2-82a9-4fd1-8ff7-87eb90d0f2b4','ac59f889-f611-460b-8bf6-6261e313aca6']],
'MA6 - Taxe':[
'2026.09.30: cererea de corectare a avizului Q3 este TRIMISA. MA6 a confirmat automat primirea la 15:38 Romania, cu posibile intarzieri. Nu exista raspuns de fond privind alocarea Q2 sau soldul. Q2 achitat 93,23 EUR la 2026.07.27; restul Q3 calculat 93,23 EUR ramane de confirmat.',
'Asteptarea confirmarii alocarii platii Q2 si a avizului/soldului Q3 actualizat, inclusiv referinta pentru plata restului; nu se repeta Q2.',
['a1382e07-b5eb-4f5b-a7fe-247121dc287e','a8b64f35-b608-4df7-82d9-2affe4108631']],
'CERHA HEMPEL':[
'2026.10.01: constatare de verificare, fara comunicare noua CERHA. Capra nu apare in CC al emailului efectiv trimis catre DONAU, desi era inclus in draftul arhivat. Ultimul raspuns CERHA ramane cel din jurnalul anterior.',
'Clarificarea cu Capra a dosarului DONAU si creantei pana la 2026.10.12; o eventuala informare necesita comunicare distincta, neefectuata aici.',
['23eaa675-18ad-45b1-b535-25f7583087b1']]}
att_old=json.loads((B/'2026.09.30 Registru atasamente.json').read_text(encoding='utf-8'))
att_new=[]
rows=[]
for e in emails:
 n=pair[e['id']]
 stem=D+' '+('TRIMIS' if e['is_sent'] else 'PRIMIT')+' '+n+' '+e['id'][:8]
 ep=lambda ext:P/(stem+ext)
 js(ep('.json'),e)
 local=datetime.fromisoformat(e['received_at'].replace('Z','+00:00')).astimezone(ZoneInfo('Europe/Bucharest'))
 head=f"{D} | Export verificare | {'TRIMIS' if e['is_sent'] else 'PRIMIT'}\nData mesaj: {local.strftime('%Y.%m.%d %H:%M:%S')} Romania\nData tehnica: {e['received_at']}\nDe la: {e['from_name']} <{e['from_address']}>\nCatre: {', '.join(e['to'])}\nCC: {', '.join(e['cc']) or '—'}\nSubiect: {e['subject']}\nID Eva-Mail: {e['id']}\nSursa: Eva-Mail, text disponibil API; nu este MIME original.\n"
 m=EmailMessage(policy=SMTP)
 for k,v in [('From',e['from_address']),('To',', '.join(e['to'])),('Subject',e['subject']),('X-Eva-Mail-ID',e['id']),('X-Eva-Status','SENT' if e['is_sent'] else 'RECEIVED'),('X-Archive-Export','EVA API text; reconstructed EML')]:m[k]=v
 if e['cc']:m['Cc']=', '.join(e['cc'])
 m['Date']=format_datetime(datetime.fromisoformat(e['received_at'].replace('Z','+00:00')))
 m.set_content(e['text'])
 paths=[]
 for a in e['attachments']:
  ap=P/f'{D} Atasamente'/(D+' '+a['id'][:8]+' '+a['name'])
  raw=ap.read_bytes();assert len(raw)==a['size']
  at={'id':a['id'],'name':a['name'],'size':len(raw),'email_id':e['id'],'date':e['received_at'],'cale':rel(ap),'sha256':hashlib.sha256(raw).hexdigest(),'stare':'salvat original integral','data':local.strftime('%Y.%m.%d'),'partener':n,'nume_original':a['name'],'octeti':len(raw)}
  att_new.append(at);paths.append(rel(ap))
  ct=(a.get('content_type') or 'application/octet-stream').split('/',1)
  m.add_attachment(raw,maintype=ct[0],subtype=ct[1],filename=a['name'])
 txt(ep('.txt'),head+'\nAtasamente:\n'+'\n'.join(paths)+'\n\n--- TEXT EVA NEMODIFICAT ---\n\n'+e['text'])
 ep('.eml').write_bytes(m.as_bytes())
 row={'data':local.strftime('%Y.%m.%d'),'ora_utc':e['received_at'][11:16],'partener':n,'domeniu':next(s['domeniu'] for s in reg['statusuri'] if s['partener']==n),'sens':'TRIMIS' if e['is_sent'] else 'PRIMIT','subiect':e['subject'],'expeditor':e['from_address'],'destinatari':'; '.join(e['to']),'cc':'; '.join(e['cc']),'id':e['id'],'cale':rel(ep('.txt')),'atasamente':len(paths),'trunchiat':e['truncated']}
 reg['comunicatii'].append(row);rows.append(row)
js(P/f'{D} Mesaje noi.json',emails)
js(B/'Surse'/f'{D} Completare verificare comunicari.json',emails)
for n,(status,nxt,sources) in states.items():
 s=next(s for s in reg['statusuri'] if s['partener']==n)
 s.update(status=status,urmatorul_pas=nxt,sursa_status=sources+[rel(P/f'{D} Dovezi Eva-Mail.json')],nivel='Sinteza verificata')
 relevant=[e for e in emails if pair[e['id']]==n]
 if relevant:
  s['ultima_comunicare']=max(e['received_at'] for e in relevant)[:10].replace('-','.')
  s['mesaje']+=len(relevant)
  incoming=[e for e in relevant if not e['is_sent']]
  if incoming:
   last=max(incoming,key=lambda e:e['received_at']);s['ultimul_primit']=last['received_at'][:10].replace('-','.');s['subiect_ultimul_primit']=last['subject']
 target=B/'Parteneri'/f'{D} Log discutii - {n}.txt'
 prior=target if target.exists() else B/'Parteneri'/f'2026.09.30 Log discutii - {n}.txt'
 check(prior);old=prior.read_text(encoding='utf-8-sig')
 note=f"{D} | Verificare comunicari Eva-Mail\n\nStatus scurt: {status}\nUltimul raspuns: {s['ultimul_primit']} | {s['subiect_ultimul_primit']}\nUrmatorul pas: {nxt}\nSurse: {'; '.join(s['sursa_status'])}\n\n"
 backup(target);txt(target,note+'ISTORIC — STATUSURILE DE DRAFT POT FI DEPASITE DE DOVEZILE DE MAI SUS\n\n'+old)
 dp=target.with_suffix('.docx');backup(dp)
 doc=Document();doc.add_paragraph(note);doc.add_paragraph('ISTORIC');doc.add_paragraph(old);doc.save(dp)
reg['data_actualizarii']=D
cover=reg['acoperire']
cover['mesaje_salvate']=len(reg['comunicatii'])
cover['mesaje_suplimentare_fata_de_lista']=cover['mesaje_salvate']-cover['mesaje_candidate']
for k in ['atasamente_salvate','atasamente_identificate']:cover[k]+=7
for k in ['atasamente_documentare','atasamente_documentare_salvate']:cover[k]+=6
cover['verificare_incrementala']={'data':D,'mesaje_noi':6,'originale_noi':7,'mesaje_office_perioada':13,'perioada':'2026.09.30–2026.10.01','cautari_globale':['Schallergasse','wohnart'],'limita':'Office sincronizat 2026.10.01 13:01:51 Romania; cosmin@ig.ro doar pana la 2026.08.12. Cautarile nu garanteaza gasirea mesajelor fara legatura textuala cu proiectul.'}
reg.setdefault('actualizari_punctuale',[]).append({'data':D,'subiect':'Verificare comunicari noi','mesaje':6,'sursa':rel(P/f'{D} Dovezi Eva-Mail.json')})
for dr in reg.get('drafturi',[]):
 if dr.get('id')=='141ce321-f1e8-4fdb-ae64-d26afeac4122':
  dr['dovada_trimiterii_ulterioare']=emails[0]['id'];dr['nota_verificare']='Draft istoric; mesaj efectiv trimis cu CC modificat, fara Capra.'
newreg=B/f'{D} Registru comunicatii.json';js(newreg,reg)
js(B/f'{D} Registru atasamente.json',att_old+att_new)
summary=f"{D} | ACTUALIZARE COMUNICARI — VERIFICARE EVA-MAIL\n\n"
for n in ['Paknehad','Donau Versicherung','TOMS','MA6 - Taxe','Maritczak - Commerz Inkasso']:
 s=next(s for s in reg['statusuri'] if s['partener']==n)
 summary+=f"{n}\nStatus scurt: {s['status']}\nUltimul raspuns: {s['ultimul_primit']} | {s['subiect_ultimul_primit']}\nUrmatorul pas: {s['urmatorul_pas']}\nSursa: {'; '.join(s['sursa_status'])}\n\n"
summary+="Acoperire: 13 mesaje office din 2026.09.30–2026.10.01, toate comparate cu registrul; 6 mesaje noi si 7 originale salvate. Nu este identificat un nou raspuns extern din 2026.10.01 pentru cladire. Sincronizare office: 2026.10.01 13:01:51 Romania; cosmin@ig.ro ramane la 2026.08.12. EML reconstruite din text API cu atasamente originale, nu export MIME integral. Nicio comunicare externa efectuata in aceasta verificare.\n\n"
txt(P/f'{D} Rezumat comunicari noi.txt',summary)
sp=R/f'{D} Status proiect.txt';old=sp.read_text(encoding='utf-8-sig');backup(sp);txt(sp,summary+'ISTORIC — STATUSURILE DE DRAFT DE MAI JOS SUNT DEPASITE DE DOVEZILE DE MAI SUS\n\n'+old)
wp=R/f'{D} Log progres proiect.xlsx';backup(wp);w=load_workbook(wp)
ws=w['Status pe scurt'];cols=[c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
 if row[0].value in states:
  s=next(s for s in reg['statusuri'] if s['partener']==row[0].value)
  for i,k in enumerate(cols,1):
   v=s.get(k);ws.cell(row[0].row,i,'; '.join(v) if isinstance(v,list) else v)
for sheet,added in [('Comunicatii',rows),('Atasamente',att_new)]:
 ws=w[sheet];cols=[c.value for c in ws[1]]
 for a in added:ws.append([a.get(k) for k in cols])
for row in w['Acoperire'].iter_rows(min_row=2):
 if row[0].value in cover:
  v=cover[row[0].value];row[1].value=str(v) if isinstance(v,dict) else v
sh=w.create_sheet('Verificare 2026.10.01')
sh.append(['Partener','Status','Ultimul raspuns','Urmatorul pas','Sursa'])
for n in states:
 s=next(s for s in reg['statusuri'] if s['partener']==n);sh.append([n,s['status'],s['ultimul_primit'],s['urmatorul_pas'],'; '.join(s['sursa_status'])])
for col in ['A','B','C','D','E']:sh.column_dimensions[col].width=25 if col in ['A','C'] else 95
for row in sh:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
sh.freeze_panes='A2';w.save(wp)
dp=R/f'{D} Log progres proiect.docx';backup(dp);doc=Document(dp);doc.paragraphs[0].insert_paragraph_before(summary);doc.save(dp)
offers=R/'04. Firme + Executie/2026.09.30 Ultimele raspunsuri ofertanti.xlsx';check(offers);ow=load_workbook(offers)
for ws in ow:
 for row in ws.iter_rows(min_row=2):
  if row[1].value in ['Paknehad','TOMS']:
   n=row[1].value;row[0].value=D;row[2].value=states[n][0];row[3].value=states[n][1];row[4].value='; '.join(states[n][2])
ow.save(offers.with_name(D+' Ultimele raspunsuri ofertanti.xlsx'))
for folder,n in [('08. Corespondenta/2026.10.01 Donau - cerere reziliere','Donau Versicherung'),('08. Corespondenta/2026.09.30 MA6 - Corectare aviz Q3','MA6 - Taxe'),('04. Firme + Executie/01. Verificator (Pruefingenieur)/TOMS/2026-10-01 Contract v3 - propunere','TOMS')]:
 txt(R/folder/f'{D} Status trimitere verificat.txt',states[n][0]+'\n\nUrmatorul pas: '+states[n][1]+'\nSurse: '+'; '.join(states[n][2])+'\nDosar dovezi: '+rel(P)+'\n')
read=R/'folder map/README.md';backup(read);content=read.read_text(encoding='utf-8')
new=f"\n\n## {D} — Verificare comunicari noi; status curent\n\nPuncte curente: ../{D} Status proiect.txt, ../{D} Log progres proiect.xlsx, ../08. Corespondenta/2026.09.30 Arhiva Eva-Mail/{D} Registru comunicatii.json si {D} Registru atasamente.json in acelasi folder. Dovezi: ../08. Corespondenta/{D} Actualizare comunicari/.\n\nSase mesaje suplimentare si sapte originale salvate. Paknehad a raspuns la 2026.09.30, cere documente si informatii, fara pret. Cererea DONAU este TRIMISA la 2026.10.01; CC efectiv Maritczak si Gruber, fara Capra. Propunerea TOMS si accesul portal sunt TRIMISE la 2026.09.30; acceptarea contractului nu este confirmata. Cererea MA6 este TRIMISA, cu confirmare automata de primire, fara solutionare. Paragrafele vechi de mai jos despre drafturi reprezinta istoricul. Nu este identificat raspuns extern nou din 2026.10.01 in mesajele disponibile. Office sincronizat 2026.10.01 13:01:51 Romania; cosmin@ig.ro ramane la 2026.08.12. Jurnalele curente ale partenerilor au prefix {D}. Oferte: ../04. Firme + Executie/{D} Ultimele raspunsuri ofertanti.xlsx.\n"
read.write_text(content.replace('# Harta arhivei Schallergasse 35','# Harta arhivei Schallergasse 35'+new,1),encoding='utf-8')
pdf=P/f'{D} Atasamente'/f'{D} 91b2e1fd 2026-10-01 Werkvertrag TOMS - Entwurf AG v3.pdf'
with fitz.open(pdf) as f:
 text='\n'.join(p.get_text() for p in f)
 assert all(v in text for v in ['37.900','25.000','1.500','8.000','3.400','Entwurf'])
audit.append({'original':rel(pdf),'verificare':'22 pagini; propunere v3; 37.900 EUR net = 34.500 + 3.400; trimiterea nu confirma semnarea.'})
js(P/f'{D} Audit verificare.json',{'data':D,'fisiere_verificate_fata_de_index':audit,'atasamente_originale':att_new,'mesaje_noi':6,'registru_curent':rel(newreg)})
assert len(reg['comunicatii'])==786
print('Actualizat: 786 mesaje, 7 originale noi; 6 jurnale parteneri; registru proiect si status curent; istoric pastrat.')
