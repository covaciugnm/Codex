from pathlib import Path
import json,shutil
from openpyxl import load_workbook
R=Path(__file__).resolve().parent.parent
P=R/'05. Asigurari/Asigurare cladire/2026.10.14 Situatie Donau'
B=R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail'
src=(P/'2026.10.14 Situatie actuala Donau.txt').relative_to(R).as_posix()
status='2026.10.14: verificare punctuală DONAU. Ultimul răspuns 2026.10.06 refuză încetarea și declară lipsa acoperirii pentru prime restante; contractul nu este confirmat încetat și nu există suspendare agreată a primelor. Nicio restabilire sau răspuns ulterior identificat. Cererea 2026.10.01 este TRIMISĂ; ciorna 2026.10.07 rămâne pending / NETRIMISĂ.'
nextstep='Clarificare urgentă scrisă a sumei/alocării/momentului restabilirii acoperirii și a datei de încetare. Actualizarea termenelor depășite din ciorna netrimisă înainte de o eventuală expediere. Ultima cerere Commerz: 5.326,15 EUR, termen 2026.10.12 depășit, fără prelungire identificată; nu reprezintă sold bancar reconciliat.'
note='2026.10.14 — DONAU: cronologie și stare acoperire\n\nStatus scurt: '+status+'\n\nUltimul răspuns: Cornelia Loschy, 2026.10.06, ID 2c273256-a0bf-408b-bb22-09718372e888. Ultimul email trimis: 2026.10.01, ID 23eaa675-18ad-45b1-b535-25f7583087b1; CC efectiv Maritczak/Gruber, fără Capra.\n\nUrmătorul pas: '+nextstep+'\n\nSursa: '+src+'\nTabel: '+(P/'2026.10.14 Cronologie Donau.xlsx').relative_to(R).as_posix()+'\n22 comunicări distincte, 24 ID-uri reexportate; 18 originale distincte reutilizate și verificate SHA-256. Ciorna este separată. Nu s-a trimis niciun mesaj nou, nu s-a efectuat o plată și nu s-a făcut reconciliere bancară.\n\n'
def backup(p):
 q=P/'Istoric'/p.relative_to(R);q.parent.mkdir(parents=True,exist_ok=True)
 if not q.exists():shutil.copy2(p,q)
def write(p,t):p.write_text(t,encoding='utf-8-sig')
for name in ['2026.10.14 Status proiect.txt','2026.10.14 Log progres proiect.txt']:
 p=R/name;backup(p);write(p,note+'ISTORIC PĂSTRAT\n\n'+p.read_text(encoding='utf-8-sig'))
for partner in ['Donau Versicherung','Maritczak - Commerz Inkasso','CERHA HEMPEL']:
 candidates=sorted((B/'Parteneri').glob('* Log discutii - '+partner+'.txt'))
 old=candidates[-1];new=B/('Parteneri/2026.10.14 Log discutii - '+partner+'.txt')
 history=old.read_text(encoding='utf-8-sig')
 if new.exists():backup(new)
 write(new,note+'ISTORIC PĂSTRAT\n\n'+history)
p=B/'2026.10.14 Registru comunicatii.json';backup(p);reg=json.loads(p.read_text(encoding='utf-8-sig'))
for s in reg['statusuri']:
 if s['partener'] in ['Donau Versicherung','Maritczak - Commerz Inkasso']:
  s['status']=status if s['partener']=='Donau Versicherung' else '2026.10.14: ultimul răspuns direct Commerz rămâne 2026.09.30: 5.326,15 EUR, inclusiv Q4; termen 2026.10.12 depășit. Nu există amânare, suspendare a recuperării sau încasare confirmată prin această verificare.'
  s.update(urmatorul_pas=nextstep,verificat='2026.10.14',sursa_status=[src,'2c273256-a0bf-408b-bb22-09718372e888','aa653a96-7eee-4133-a41d-5892ea79788d'])
reg.setdefault('actualizari_punctuale',[]).append({'data':'2026.10.14','subiect':'DONAU cronologie și acoperire','status':status,'sursa':src,'nota':'Reverificare și reexport; fără numărarea dublă a comunicărilor sau anexelor deja salvate.'})
reg['acoperire']['verificare_punctuala_Donau_2026_10_14']={'office_last_sync':'2026-10-14T09:32:19.699001Z','cosmin_ig_last_sync':'2026-10-09T23:30:10.009451Z','mesaje_noi':0,'emailuri_ID_reexportate':24,'randuri_cronologie':22,'originale_unice_reutilizate':18,'draft':'aa653a96-7eee-4133-a41d-5892ea79788d','draft_status':'pending','sursa':src}
p.write_text(json.dumps(reg,ensure_ascii=False,indent=2),encoding='utf-8')
p=R/'folder map/README.md';backup(p)
p.write_text('# 2026.10.14 — DONAU: stare actuală și tabel de corespondență\n\nDosar: `../05. Asigurari/Asigurare cladire/2026.10.14 Situatie Donau/`. Citește `2026.10.14 Situatie actuala Donau.txt` și `2026.10.14 Cronologie Donau.xlsx` (22 de comunicări, cele cinci coloane cerute).\n\n'+status+' '+nextstep+' Jurnalele proiectului și partenerilor sunt actualizate; originalele și exporturile sunt în subfolderele Anexe/Surse. Acoperirea generală anterioară rămâne distinctă; această verificare privește DONAU.\n\nISTORIC PĂSTRAT\n\n'+p.read_text(encoding='utf-8'),encoding='utf-8')
write(P/'2026.10.14 Jurnal actualizare.txt',note)
# Read-only verification of the exported workbook.
w=load_workbook(P/'2026.10.14 Cronologie Donau.xlsx');s=w.active
assert s.max_row==30 and s.max_column==5 and len(s.tables)==1
assert s['A9'].value.strftime('%Y.%m.%d')=='2026.05.27' and s['A30'].value.strftime('%Y.%m.%d')=='2026.10.06'
assert all(s.cell(i,1).data_type=='d' for i in range(9,31))
assert sum(c.comment is not None for row in s for c in row)==22
assert all(s.cell(i,2).value in ['Email trimis','Email primit'] for i in range(9,31))
print(json.dumps({'status':'jurnale actualizate, verificare Excel reusita','randuri':22,'coloane':5,'data_format':s['A9'].number_format},ensure_ascii=False))
