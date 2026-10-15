from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parent.parent
p=R/'05. Asigurari/Asigurare cladire/2026.10.14 Situatie Donau/2026.10.14 Cronologie Donau.xlsx'
w=load_workbook(p);s=w.active
for c in ['A9','A30','D30']:print(c,repr(s[c].value),s[c].data_type,s[c].number_format)
print('tables',list(s.tables),'notes',sum(c.comment is not None for row in s for c in row))
