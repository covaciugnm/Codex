from pathlib import Path
import json, shutil, hashlib, datetime, re
from email.message import EmailMessage
from email.policy import SMTP
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from docx import Document

R=Path(__file__).resolve().parent.parent
B=R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail'
P=R/'08. Corespondenta/2026.09.30 Revenire Donau - Maritczak 2616052'
P.mkdir(exist_ok=True)
S=json.loads((R/'folder map/2026.09.30 Donau verificare surse.json').read_text(encoding='utf-8'))
D=json.loads((R/'folder map/2026.09.30 Donau ciorna verificata.json').read_text(encoding='utf-8'))
assert D['status']=='pending' and D['to']==['office@maritczak.at'] and D['cc']==['bogdan.capra@cerhahempel.com']
H=P/'2026.09.30 Istoric inainte de revenire'
H.mkdir(exist_ok=True)
def backup(p):
    dest=H/p.relative_to(R)
    dest.parent.mkdir(parents=True,exist_ok=True)
    if p.exists() and not dest.exists(): shutil.copy2(p,dest)
def savej(p,obj): p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def txt(p,t): p.write_text(t,encoding='utf-8-sig')
for e in S['emails']:
    stamp=e['received_at'][:10].replace('-','.')
    status='TRIMIS' if e['is_sent'] else 'PRIMIT'
    name=f'{stamp} {status} '+('Cerere justificare Maritczak' if e['is_sent'] else 'Capra notificare somatie')
    savej(P/(name+'.json'),dict(e,data_lizibila=stamp,status_arhiva=status))
    header=f"{stamp} | {status}\nData export: 2026.09.30\nSubiect: {e['subject']}\nDe la: {e['from_address']}\nCatre: {', '.join(e['to'])}\nCC: {', '.join(e.get('cc',[])) or '—'}\nID Eva-Mail: {e['id']}\nSursa: export text Eva-Mail; MIME original indisponibil. Textul sursa de mai jos este nemodificat.\n\n"
    txt(P/(name+'.txt'),header+e['text'])
pdf=R/'05. Asigurari/Asigurare cladire/2026.09.14 Somatie Commerz-Inkasso 2616052 (Mitteilung Commerz-Inkasso_A&C Wohnart).pdf'
dest=P/'2026.09.09 Somatie Commerz-Inkasso 2616052.pdf'
shutil.copy2(pdf,dest)
sha=hashlib.sha256(dest.read_bytes()).hexdigest()
assert sha=='2c1c0e9e8f85447b0b699941f23e19d3cc070e56f43af42bb2589528f2b33025'
savej(P/'2026.09.30 Registru atasament justificativ.json',{'email_id':'6b7b1131-4ea1-482b-985c-deb6ec06fae8','attachment_id':'589967a1-945d-4cee-8e58-53d111190749','nume_original':'Mitteilung Commerz-Inkasso_A&C Wohnart.pdf','cale':dest.relative_to(R).as_posix(),'sha256':sha,'statut':'PRIMIT; original local verificat fata de index; pagina 1 verificata vizual','data_document':'2026.09.09','data_primirii':'2026.09.14','nota':'Cele trei imagini mici din semnatura nu sunt anexe justificative; metadatele lor sunt pastrate in export.'})
savej(P/'2026.09.30 Verificare Eva-Mail.json',S)
draftfile=P/'2026.09.30 DRAFT NETRIMIS - Revenire Donau.txt'
head=f"2026.09.30 | DRAFT, NETRIMIS | EVA status: {D['status']}\nDe la: {D['account_email']}\nCatre: {', '.join(D['to'])}\nCC: {', '.join(D['cc'])}\nSubiect: {D['subject']}\nID draft Eva-Mail: {D['id']}\nRaspuns la ID: {D['reply_to_email_id']}\nAtasamente: niciunul\nTrimiterea se face din EVA; nu exista confirmare de trimitere.\n\n"
txt(draftfile,head+D['body'])
savej(P/'2026.09.30 DRAFT NETRIMIS - Revenire Donau.json',dict(D,data_lizibila='2026.09.30',status_arhiva='DRAFT, NETRIMIS'))
em=EmailMessage(policy=SMTP)
em['From']=D['account_email'];em['To']=', '.join(D['to']);em['Cc']=', '.join(D['cc']);em['Subject']=D['subject'];em['X-Unsent']='1';em['X-Eva-Draft-ID']=D['id']
em['In-Reply-To']=S['emails'][0]['thread_id'];em['References']=S['emails'][0]['thread_id']
em.set_content(D['body'])
(P/'2026.09.30 DRAFT NETRIMIS - Revenire Donau.eml').write_bytes(em.as_bytes())
status='2026.09.30: verificare directa Eva-Mail; niciun raspuns identificat la cererea trimisa in 2026.09.15 pentru dosarul 2616052 / polita 2044001194, suma solicitata 3.734,34 EUR. Revenire in germana salvata in EVA ca DRAFT, NETRIMIS, catre office@maritczak.at, CC bogdan.capra@cerhahempel.com. Termen solicitat pentru acte: 2026.10.07. Prelungirea termenului de plata 2026.09.21 si suspendarea demersurilor nu sunt confirmate.'
nextstep='Trimiterea ciornei din EVA; apoi urmarirea documentelor si confirmarii scrise privind termenul de analiza. Termenul 2026.10.07 este doar propus in draft. Verificarea nu reprezinta confirmare a datoriei sau a acoperirii.'
scope='Cautare in toate casutele accesibile Eva-Mail, dupa dosar, polita, nume si expeditor. Office sincronizat 2026.09.30 11:30 UTC / 14:30 Romania; Gmail 2026.09.30 11:46 UTC; ipec 2026.09.30 11:28 UTC. cosmin@ig.ro are ultima sincronizare 2026.08.12; mesajele ulterioare din acea casuta nu pot fi excluse. Arhivarea generala a proiectului ramane IN CURS; cautarea directa din acest dosar este distincta.'
source=draftfile.relative_to(R).as_posix()
audit='2026.09.30 | Verificare si revenire Donau\n\nStatus scurt: '+status+'\n\nUltimul raspuns al avocatului/Inkasso: niciunul identificat.\nMesaj conex: Capra, 2026.09.17, ID b8bab41b-0cd5-40a8-8467-ca9c98c344d4, aminteste copia politei transmisa in 2026.03.19; nu raspunde cererii de justificare si prelungire.\n\nUrmatorul pas: '+nextstep+'\n\nAcoperire: '+scope+'\n\nSurse: cerere trimisa b487e95f-8a9f-4865-9836-ae95481ce360; notificare Capra 6b7b1131-4ea1-482b-985c-deb6ec06fae8; draft '+D['id']+' (recitit si verificat dupa salvare). PDF somație verificat vizual si SHA-256 fata de index. Rezultatele cautarilor sunt salvate in 2026.09.30 Verificare Eva-Mail.json.\n'
txt(P/'2026.09.30 Jurnal verificare si revenire.txt',audit)
partner='Maritczak - Commerz Inkasso'
regpath=B/'2026.09.30 Registru comunicatii.json'
manualpath=R/'folder map/2026.09.30 statusuri verificate.json'
reg=json.loads(regpath.read_text(encoding='utf-8'));manual=json.loads(manualpath.read_text(encoding='utf-8'))
backup(regpath);backup(manualpath)
old=next(s for s in reg['statusuri'] if s['partener']==partner)
oldstatus=old['status'];oldnext=old['urmatorul_pas']
old.update(status=status,urmatorul_pas=nextstep,ultima_comunicare='2026.09.30 (draft)',sursa_status=[source,D['id'],'b487e95f-8a9f-4865-9836-ae95481ce360'])
manual[partner]={'status':status,'urmatorul_pas':nextstep,'surse':old['sursa_status']}
row={'data':'2026.09.30','ora_utc':D['created_at'][11:16],'partener':partner,'domeniu':'Creante - Asigurare','sens':'DRAFT, NETRIMIS','subiect':D['subject'],'expeditor':D['account_email'],'destinatari':'; '.join(D['to']),'cc':'; '.join(D['cc']),'id':D['id'],'cale':source,'atasamente':0,'trunchiat':False}
reg.setdefault('ciorne',[]).append(row)
reg.setdefault('verificari_punctuale',[]).append({'data':'2026.09.30','partener':partner,'rezultat':status,'acoperire':scope,'sursa':(P/'2026.09.30 Verificare Eva-Mail.json').relative_to(R).as_posix()})
savej(regpath,reg);savej(manualpath,manual)
targets=[R/'2026.09.30 Status proiect.txt',B/'Parteneri/2026.09.30 Log discutii - Maritczak - Commerz Inkasso.txt']
for p in targets:
    backup(p);t=p.read_text(encoding='utf-8-sig');t=t.replace(oldstatus,status).replace(oldnext,nextstep)
    txt(p,t+'\n\n2026.09.30 | Surse revenire Donau: '+source+' | ID draft '+D['id']+'\n'+scope+'\n')
for p in [R/'2026.09.30 Log progres proiect.docx',B/'Parteneri/2026.09.30 Log discutii - Maritczak - Commerz Inkasso.docx']:
    if p.exists():
        backup(p);doc=Document(p)
        for par in doc.paragraphs:
            if oldstatus in par.text or oldnext in par.text: par.text=par.text.replace(oldstatus,status).replace(oldnext,nextstep)
        doc.add_paragraph('2026.09.30 | Sursa revenire Donau: '+source+' | ID draft '+D['id'])
        doc.add_paragraph(scope);doc.save(p)
for name in ['Donau Versicherung','CERHA HEMPEL']:
    note='2026.09.30 | Dosar Donau 2616052: '+status+'\nUrmatorul pas: '+nextstep+'\nSursa: '+source
    p=B/('Parteneri/2026.09.30 Log discutii - '+name+'.txt');backup(p);txt(p,note+'\n\n'+p.read_text(encoding='utf-8-sig'))
    p=p.with_suffix('.docx')
    if p.exists():
        backup(p);doc=Document(p);doc.paragraphs[0].insert_paragraph_before(note);doc.save(p)
wbpath=R/'2026.09.30 Log progres proiect.xlsx';backup(wbpath);wb=load_workbook(wbpath)
ws=wb['Status pe scurt'];cols={c.value:c.column for c in ws[1]}
for cells in ws.iter_rows(min_row=2):
    if cells[0].value==partner:
        for k in ['status','urmatorul_pas','ultima_comunicare','sursa_status']: ws.cell(cells[0].row,cols[k],'; '.join(old[k]) if isinstance(old[k],list) else old[k])
ws=wb['Comunicatii'];headers=[c.value for c in ws[1]]
if not any(c.value==D['id'] for c in ws['J']):
    ws.append([row.get(k) for k in headers])
    for c in ws[ws.max_row]:c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(ws.max_row,headers.index('cale')+1).hyperlink=draftfile.as_uri();ws.auto_filter.ref=ws.dimensions
wb['Acoperire'].append(['2026.09.30 Verificare punctuala Donau',scope+' Ciorna noua este NETRIMISA; nu este inclusa in numarul mesajelor primite/trimise salvate.'])
wb.save(wbpath)
check=load_workbook(wbpath);assert any(c.value==D['id'] for c in check['Comunicatii']['J'])
assert 'DRAFT, NETRIMIS' in draftfile.read_text(encoding='utf-8-sig')
print(json.dumps({'dosar':str(P),'draft_id':D['id'],'status':D['status'],'pdf_sha256':sha,'fisiere':len(list(P.glob('*.*'))),'jurnale_actualizate':True},ensure_ascii=False))
