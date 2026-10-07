from pathlib import Path
import json,shutil,copy,hashlib
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from docx import Document
R=Path(__file__).resolve().parent.parent
B=R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail'
P=R/'08. Corespondenta/2026.09.30 Raspuns Commerz - analiza Donau 2616052'
H=P/'2026.09.30 Istoric inainte de raspuns';H.mkdir(exist_ok=True)
def backup(p):
 q=H/p.relative_to(R);q.parent.mkdir(parents=True,exist_ok=True)
 if p.exists() and not q.exists():shutil.copy2(p,q)
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def txt(p,x):p.write_text(x,encoding='utf-8-sig')
emails=[json.loads(p.read_text(encoding='utf-8')) for p in [P/'2026.09.30 PRIMIT Raspuns Stefanie Gruber.json',P/'2026.09.30 TRIMIS Revenire Maritczak.json']]
recv=emails[0];sent=emails[1];partner='Maritczak - Commerz Inkasso'
source=(P/'2026.09.30 Analiza juridica Donau.txt').relative_to(R).as_posix()
status='2026.09.30: Commerz/Stefanie Gruber raspunde; solicita 5.326,15 EUR, inclusiv prima octombrie 2026–ianuarie 2027 de 1.591,81 EUR, si comunica termenul 2026.10.12. Total: 4.775,43 EUR prime + 550,72 EUR costuri/dobanda. Invoca doua emailuri nelivrate si scrisoare postala din 2026.09.17. Revenirea este confirmata TRIMISA la 2026.09.30, cu Capra in CC. Raspunsul primit nu are CC. Acoperirea actuala, reducerea politei si inghetarea costurilor nu sunt confirmate.'
nextstep='Pana la 2026.10.12: clarificare cu Capra a principalului/alocarii si a acoperirii, verificare costuri si prescriptie Q4; completare documente Pfeiffer (autorizatie, descriere lucrari, confirmare instalator) si cerere decizie asupra reducerii. Nu dubla prima Q4 deja inclusa in 5.326,15 EUR.'
regpath=B/'2026.09.30 Registru comunicatii.json';mp=R/'folder map/2026.09.30 statusuri verificate.json';ap=B/'2026.09.30 Registru atasamente.json'
for p in [regpath,mp,ap]:backup(p)
reg=json.loads(regpath.read_text(encoding='utf-8'));manual=json.loads(mp.read_text(encoding='utf-8'));ats=json.loads(ap.read_text(encoding='utf-8'))
old={x['partener']:copy.deepcopy(x) for x in reg['statusuri'] if x['partener'] in [partner,'Donau Versicherung']}
for s in reg['statusuri']:
 if s['partener'] in old:
  s.update(status=status,urmatorul_pas=nextstep,ultima_comunicare='2026.09.30',sursa_status=[recv['id'],sent['id'],source])
  if s['partener']==partner:s.update(ultimul_primit='2026.09.30',subiect_ultimul_primit=recv['subject'])
  manual[s['partener']]={'status':status,'urmatorul_pas':nextstep,'surse':s['sursa_status']}
rows=[]
for e in emails:
 direction='TRIMIS' if e['is_sent'] else 'PRIMIT'
 p=P/('2026.09.30 '+direction+' '+('Revenire Maritczak' if e['is_sent'] else 'Raspuns Stefanie Gruber')+'.txt')
 rows.append({'data':'2026.09.30','ora_utc':e['received_at'][11:16],'partener':partner,'domeniu':'Creante - Asigurare','sens':direction,'subiect':e['subject'],'expeditor':e['from_address'],'destinatari':'; '.join(e['to']),'cc':'; '.join(e['cc']),'id':e['id'],'cale':p.relative_to(R).as_posix(),'atasamente':len(e['attachments']),'trunchiat':e['truncated']})
added=[]
for row in rows:
 if not any(x['id']==row['id'] for x in reg['comunicatii']):reg['comunicatii'].append(row);added.append(row)
for s in reg['statusuri']:
 if s['partener']==partner:s['mesaje']=sum(x['partener']==partner for x in reg['comunicatii'])
reg['acoperire']['mesaje_salvate']=len(reg['comunicatii'])
reg['acoperire']['mesaje_suplimentare_fata_de_lista']=reg['acoperire'].get('mesaje_suplimentare_fata_de_lista',0)+len(added)
reg['acoperire']['corpuri_trunchiate_de_Eva']=sum(bool(x.get('trunchiat')) for x in reg['comunicatii'])
regs=json.loads((P/'2026.09.30 Registru anexe.json').read_text(encoding='utf-8'));byname={x['nume_original']:x for x in regs};newatts=[]
for a in recv['attachments']:
 if any(x['id']==a['id'] for x in ats):continue
 z=byname[a['name']];p=R/z['cale'];assert p.stat().st_size==a['size']
 row={'id':a['id'],'name':a['name'],'size':a['size'],'email_id':recv['id'],'date':recv['received_at'],'cale':z['cale'],'sha256':z['sha256'],'stare':'salvat original integral','data':'2026.09.30','partener':partner,'nume_original':a['name'],'octeti':a['size']}
 ats.append(row);newatts.append(row)
reg['acoperire']['atasamente_salvate']=sum(x.get('stare')=='salvat original integral' for x in ats)
reg['acoperire']['atasamente_identificate']+=len(newatts)
reg['acoperire']['atasamente_documentare']+=sum(not x['name'].lower().endswith('.jpg') for x in newatts)
reg['acoperire']['atasamente_documentare_salvate']+=sum(not x['name'].lower().endswith('.jpg') for x in newatts)
reg.setdefault('actualizari_punctuale',[]).append({'data':'2026.09.30','subiect':'Raspuns Commerz Donau 2616052','sursa':source,'status':status,'nota':'Doua mesaje suplimentare; 10 ID-uri anexe reprezentate de 5 fisiere originale. Cautarea generala anterioara ramane completa pentru lista sa.'})
save(regpath,reg);save(mp,manual);save(ap,ats)
save(B/'Surse/2026.09.30 Actualizare raspuns Donau.json',emails)
notes='2026.09.30 | ACTUALIZARE ULTERIOARA RASPUNSULUI COMMERZ\n\nStatus scurt: '+status+'\n\nUltimul raspuns: 2026.09.30, 15:51 Romania, Stefanie Gruber, '+recv['id']+'\n\nUrmatorul pas: '+nextstep+'\n\nSursa: '+source+'\n\n'
for name in [partner,'Donau Versicherung','CERHA HEMPEL']:
 p=B/('Parteneri/2026.09.30 Log discutii - '+name+'.txt');backup(p);txt(p,notes+'ISTORIC PASTRAT MAI JOS\n\n'+p.read_text(encoding='utf-8-sig'))
 p=p.with_suffix('.docx')
 if p.exists():
  backup(p);d=Document(p);d.paragraphs[0].insert_paragraph_before(notes);d.save(p)
p=R/'2026.09.30 Status proiect.txt';backup(p);t=p.read_text(encoding='utf-8-sig')
for o in old.values():t=t.replace(o['status'],status).replace(o['urmatorul_pas'],nextstep)
txt(p,notes+t)
p=R/'2026.09.30 Log progres proiect.docx';backup(p);d=Document(p)
for par in d.paragraphs:
 for o in old.values():
  if o['status'] in par.text or o['urmatorul_pas'] in par.text:par.text=par.text.replace(o['status'],status).replace(o['urmatorul_pas'],nextstep)
d.paragraphs[0].insert_paragraph_before(notes);d.save(p)
p=R/'2026.09.30 Log progres proiect.xlsx';backup(p);wb=load_workbook(p)
ws=wb['Status pe scurt'];cols=[c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
 if row[0].value in old:
  s=next(x for x in reg['statusuri'] if x['partener']==row[0].value)
  for i,k in enumerate(cols,1):v=s.get(k);ws.cell(row[0].row,i,'; '.join(v) if isinstance(v,list) else v)
for name,data in [('Comunicatii',rows),('Atasamente',newatts)]:
 ws=wb[name];cols=[c.value for c in ws[1]];idcol=cols.index('id')+1;existing={ws.cell(i,idcol).value for i in range(2,ws.max_row+1)}
 for row in data:
  if row['id'] not in existing:
   ws.append([row.get(k) for k in cols]);ws.cell(ws.max_row,cols.index('cale')+1).hyperlink=(R/row['cale']).as_uri()
   for c in ws[ws.max_row]:c.alignment(wrap_text=True) if False else None;c.alignment=Alignment(wrap_text=True,vertical='top')
 ws.auto_filter.ref=ws.dimensions
ws=wb['Acoperire']
for row in ws.iter_rows(min_row=2):
 if row[0].value in reg['acoperire']:ws.cell(row[0].row,2,str(reg['acoperire'][row[0].value]))
ws.append(['2026.09.30 Raspuns Commerz',source+' | 5.326,15 EUR ceruti, termen 2026.10.12; include Q4.'])
wb.save(p)
# Preserve dated prior audit; add a clear update sheet to both current financial workbooks.
for rel in ['10. Banci + Extrase de cont/2026.09.30 Reconciliere completa facturi si plati.xlsx','08. Corespondenta/2026.09.30 De platit - Schallergasse 35 - verificat.xlsx']:
 p=R/rel
 if not p.exists():continue
 backup(p);w=load_workbook(p);sh=w.create_sheet('Donau raspuns 2026.09.30')
 for row in [['Actualizare documentara','Valoare'],['Data','2026.09.30'],['Status',status],['Prime 3 trimestre',4775.43],['Accesorii',550.72],['Total solicitat - NU sold acceptat',5326.15],['Termen comunicat','2026.10.12'],['Atentie','Include deja prima Q4 1.591,81 EUR; nu adunati din nou la total. Nu exista plata confirmata in aceasta analiza.'],['Urmatorul pas',nextstep],['Sursa',source]]:sh.append(row)
 sh.column_dimensions['A'].width=42;sh.column_dimensions['B'].width=110
 for row in sh.iter_rows():
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 sh.cell(sh.max_row,2).hyperlink=(R/source).as_uri();w.save(p)
txt(P/'2026.09.30 Jurnal actualizare dosar.txt',notes+'Arhiva generala: '+str(reg['acoperire']['mesaje_salvate'])+' mesaje salvate. Documentele de analiza nu sunt acte ale creditorului. Nu s-a trimis un mesaj nou si nu s-a efectuat plata.\n')
report=P/'2026.09.30 Analiza juridica Donau.txt';d=Document();d.add_heading('2026.09.30 | Analiza Donau / Commerz 2616052',0)
for block in report.read_text(encoding='utf-8').split('\n\n'):d.add_paragraph(block)
d.save(report.with_suffix('.docx'))
print(json.dumps({'status':'actualizat','mesaje_noi':len(added),'anexe_id_noi':len(newatts),'total_mesaje':reg['acoperire']['mesaje_salvate'],'dosar':str(P)},ensure_ascii=False))
