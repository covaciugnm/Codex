from pathlib import Path
import json, shutil, hashlib
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from datetime import datetime
from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import Alignment
import fitz

R=Path(__file__).resolve().parent.parent
P=R/'08. Corespondenta/2026.10.01 Donau - cerere reziliere'
B=R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail'
d=json.loads((P/'2026.10.01 DRAFT Reziliere Donau.json').read_text(encoding='utf-8'))
assert d['status']=='pending' and d['id']=='141ce321-f1e8-4fdb-ae64-d26afeac4122'
source=(P/'2026.10.01 DRAFT Reziliere Donau.txt').relative_to(R).as_posix()
status='2026.10.01: cerere de incetare a politei 2044001194 pregatita in Eva-Mail, DRAFT NETRIMIS. Solicita incetare imediata prin acord; independent, notificare de incetare la prima data legal/contractual permisa. Catre DONAU si Loschy; CC Capra, Maritczak si Commerz/Gruber. Confirmare ceruta pana la 2026.10.07; incetarea nu este confirmata.'
nextstep='Trimiterea ciornei din Eva-Mail; apoi verificarea dovezii de trimitere si a confirmarii DONAU privind data incetarii, acoperirea si decontul final. Termenul Commerz 2026.10.12 ramane distinct; cererea nu il suspenda. Suma solicitata anterior: 5.326,15 EUR, include Q4; nu constituie sold acceptat.'
last='Ultimul raspuns relevant arhivat: Commerz/Stefanie Gruber, 2026.09.30, ID 1fa19552-d1fd-459e-92b3-57ed5319d18c. Nu s-a efectuat o noua verificare generala a corespondentei in aceasta actiune.'
note='2026.10.01 | DONAU - CERERE REZILIERE\n\nStatus scurt: '+status+'\n\n'+last+'\n\nUrmatorul pas: '+nextstep+'\n\nSursa: '+source+'\nID Eva-Mail: '+d['id']+'\n\n'
def txt(p,t):p.write_text(t,encoding='utf-8-sig')
def js(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def backup(p):
 q=P/'2026.10.01 Istoric inainte de actualizare'/p.relative_to(R)
 q.parent.mkdir(parents=True,exist_ok=True)
 if not q.exists():shutil.copy2(p,q)
headers='2026.10.01 | DRAFT NETRIMIS\nID Eva-Mail: '+d['id']+'\nCreat la (tehnic): '+d['created_at']+'\nDe la: '+d['account_email']+'\nCatre: '+', '.join(d['to'])+'\nCC: '+', '.join(d['cc'])+'\nSubiect: '+d['subject']+'\nAtasamente: niciunul\n\n'
txt(R/source,headers+d['body'])
m=EmailMessage(policy=SMTP)
for k,v in [('From',d['account_email']),('To',', '.join(d['to'])),('Cc',', '.join(d['cc'])),('Subject',d['subject']),('X-Unsent','1'),('X-Eva-Draft-ID',d['id'])]:m[k]=v
m['Date']=format_datetime(datetime.fromisoformat(d['created_at'].replace('Z','+00:00')))
m.set_content(d['body'])
(P/'2026.10.01 DRAFT Reziliere Donau.eml').write_bytes(m.as_bytes())
pol=R/'05. Asigurari/Asigurare cladire/814BC40FC7831FE197E21FDA1089ECC3_Polizzenkopie.pdf'
idx=json.loads((R/'folder map/inventar.json').read_text(encoding='utf-8'))
entry=next(x for x in idx if x['cale']==pol.relative_to(R).as_posix())
assert hashlib.sha256(pol.read_bytes()).hexdigest()==entry['sha256']
with fitz.open(pol) as pdf:
 assert '2044001194' in pdf[0].get_text() and '2900010498' in pdf[0].get_text()
txt(P/'2026.10.01 Jurnal actualizare dosar.txt',note+'Polita originala verificata: '+pol.relative_to(R).as_posix()+'\nSHA-256: '+entry['sha256']+'\nAdresa generala DONAU verificata: https://www.donauversicherung.at/kontakt\nFara atasamente noi; documentele sursa sunt pastrate in arhiva existenta.\n')
# New dated editions preserve the previous journals unchanged.
txt(R/'2026.10.01 Status proiect.txt',note+'ISTORIC PASTRAT\n\n'+(R/'2026.09.30 Status proiect.txt').read_text(encoding='utf-8-sig'))
for name in ['Donau Versicherung','Maritczak - Commerz Inkasso','CERHA HEMPEL']:
 old=B/('Parteneri/2026.09.30 Log discutii - '+name+'.txt')
 new=old.with_name('2026.10.01 Log discutii - '+name+'.txt')
 txt(new,note+'ISTORIC PASTRAT\n\n'+old.read_text(encoding='utf-8-sig'))
 doc=Document(old.with_suffix('.docx'));doc.paragraphs[0].insert_paragraph_before(note);doc.save(new.with_suffix('.docx'))
doc=Document(R/'2026.09.30 Log progres proiect.docx');doc.paragraphs[0].insert_paragraph_before(note);doc.save(R/'2026.10.01 Log progres proiect.docx')
regp=B/'2026.09.30 Registru comunicatii.json';backup(regp)
reg=json.loads(regp.read_text(encoding='utf-8'))
for s in reg['statusuri']:
 if s['partener'] in ['Donau Versicherung','Maritczak - Commerz Inkasso']:
  s['status']=status+' Ultima creanta comunicata: 5.326,15 EUR, termen 2026.10.12.'
  s['urmatorul_pas']=nextstep
  s['sursa_status']=[d['id'],source,'1fa19552-d1fd-459e-92b3-57ed5319d18c']
reg.setdefault('drafturi',[]).append({'data':'2026.10.01','id':d['id'],'statut':'DRAFT NETRIMIS','cale':source,'expeditor':d['account_email'],'destinatari':d['to'],'cc':d['cc'],'subiect':d['subject']})
reg.setdefault('actualizari_punctuale',[]).append({'data':'2026.10.01','subiect':'Cerere reziliere Donau','status':status,'sursa':source,'nota':'Draft verificat; numarul mesajelor primite/trimise si acoperirea preluarii nu sunt modificate.'})
js(regp,reg)
mp=R/'folder map/2026.09.30 statusuri verificate.json';backup(mp);manual=json.loads(mp.read_text(encoding='utf-8'))
for name in ['Donau Versicherung','Maritczak - Commerz Inkasso']:manual[name]={'status':status,'urmatorul_pas':nextstep,'surse':[d['id'],source]}
js(mp,manual)
w=load_workbook(R/'2026.09.30 Log progres proiect.xlsx');ws=w['Status pe scurt'];cols=[c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
 if row[0].value in ['Donau Versicherung','Maritczak - Commerz Inkasso']:
  s=next(x for x in reg['statusuri'] if x['partener']==row[0].value)
  for i,k in enumerate(cols,1):
   v=s.get(k);ws.cell(row[0].row,i,'; '.join(v) if isinstance(v,list) else v)
sh=w.create_sheet('Donau draft 2026.10.01')
for row in [['Camp','Valoare'],['Data','2026.10.01'],['Status',status],['Ultimul raspuns',last],['Urmatorul pas',nextstep],['ID Eva-Mail',d['id']],['De la',d['account_email']],['Catre',', '.join(d['to'])],['CC',', '.join(d['cc'])],['Subiect',d['subject']],['Sursa',source]]:sh.append(row)
sh.column_dimensions['A'].width=25;sh.column_dimensions['B'].width=115
for row in sh.iter_rows():
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
sh.cell(sh.max_row,2).hyperlink=(R/source).as_uri()
w.save(R/'2026.10.01 Log progres proiect.xlsx')
read=R/'folder map/README.md';backup(read)
t=read.read_text(encoding='utf-8');head='# Harta arhivei Schallergasse 35'
new='\n\n## 2026.10.01 — Cerere reziliere DONAU\n\nPunctul curent de intrare: `../2026.10.01 Status proiect.txt` si `../2026.10.01 Log progres proiect.xlsx`. Jurnalele Donau, Maritczak/Commerz si CERHA HEMPEL au versiuni `2026.10.01` in acelasi folder Parteneri. Editiile `2026.09.30` sunt pastrate ca istoric.\n\n'+status+' ID EVA `'+d['id']+'`. Dosar: `../08. Corespondenta/2026.10.01 Donau - cerere reziliere/`. '+nextstep+' Acoperirea generala a arhivei ramane cea din registrul central; aceasta actiune adauga un draft, fara o noua preluare generala.\n'
read.write_text(t.replace(head,head+new,1),encoding='utf-8')
assert 'X-Unsent: 1' in (P/'2026.10.01 DRAFT Reziliere Donau.eml').read_text(encoding='utf-8')
print(json.dumps({'status':'arhivat, DRAFT NETRIMIS','id':d['id'],'dosar':str(P),'jurnal':'2026.10.01 Log progres proiect.xlsx'},ensure_ascii=False))
