from pathlib import Path
import json,hashlib,shutil,csv,datetime,html
from openpyxl import load_workbook
from pypdf import PdfReader
ROOT=Path('D:/00. Downloads/Apartamente Viena/00. Cladire 1 Schallergasse 35, 1120 Wien (Meidling)')
WORK=Path(__file__).parent
D=ROOT/'04. Firme + Executie/08. Ofertanti electrice/Tongou - Conex Electronic/2026.10.08 Catalog comparativ'
data=json.loads((D/'Date structurate/2026.10.08 produse-normalizate.json').read_text(encoding='utf-8'))
docs=json.loads((D/'Date structurate/2026.10.08 registru manuale oficiale.json').read_text(encoding='utf-8'))
broken=json.loads((D/'Date structurate/2026.10.08 registru manuale.json').read_text(encoding='utf-8'))
blocks=json.loads((D/'Lucru/2026.10.08 blocuri verificare.json').read_text(encoding='utf-8'))
book=D/'2026.10.08 Tongou comparativ produse.xlsx'
w=load_workbook(book,data_only=False)
assert len(w.sheetnames)==6
found=[]
for b in blocks:
    s=w[b['sheet']];actual=[str(s.cell(b['start']+1,c).value) for c in range(2,2+len(b['skus']))]
    assert actual==b['skus'];found.extend(actual)
    for c,sku in enumerate(b['skus'],2):
        p=next(x for x in data if x['sku']==sku)
        assert s.cell(b['start']+2,c).value==p['name']
        for r in range(b['start']+3,b['end']-1):
            k=s.cell(r,1).value
            assert s.cell(r,c).value==p['technical'][k],(sku,k,s.cell(r,c).value,p['technical'][k])
assert len(found)==29 and len(set(found))==29
stock=[]
for r in range(6,35):
    s=w['Catalog si stoc'];sku=str(s.cell(r,1).value);p=next(x for x in data if x['sku']==sku)
    assert s.cell(r,7).value==p['technical']['Cantitate disponibilă (buc.)']
    assert s.cell(r,6).value==p['technical']['Preț cu TVA (RON)']
    stock.append({'sku':sku,'cantitate':s.cell(r,7).value,'status':s.cell(r,8).value,'sursa':p['source'],'metoda':p['stock_source'],'data':'2026.10.08'})
assert all(isinstance(x['cantitate'],int) for x in stock)
assert [x['sku'] for x in stock if x['cantitate']==0]==['43074']
for s in w:
    for row in s:
        assert all(c.data_type!='e' for c in row)
for d in docs:
    if d['status']=='DESCARCAT':
        p=D/d['path'];assert p.exists() and hashlib.sha256(p.read_bytes()).hexdigest()==d['sha256'];assert len(PdfReader(p).pages)==d['pages']
qa={'data':'2026.10.08','produse':29,'sku_unice_in_matrice':29,'produse_in_stoc':28,'produse_epuizate':1,'cantitati_exacte_identificate':29,'verificare_cos':'Nu a fost necesară: cantitatea numerică este publicată în HTML-ul fiecărei pagini; pentru SKU epuizat, 0 din statutul explicit.','pdf_oficiale_salvate':sum(x['status']=='DESCARCAT' for x in docs),'linkuri_conex_pdf_indisponibile':len(broken),'foi':w.sheetnames,'panouri_blocate':{s.title:str(s.freeze_panes) for s in w},'verificari':['29 SKU unice, fiecare prezent o singură dată în matrice','Denumiri și toate celulele tehnice comparate cu datele normalizate','Prețuri și stocuri din registrul Excel comparate cu extragerea sursă','EAN păstrat ca text','PDF-uri verificate prin hash și număr de pagini','Șase foi randate și inspectate vizual; observațiile și neconcordanța RCCB verificate','Fără celule Excel de eroare']}
(D/'Lucru/2026.10.08 Verificare finala.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
(D/'Date structurate/2026.10.08 Registru stocuri.json').write_text(json.dumps(stock,ensure_ascii=False,indent=2),encoding='utf-8')
guide='''2026.10.08 | CATALOG TONGOU — CONEX ELECTRONIC

Începe cu: 2026.10.08 Tongou comparativ produse.xlsx
Sursa catalogului: https://www.conexelectronic.ro/catalog/q/tongou?sort_by=price_desc
Paginare verificată: 24 + 5 = 29 rezultate. Fiecare SKU este inclus o singură dată.

ORGANIZARE
1 pol: niciun 1P simplu în rezultatele actuale. Subdiviziunea 1P+N cuprinde 6 produse.
2 poli: 10 produse, după titlul comercial; 43067 are descriere 1P+N și observație explicită.
3 poli: 7 produse. 4 poli: 6 produse.
Subdiviziuni pe tip: disjunctoare MCB, diferențiale RCBO, diferențial de clarificat RCCB/RCBO, întrerupătoare smart, SPD.
Modelele sunt pe coloane: prima linie cod produs, a doua linie denumire originală, apoi caracteristici și observații.

LEGENDĂ
✓ verde: funcție confirmată de pagina produsului sau de tipul tehnic declarat, unde este precizat.
✓ albastru: funcție confirmată față de o absență/altă variantă explicită din grupă; rândul de diferențe include și specificații diferite.
✓ portocaliu: observație, neconcordanță sau document indisponibil; nu certifică o funcție incertă.
? = neprecizat/neconfirmat. Nu = funcție absentă pentru tipul aparatului ori pentru varianta de comunicație identificată. — = nu se aplică.
Lipsa unei funcții în descriere nu este transformată automat în Nu.
1P+N nu este echivalat automat cu 1P simplu sau cu 2 poli protejați.
Protecția electronică la supratensiune și protecția de impuls SPD sunt caracteristici distincte.

STOC ȘI PREȚ
28 produse în stoc; SKU 43074 este epuizat, 0 bucăți. Cantitățile numerice sunt disponibile în paginile publice, deci verificarea prin coș nu a fost necesară.
Snapshot 2026.10.08; valorile se pot modifica. Prețurile sunt RON, cu TVA, fără costul transportului. Sunt prețuri de catalog, nu ofertă fermă personalizată.

DOCUMENTAȚIE
Toate cele 5 URL-uri PDF distincte legate în cele 29 pagini Conex au răspuns HTTP 404: Smart Breaker, SMR1, TO-Q-SY2-JWT, Zigbee 2-3-4P și manualul etichetat 4G/LTE.
Au fost descărcate 12 PDF-uri oficiale alternative/suplimentare de la Tongou/Chayo/elcb.net: manuale, fișe de familie și documente de conformitate.
Două fișiere au aceeași denumire ca linkurile Conex, dar identitatea binară nu poate fi dovedită fără originale. Celelalte sunt alternative de familie, nu copii pretins identice.
Unele produse nu au PDF atașat pe Conex. Nu a fost inventată o fișă specifică fiecărui SKU.
Documentele de familie pot enumera variante de protocol, curbe sau tip diferențial; acestea nu sunt atribuite simultan tuturor produselor.
Registrul păstrează URL, status, calea locală, pagini, dimensiune, SHA-256 și limitele asocierii la SKU. O sursă suplimentară de catalog TORD4 nu a putut fi descărcată și este marcată EROARE.

NECONCORDANȚE IMPORTANTE
43092: Conex numește RCBO și curba C, dar declarația oficială TORD4-63 indică RCCB. Protecția la supracurent/scurtcircuit nu este confirmată; modelul livrat trebuie clarificat.
43073: 1–20 A în titlu versus 1–40 A în descriere; TOQCB2L în text versus SMR1 în atașament.
43067: 2P în titlu versus 1P+N în descriere.
43077/43078: 400 V în descrierea scurtă versus AC 90–295 V în specificații, fără definirea referinței tensiunilor.
43084–43087: neconcordanțe de familie TOQCB2/TO-Q-SY2 și coduri 2C63; vezi observațiile individuale.

FIȘIERE
Surse web: cele două pagini de catalog și 29 pagini de produs, HTML original și text relevant.
Date structurate: JSON original, JSON normalizat, CSV de caracteristici, stocuri și registre documente, texte extrase din PDF.
Datasheet si manuale: 12 PDF-uri originale descărcate, nemodificate.
Lucru: scripturi, previzualizări și verificări finale.
Actualizarea este documentare de catalog; nu reprezintă comandă, acceptare comercială sau selecție tehnică pentru proiect.
'''
(D/'2026.10.08 Ghid dosar si limite.txt').write_text(guide,encoding='utf-8')
request='''2026.10.08 | CERINȚE UTILIZATOR ÎN CHAT
Expeditor: utilizator. Destinatar: asistent. CC: nu se aplică. ID Eva-Mail: nu se aplică.
Statut: instrucțiuni în chat; nu reprezintă email primit/trimis/draft.
Subiect: catalog comparativ Tongou și documentație, Schallergasse, ofertanți electrice.

1. Extragerea tuturor produselor Tongou din linkul Conex și gruparea pe tip și număr de poli, modele pe coloane, cod produs prima linie, denumire a doua linie, apoi caracteristici și prezență/absență.
2. Diferențele se notează la observații, cu bifă de altă culoare și explicații.
3. Crearea dosarului în Schallergasse, la ofertanți electrice; salvarea lucrului, descărcarea tuturor datasheet-urilor și structurarea informațiilor.
4. Marcarea stocului epuizat și identificarea cantității, inclusiv prin creșterea cantității în coș dacă este necesar.

Implementare: cantitățile sunt publicate în pagină, astfel încât nu a fost necesar coșul. Limitările descărcării PDF sunt explicit documentate.
'''
(D/'2026.10.08 Cerinte utilizator.txt').write_text(request,encoding='utf-8')
rel=str(D.relative_to(ROOT)).replace('\\','/')
entry=f'''2026.10.08 | TONGOU / CONEX ELECTRONIC — CATALOG ECHIPAMENTE ELECTRICE

Status scurt: 29 produse catalogate și comparate pe poli/tip; stoc numeric documentat pentru toate. 28 disponibile, SKU 43074 epuizat. Excel cu bife colorate și observații; pagini sursă și date structurate salvate. 12 PDF-uri oficiale alternative descărcate; toate cele 5 linkuri PDF Conex returnează 404. SKU 43092 are conflict RCBO la Conex / RCCB în declarația TORD4 a fabricantului.
Ultimul răspuns: nu a fost consultată corespondența furnizorului în această lucrare; sursa este catalogul web accesat la 2026.10.08. Niciun email primit/trimis/draft în această documentare. ID Eva-Mail, expeditor, destinatari și CC: nu se aplică sursei web.
Următorul pas: clarificarea cu furnizorul a codurilor și specificațiilor neconcordante, obținerea fișelor specifice SKU și verificarea selecției de către proiectantul electric. Fără comandă sau acceptare de ofertă.
Sursă: {rel}/2026.10.08 Tongou comparativ produse.xlsx; registrele și originalele din același dosar; https://www.conexelectronic.ro/catalog/q/tongou?sort_by=price_desc
Acoperire: 29/29 rezultate pe două pagini, 5/5 URL-uri Conex documentate ca indisponibile; 12 PDF-uri alternative salvate. Acoperirea Eva-Mail anterioară nu este extinsă sau modificată de această lucrare.
'''
(D/'2026.10.08 Jurnal Tongou - Conex Electronic.txt').write_text(entry,encoding='utf-8')
partner=ROOT/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail/Parteneri/2026.10.08 Log discutii - Tongou Conex Electronic.txt'
partner.write_text(entry,encoding='utf-8')
for name in ['2026.10.08 Log progres proiect.txt','2026.10.08 Status proiect.txt']:
    p=ROOT/name;old=p.read_text(encoding='utf-8-sig') if p.exists() else ''
    marker='TONGOU / CONEX ELECTRONIC — CATALOG ECHIPAMENTE ELECTRICE'
    if marker not in old:
        backup=D/'Lucru'/('2026.10.08 Istoric anterior Tongou - '+name)
        if p.exists():shutil.copy2(p,backup)
        with p.open('a',encoding='utf-8') as f:f.write('\n\n'+entry)
readme=ROOT/'folder map/README.md';old=readme.read_text(encoding='utf-8-sig')
section=f'''\n\n## 2026.10.08 — Tongou / Conex Electronic, ofertanți electrice

Dosar: `../{rel}/`. Începe cu `2026.10.08 Tongou comparativ produse.xlsx` și `2026.10.08 Ghid dosar si limite.txt`. Toate cele 29 produse din catalog, 4 foi comparative, stocuri numerice, bife și observații. SKU 43074 epuizat. SKU 43092: conflict RCBO/RCCB, fără protecție la supracurent confirmată. 12 PDF-uri oficiale alternative salvate, 5 linkuri PDF Conex indisponibile HTTP 404; registrul distinge originalele indisponibile de documentele alternative. 1P+N este subdiviziune explicită, nu 1P simplu. Jurnalele proiectului și Tongou/Conex actualizate. Documentare web, fără email, comandă sau ofertă acceptată.\n'''
if '## 2026.10.08 — Tongou / Conex Electronic, ofertanți electrice' not in old:
    with readme.open('a',encoding='utf-8') as f:f.write(section)
for f in WORK.glob('2026.10.08 *'):
    if f.suffix in ['.py','.mjs']:shutil.copy2(f,D/'Lucru'/f.name)
for f in WORK.glob('2026.10.08 fabricant*.html'):shutil.copy2(f,D/'Surse web'/f.name)
manifest=[]
for p in sorted(D.rglob('*')):
    if p.is_file() and p.name!='2026.10.08 Manifest fisiere.json':manifest.append({'cale':str(p.relative_to(D)),'octeti':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(D/'Date structurate/2026.10.08 Manifest fisiere.json').write_text(json.dumps({'data':'2026.10.08','fisiere':manifest},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(qa,ensure_ascii=False,indent=2))
