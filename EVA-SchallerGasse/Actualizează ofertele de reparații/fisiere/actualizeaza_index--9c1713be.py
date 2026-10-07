"""Index local incremental. Nu modifica documentele sursa. Python 3.12."""
from pathlib import Path
import hashlib, json, csv, re, zipfile, tarfile, subprocess, datetime, collections, html
import fitz
from docx import Document
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image
import sqlite3, sys

BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE/'runtime'))
ROOT=BASE.parent
CACHE=BASE/'texte'
CACHE.mkdir(exist_ok=True)
STAMP=datetime.datetime.now().astimezone().isoformat(timespec='seconds')
OLD={}
if (BASE/'inventar.json').exists():
    OLD={r['cale']:r for r in json.loads((BASE/'inventar.json').read_text(encoding='utf-8'))}
SEVEN=Path('C:/Program Files/7-Zip/7z.exe')
TEXT={'.txt','.md','.csv','.html','.css','.py','.ics','.vcf','.kml','.ini','.json','.ps1'}

def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(4*1024*1024),b''):h.update(b)
    return h.hexdigest()

def decode(b):
    for enc in ('utf-8-sig','utf-16','cp1252','latin1'):
        try:return b.decode(enc)
        except (UnicodeError,LookupError):pass

def extract(p):
    e=p.suffix.lower(); info={}; text=''; status=''
    with p.open('rb') as f:magic=f.read(8)
    if e=='.tar' and magic.startswith(b'Rar!'):e='.rar';info['format_detectat']='RAR cu extensie TAR'
    if e=='.xlsx' and magic.startswith(bytes.fromhex('d0cf11e0')):e='.xls';info['format_detectat']='Office OLE cu extensie XLSX'
    if (not e or (e=='.doc' and magic.startswith(b'PK'))) and zipfile.is_zipfile(p):
        with zipfile.ZipFile(p) as z:
            names=z.namelist()
        if 'word/document.xml' in names:e='.docx';info['format_detectat']='DOCX cu extensie absentă sau DOC'
        else:e='.zip';info['format_detectat']='ZIP fără extensie'
    if p.name.startswith('~$'):return 'Fișier temporar Office; nu reprezintă documentul original.','temporar',info
    if e=='.pdf':
        with fitz.open(p) as doc:
            info['pagini']=len(doc); pages=[]; low=[]
            for i,page in enumerate(doc):
                t=page.get_text(sort=True)
                widgets=list(page.widgets() or [])
                if widgets:t+='\nCÂMPURI FORMULAR:\n'+'\n'.join(f'{w.field_name}: {w.field_value}' for w in widgets)
                if len(t.strip())<40:low.append(i+1)
                pages.append(f'\n--- PAGINA {i+1} ---\n{t}')
            info['pagini_fara_text_suficient']=low
            text='\n'.join(pages)
            status='text extras; verificare vizuală necesară' if low else 'text extras'
    elif e=='.docx':
        d=Document(p); a=[x.text for x in d.paragraphs]
        for ti,t in enumerate(d.tables):a.append(f'TABEL {ti+1}\n'+'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows))
        for s in d.sections:a.extend(x.text for x in s.header.paragraphs);a.extend(x.text for x in s.footer.paragraphs)
        text='\n'.join(a);status='text și tabele extrase'
    elif e=='.xlsx':
        wb=load_workbook(p,read_only=True,data_only=False);a=[];info['foi']=wb.sheetnames
        for ws in wb:
            a.append('FOAIE: '+ws.title)
            for ri,r in enumerate(ws.iter_rows(),1):
                vals=[f'{c.column_letter}={c.value}' for c in r if c.value is not None]
                if vals:a.append(f'Rând {ri}: '+' | '.join(vals))
        wb.close();text='\n'.join(a);status='celule și formule extrase; imaginile nu sunt interpretate'
    elif e in TEXT or (not e and p.stat().st_size<2_000_000):
        text=decode(p.read_bytes());status='text citit'
    elif e in {'.zip','.tar','.7z','.rar'}:
        if e=='.zip':
            with zipfile.ZipFile(p) as z:
                entries=[{'cale':x.filename,'octeti':x.file_size} for x in z.infolist()]
            text='\n'.join(f"{x['octeti']}\t{x['cale']}" for x in entries);info['membri']=len(entries)
        elif e=='.tar':
            with tarfile.open(p) as z:entries=[{'cale':x.name,'octeti':x.size} for x in z.getmembers()]
            text='\n'.join(f"{x['octeti']}\t{x['cale']}" for x in entries);info['membri']=len(entries)
        elif SEVEN.exists():
            r=subprocess.run([str(SEVEN),'l','-slt',str(p)],capture_output=True,timeout=90)
            text=decode(r.stdout);info['cod_listare']=r.returncode
        status='conținut arhivă listat; documentele interne nu sunt citite'
    elif e in {'.jpg','.jpeg','.png','.webp'}:
        with Image.open(p) as im:info.update(latime=im.width,inaltime=im.height,format=im.format)
        text=json.dumps(info,ensure_ascii=False);status='metadate imagine; interpretare vizuală necesară'
    elif e in {'.dwg','.dxf','.shx'}:
        text='Fișier CAD sau font CAD. Necesită aplicație CAD pentru interpretarea geometriei.';status='inventariat; interpretare CAD necesară'
    elif e in {'.mp4','.mov'}:
        text='Înregistrare video. Conținutul nu a fost vizionat/transcris.';status='inventariat; vizionare necesară'
    elif e=='.xls':
        import xlrd
        wb=xlrd.open_workbook(str(p));a=[];info['foi']=wb.sheet_names()
        for ws in wb.sheets():
            a.append('FOAIE: '+ws.name)
            for ri in range(ws.nrows):
                vals=[f'C{ci+1}={v}' for ci,v in enumerate(ws.row_values(ri)) if v!='']
                if vals:a.append(f'Rând {ri+1}: '+' | '.join(vals))
        text='\n'.join(a);status='celule XLS extrase'
    elif e=='.doc':
        import win32com.client
        app=win32com.client.DispatchEx('Word.Application');app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3
        try:
            doc=app.Documents.Open(str(p),ConfirmConversions=False,ReadOnly=True,AddToRecentFiles=False,Visible=False)
            text=doc.Content.Text;doc.Close(False);status='text DOC extras; macrocomenzi dezactivate'
        finally:app.Quit(False)
    else:text='Fișier tehnic sau format fără extractor disponibil.';status='inventariat; extragere indisponibilă'
    return text,status,info

RULES=[
 ('ciorna|draft|entwurf','Ciornă sau proiect de document; nu dovedește trimiterea ori acceptarea'),
 ('email-text_anfrage|email-text_angebotsanfrage|angebotsanfrage|cerere oferta','Cerere de ofertă pregătită sau salvată; trimiterea se verifică separat'),
 ('fragenkatalog|fragekatalog','Chestionar de clarificări tehnice și contractuale; verifică data și dacă este completat'),
 ('ang 891|bauwerksbuch','Ofertă sau documentație pentru cartea tehnică a clădirii'),
 ('ang 855|tragwerks|statik|static','Proiectare structurală, calcule sau ofertă de statică'),
 ('ang 839|prüfingenieur|pruefingenieur','Verificarea lucrărilor prin Prüfingenieur'),
 ('ang 840|baukg|sige|koordination','Coordonarea securității proiectării și șantierului'),
 ('fragenkatalog|fragekatalog','Chestionar de clarificări tehnice și contractuale'),
 ('werkvertrag|vertrag|contract','Contract, proiect de contract sau condiții contractuale'),
 ('angebot|anbot|ofert|preisindikation|honorar','Ofertă comercială, cerere de ofertă sau comparație de preț'),
 ('rechnung|invoice|factur|zahlungs','Factură sau document de plată'),
 ('email|e-mail|raspuns|răspuns|thread','Corespondență sau rezumat al comunicărilor'),
 ('befugnis|gewerbe|gisa|nachweis|firmenbuch|prüfung','Dovadă de autorizare, calificare sau înregistrare'),
 ('besprechung|intalnir|întâlnir|program','Pregătirea și programarea întâlnirilor'),
 ('plan|grundriss|schnitt|ansicht','Planșă, planificare sau documentație de proiect'),
 ('versicherung|asigur','Asigurare sau condiții de acoperire'),
 ('grundbuch|cadastr|kauf|eigentum','Proprietate, achiziție sau cadastru'),
 ('auszug|umsatz|extrase|banci','Extras sau document bancar'),
 ('supraf|fläche|flache','Suprafețe și măsurători'),
]
def describe(rel,text,status):
    n=Path(rel).name
    kind=next((v for k,v in RULES if re.search(k,n,re.I)),None)
    if not kind:
        e=Path(rel).suffix.lower()
        kind={'temporar':'Fișier temporar Office'}.get(status)
        if not kind:kind='Fotografie/imagine' if e in {'.jpg','.jpeg','.png','.webp'} else 'Arhivă de documente' if e in {'.zip','.rar','.tar','.7z'} else 'Înregistrare video' if e in {'.mp4','.mov'} else 'Desen sau resursă CAD' if e in {'.dwg','.dxf','.shx'} else 'Document de proiect'
    clean=re.sub(r'--- PAGINA \d+ ---','',text)
    excerpt=' '.join(clean.split())[:850] if ('extras' in status or 'citit' in status) else ''
    return kind+' — '+n,excerpt

def main():
    files=sorted((p for p in ROOT.rglob('*') if p.is_file() and BASE not in p.parents and not any(x in {'.git','.codex','.agents'} for x in p.relative_to(ROOT).parts)),key=lambda p:str(p).lower())
    rows=[]; hashes={}; errors=[]
    verified={}
    situation=ROOT/'04. Firme + Executie'/'Actualizare oferte 2026.09.30'/'situatie_structurata.json'
    if situation.exists():
        for item in json.loads(situation.read_text(encoding='utf-8'))['inregistrari']:
            for source in item['surse']:verified.setdefault(source,[]).append(item['firma']+' / '+item['serviciu']+': '+item['ultimul_raspuns'])
    for i,p in enumerate(files):
        rel=p.relative_to(ROOT).as_posix();st=p.stat();old=OLD.get(rel,{})
        digest=old.get('sha256') if old.get('octeti')==st.st_size and old.get('modificat_ns')==st.st_mtime_ns else sha(p)
        cp=CACHE/(digest+'.txt');meta=BASE/'texte'/(digest+'.json')
        if cp.exists() and meta.exists() and p.suffix:
            inf=json.loads(meta.read_text(encoding='utf-8'));t=cp.read_text(encoding='utf-8');status=inf.pop('status')
            if (status=='eroare extragere' and p.suffix.lower() in {'.xlsx','.tar','.doc'}) or (status=='inventariat; extragere indisponibilă' and p.suffix.lower() in {'.xls','.doc'}) or (status.startswith('conținut arhivă') and p.suffix.lower()=='.doc'):
                try:
                    t,status,inf=extract(p)
                    cp.write_text(t,encoding='utf-8');meta.write_text(json.dumps(dict(inf,status=status),ensure_ascii=False,indent=2),encoding='utf-8')
                except Exception as exc:t=f'{type(exc).__name__}: {exc}';status='eroare extragere';inf={}
        else:
            try:t,status,inf=extract(p)
            except Exception as exc:t=f'{type(exc).__name__}: {exc}';status='eroare extragere';inf={};errors.append({'cale':rel,'eroare':t})
            cp.write_text(t,encoding='utf-8');meta.write_text(json.dumps(dict(inf,status=status),ensure_ascii=False,indent=2),encoding='utf-8')
        if status=='eroare extragere' and not any(x['cale']==rel for x in errors):errors.append({'cale':rel,'eroare':t})
        desc,excerpt=describe(rel,t,status)
        r={'cale':rel,'folder':p.parent.relative_to(ROOT).as_posix(),'extensie':p.suffix.lower(),'octeti':st.st_size,'modificat':datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec='seconds'),'modificat_ns':st.st_mtime_ns,'sha256':digest,'duplicat_al':hashes.get(digest,''),'ce_reprezinta':desc,'baza_descrierii':'tip dedus din nume; extras de conținut separat','extras_continut':excerpt,'stare_citire':status,'text_cache':cp.relative_to(ROOT).as_posix(),'detalii':inf}
        r['rezumat_verificat']='\n'.join(verified.get(rel,[]))
        r['ocr_cache']=''
        op=BASE/'ocr'/(digest+'.txt')
        if op.exists():r['ocr_cache']=op.relative_to(ROOT).as_posix()
        hashes.setdefault(digest,rel);rows.append(r)
        if (i+1)%100==0:print(f'{i+1}/{len(files)} inventariate',flush=True)
    (BASE/'inventar.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
    fields=[k for k in rows[0] if k!='detalii']
    with (BASE/'inventar.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fields,extrasaction='ignore');w.writeheader();w.writerows(rows)
    folders=sorted([p for p in ROOT.rglob('*') if p.is_dir() and BASE not in p.parents and p!=BASE and not any(x in {'.git','.codex','.agents'} for x in p.relative_to(ROOT).parts)],key=str)
    dirs=[]
    roles={'00.Claude':'Analize de lucru, pași ai proiectului și rolurile echipei','00.Proiect':'Dosarul tehnic și administrativ structurat pe fazele proiectului','01. Proprietate + Acte':'Proprietate, societate, cadastru, achiziție și documente juridice','02. Autorizatie + Planse oficiale':'Autorizații și planșe oficiale ale clădirii','03. Proiectare':'Arhitectură, structură, studii, cantități și versiuni de proiect','04. Firme + Executie':'Firme contactate, cereri, oferte, clarificări și pregătirea execuției','05. Asigurari':'Asigurările clădirii și documentele aferente','06. Utilitati':'Contracte, contoare, facturi și corespondență pentru utilități','07. Poze + Video':'Documentarea fotografică și video a imobilului','08. Corespondenta':'Corespondență și fișiere recuperate; include un subfolder cu alte proiecte','09. Arhiva ZIP-uri mari':'Pachete originale de documentație primite prin transfer','10. Banci + Extrase de cont':'Extrase, facturi și reconcilierea plăților','Inbox':'Documente în așteptarea clasificării definitive','Website':'Fișierele site-ului local și resursele sale'}
    for p in folders:
        rel=p.relative_to(ROOT).as_posix();sub=[r for r in rows if r['cale'].startswith(rel+'/')];direct=[r for r in sub if r['folder']==rel]
        parts=p.relative_to(ROOT).parts
        role=roles.get(parts[0],'Planșe și documentație indexată')+'; subiect: '+' / '.join(parts[1:] or parts)
        kinds=collections.Counter(r['extensie'] or 'fără extensie' for r in sub)
        role+='; conține '+', '.join(f'{v} {k}' for k,v in kinds.most_common(4)) if kinds else '; folder fără fișiere în prezent'
        dirs.append({'cale':rel,'fisiere_directe':len(direct),'fisiere_recursive':len(sub),'rol':role,'baza':'rol dedus din structura și denumirea folderului; tipuri de fișiere inventariate'})
    (BASE/'foldere.json').write_text(json.dumps(dirs,ensure_ascii=False,indent=2),encoding='utf-8')
    state=collections.Counter(r['stare_citire'] for r in rows)
    stats={'generat':STAMP,'radacina':str(ROOT),'fisiere':len(rows),'foldere':len(dirs),'continuturi_unice':len(hashes),'copii_identice':len(rows)-len(hashes),'stari':dict(state),'erori':errors}
    (BASE/'statistici.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    wb=Workbook();ws=wb.active;ws.title='Fișiere'
    cols=['cale','ce_reprezinta','extras_continut','stare_citire','duplicat_al','octeti','modificat','sha256','text_cache','ocr_cache','rezumat_verificat']
    ws.append(cols)
    for r in rows:
        ws.append([re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',str(r[k]))[:32700] for k in cols])
        for c in ws[ws.max_row]:c.data_type='s'
        ws.cell(ws.max_row,1).hyperlink=(ROOT/r['cale']).as_uri();ws.cell(ws.max_row,9).hyperlink=(ROOT/r['text_cache']).as_uri()
        if r['ocr_cache']:ws.cell(ws.max_row,10).hyperlink=(ROOT/r['ocr_cache']).as_uri()
    ds=wb.create_sheet('Foldere');ds.append(['Cale','Fișiere directe','Fișiere recursive','Rol conform structurii'])
    for r in dirs:ds.append([r['cale'],r['fisiere_directe'],r['fisiere_recursive'],r['rol']])
    ss=wb.create_sheet('Acoperire');ss.append(['Indicator','Valoare'])
    for k,v in stats.items():
        if k not in {'stari','erori'}:ss.append([k,str(v)])
    for k,v in state.items():ss.append([k,v])
    for sh in wb:
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24486B')
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=55 if col[0].column<=3 else 30
        for row in sh.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(BASE/'Folder map.xlsx')
    # A navigable index includes every folder, including empty directories.
    byfolder=collections.defaultdict(list)
    for r in rows:byfolder[r['folder']].append(r)
    lines=['# Harta documentelor Schallergasse 35','',f'Generat: {STAMP}. {len(rows)} fișiere, {len(dirs)} foldere, {len(hashes)} conținuturi distincte.','', 'Descrierile de tip sunt deduse din nume. Extrasele sunt text din documente. Starea de citire arată limitele fiecărui fișier.','']
    from urllib.parse import quote
    for folder in ['.']+[d['cale'] for d in dirs]:
        lines+=['## '+folder,'']
        if folder!='.':lines.extend([next(d['rol'] for d in dirs if d['cale']==folder),''])
        for r in byfolder[folder]:
            link='../'+quote(r['cale'],safe='/')
            lines.append(f"- [{Path(r['cale']).name}]({link}) — {r['ce_reprezinta']}. Stare: {r['stare_citire']}."+(f" Copie identică: `{r['duplicat_al']}`." if r['duplicat_al'] else ''))
            if r['extras_continut']:lines.append('  Extras: '+r['extras_continut'].replace('\n',' '))
            if r['ocr_cache']:lines.append('  OCR automat disponibil: `'+r['ocr_cache']+'`. Necesită verificare înaintea folosirii sumelor sau termenelor.')
            if r['rezumat_verificat']:lines.append('  Situație documentată: '+r['rezumat_verificat'].replace('\n',' '))
        if not byfolder[folder]:lines.append('Fără fișiere directe; consultați subfolderele.')
        lines.append('')
    (BASE/'HARTA_COMPLETA.md').write_text('\n'.join(lines),encoding='utf-8')
    db=sqlite3.connect(BASE/'cautare.sqlite')
    db.execute('CREATE VIRTUAL TABLE IF NOT EXISTS documente USING fts5(cale, descriere, text, sha256 UNINDEXED)')
    db.execute('DELETE FROM documente')
    for r in rows:
        tx=(ROOT/r['text_cache']).read_text(encoding='utf-8')
        if r['ocr_cache']:tx+='\n\nOCR AUTOMAT NEVERIFICAT\n'+(ROOT/r['ocr_cache']).read_text(encoding='utf-8')
        db.execute('INSERT INTO documente VALUES (?,?,?,?)',(r['cale'],r['ce_reprezinta']+'\n'+r['rezumat_verificat'],tx,r['sha256']))
    db.commit();db.close()
    print(json.dumps(stats,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
