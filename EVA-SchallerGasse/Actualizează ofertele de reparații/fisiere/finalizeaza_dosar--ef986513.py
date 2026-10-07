from pathlib import Path
import json, shutil, hashlib, re
from openpyxl import load_workbook
from docx import Document
import genereaza_documente as g
W=Path(__file__).resolve().parent;OUT=W.parent;ROOT=OUT.parents[3]
eb=g.parse((W/'EMAIL_DE.md').read_text(encoding='utf-8'))
g.word(eb,OUT/'02 E-Mail Vertragsvorschlag BAU-WERTE.docx','E-Mail Vertragsvorschlag BAU-WERTE')
(OUT/'02 E-Mail Vertragsvorschlag BAU-WERTE.txt').write_text('\n\n'.join(g.plain(v) for t,v in eb if isinstance(v,str)),encoding='utf-8-sig')
ev=json.loads((OUT/'04 Eva-Mail BAU-WERTE 2026-09-25.json').read_text(encoding='utf-8'))
(OUT/'04 Eva-Mail BAU-WERTE 2026-09-25.txt').write_text('Mesaj primit, export text Eva-Mail, verificat la 30.09.2026.\nID: '+ev['id']+'\nData UTC: '+ev['received_at']+'\nDe la: '+ev['from_address']+'\nCatre: '+', '.join(ev['to'])+'\nSubiect: '+ev['subject']+'\n\n'+ev['text'],encoding='utf-8-sig')
note="""DOSAR BAU-WERTE — PROPUNERE DE CONTRACT — 30.09.2026
Documentele sunt în germană. Nu sunt semnate și emailul nu a fost trimis.

01 Contract: Word editabil și PDF de 19 pagini. Textul contractual și anexele 1/3 sunt editabile în Word; cele trei pagini ale ofertei originale sunt facsimile. PDF-ul include paginile originale 15–17, identice vizual cu sursa.
02 Email: text și Word, pregătite pentru baumeister@bau-werte.biz.
03 Oferta originală: copie integrală, identică prin SHA-256.
04 Eva-Mail: mesajul primit la 25.09, export JSON cu metadate și text lizibil.
05 Audit: raportul verificărilor juridice, tehnice și ale documentelor.
_lucru: sursele de redactare, matricea cerințelor TOMS, sursele legale, auditul detaliat și verificările tehnice.

CONDIȚII PROPUSE
8.250 EUR net: 750 planificare + 7.500 coordonare șantier; 9.900 EUR cu TVA 20%.
12 luni + până la 3 luni fără onorariu suplimentar pentru același proiect. Vizită personală cel puțin săptămânal în perioadele active și vizite suplimentare necesare riscurilor; circa 50 nu reprezintă plafon.
SiGe-Plan, notificare când legea o impune, regulament, documentația pentru lucrări ulterioare și actualizări incluse. Cel puțin două ședințe de planificare. Rapoarte în două zile lucrătoare, pericole comunicate imediat.
De la luna 16: continuare aprobată, propusă la 625 EUR net/lună proporțional. Vizită opțională în afara sarcinilor incluse: 200 EUR net; alte suplimente aprobate: 135 EUR net/oră, fără dublă facturare.
Asigurare propusă 2 milioane EUR/caz, de dovedit și acceptat. Valabilitate nouă propusă până la 31.12.2026.
Aceste completări necesită acceptarea expresă BAU-WERTE. Nu sunt drepturi deja convenite.

ULTIMUL RĂSPUNS VERIFICAT
Emailul Lechner din 25.09.2026, 11:04 UTC, confirmă întâlnirea și recomandă Thomas Schauer / SCHAUERLEUTE GmbH. Nu confirmă prelungirea ofertei de 8.250 EUR sau acceptarea condițiilor suplimentare. Corectează situația locală anterioară, în care întâlnirea era doar programată.

ÎNAINTE DE SEMNARE
Se completează datele de registru/UID, consimțământul coordonatorului, înlocuitorul, calificările și polița; se confirmă cronologia șantierului și termenele legale. Agentul juridic, specialistul tehnic și auditorul sunt agenți AI. Auditul nu este aviz semnat de avocat; recomandăm verificarea finală de un avocat austriac înainte de semnare.
PDF-ul a fost verificat vizual. Word-ul a fost verificat structural și pentru integritatea textului; paginarea sa în Microsoft Word nu a fost randată în această sesiune.
"""
(OUT/'00 Citeste-ma - situatie si continut dosar.txt').write_text(note,encoding='utf-8-sig')
DEST=ROOT/'04. Firme + Executie'/'Actualizare oferte 2026.09.30'
hist=ROOT/'folder map'/'istoric'/'2026-09-30 inainte completare BAU-WERTE';hist.mkdir(parents=True,exist_ok=True)
def back(p):
    q=hist/p.name
    if p.exists() and not q.exists():shutil.copy2(p,q)
source_rel=(OUT/'04 Eva-Mail BAU-WERTE 2026-09-25.json').relative_to(ROOT).as_posix()
reply='Email primit 25.09 confirmă întâlnirea și recomandă Thomas Schauer / SCHAUERLEUTE GmbH. Nu confirmă reînnoirea ofertei. Contract nou pregătit la 8.250 EUR net, 12+3 luni propuse, neacceptat.'
nextstep='Transmite propunerea Word/PDF și cere acceptarea expresă a prețului, valabilității până la 31.12 și condițiilor incluse; completează dovezile din Anexa 3 înainte de semnare. Emailul este pregătit, netrimis.'
cover='Etapa inițială: surse locale. Completare BAU-WERTE prin Eva-Mail la 30.09.2026; ceilalți ofertanți nu au fost reverificați în această completare.'
jp=DEST/'situatie_structurata.json';back(jp);data=json.loads(jp.read_text(encoding='utf-8'));data['acoperire']=cover
for r in data['inregistrari']:
    if r['firma']=='BAU-WERTE':
        r.update(data_ultimei_dovezi='2026-09-25',status='Întâlnire confirmată; propunere de contract nouă, neacceptată',ultimul_raspuns=reply,pas_urmator=nextstep,nivel_dovada='Email original Eva-Mail exportat; ofertă originală verificată',verificare_mail='Eva-Mail verificat 30.09.2026 pentru BAU-WERTE; mesaj ID '+ev['id'])
        if source_rel not in r['surse']:r['surse'].append(source_rel)
jp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
xp=ROOT/'04. Firme + Executie'/'Registru Comunicari si Oferte - Schallergasse 35 - 2026-09-30.xlsx';back(xp);wb=load_workbook(xp);ws=wb['Situație curentă'];headers={c.value:c.column for c in ws[1]}
for row in ws.iter_rows(min_row=2):
    if row[headers['firma']-1].value=='BAU-WERTE':
        rr=next(r for r in data['inregistrari'] if r['firma']=='BAU-WERTE')
        for k,col in headers.items():ws.cell(row[0].row,col,'\n'.join(rr[k]) if isinstance(rr[k],list) else rr[k])
        ws.cell(row[0].row,headers['surse']).hyperlink=(OUT/'04 Eva-Mail BAU-WERTE 2026-09-25.json').as_uri()
for sh in ['Acoperire','Valabilități']:
    ws=wb[sh]
    if sh=='Acoperire':
        for row in ws:
            if row[0].value=='Email Eva-Mail':row[1].value=cover
    else:
        for row in ws.iter_rows(min_row=2):
            if any('BAU-WERTE' in str(c.value) for c in row):
                row[3].value='25.09 întâlnire confirmată; fără prelungire expresă'
                row[4].value='Contract nou propus 30.09; acceptare până la 31.12 propusă, nu confirmată'
wb.save(xp)
addendum='COMPLETARE BAU-WERTE, verificare Eva-Mail la 30.09.2026: '+reply+' '+nextstep+' Restul raportului păstrează situația verificată anterior, pe surse locale.'
mp=DEST/'Situatie oferte 2026-09-30.md';back(mp);s=mp.read_text(encoding='utf-8')
if '## Completare BAU-WERTE din Eva-Mail' not in s:s='## Completare BAU-WERTE din Eva-Mail\n\n'+addendum+'\n\n---\n\n'+s
mp.write_text(s,encoding='utf-8')
dp=DEST/'Situatie oferte 2026-09-30.docx';back(dp);dd=Document(dp)
if not any(p.text.startswith('COMPLETARE BAU-WERTE') for p in dd.paragraphs):dd.paragraphs[0].insert_paragraph_before(addendum)
dd.save(dp)
fp=DEST/'BAU-WERTE'/'Situatie si surse.md';back(fp);s=fp.read_text(encoding='utf-8')
if '## Completare Eva-Mail' not in s:s='## Completare Eva-Mail\n\n'+addendum+'\n\n---\n\n'+s
fp.write_text(s,encoding='utf-8')
legal=W/'SURSE_JURIDICE.md';s=legal.read_text(encoding='utf-8');s=s.replace('ori un rezultat al întâlnirilor programate.','a condițiilor financiare. Ulterior redactării, emailul Eva-Mail din 25.09.2026 (ID f3943f2d-516e-4330-8b53-3fd9421d07e1), salvat în dosar, a confirmat întâlnirea și recomandarea SCHAUERLEUTE, fără acceptarea condițiilor contractuale.');legal.write_text(s,encoding='utf-8')
print('Email regenerat. Sursa nouă salvată. Registrul BAU-WERTE și rezumatele completate; istoricul păstrat în folder map/istoric.')
