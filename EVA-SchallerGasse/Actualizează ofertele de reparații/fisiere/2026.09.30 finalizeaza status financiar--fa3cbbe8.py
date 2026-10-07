from pathlib import Path
import json,shutil
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
R=Path(__file__).resolve().parent.parent;D='2026.09.30';P=R/'folder map'/f'{D} statusuri verificate.json'
m=json.loads(P.read_text(encoding='utf-8'))
audit='10. Banci + Extrase de cont/2026.09.30 Audit facturi si plati/2026.09.30 Registru verificat.json'
def update(k,s,a):
 old=m.get(k,{})
 m[k]={'status':s,'urmatorul_pas':a,'surse':list(dict.fromkeys(old.get('surse',[])+[audit]))}
update('Attensam','2026.09.30: factura 34642 deszapezire, 653,12 EUR, stinsa prin plata firmei 640,06 EUR din 2026.09.28 si Skonto 13,06 EUR. Factura 34648 deratizare, 156,53 EUR, debit personal din 2026.09.09; furnizorul ceruse dovada la 2026.09.16. Versiunile vechi 6253/25777 nu se aduna. Identificata si factura 14983/1001371 din 2026.07.20, 156,53 EUR, perioada ambigua Juli–Juni 2026.','Confirmarea alocarii 34648 si inchiderii versiunilor vechi; clarificarea perioadei/stornarii 14983. Nu se repeta platile deja debitate.')
update('CERHA HEMPEL','2026.09.30: HN 25/31404, 6.147,70 EUR, achitata 2026.02.16; HN 26/30186, 24.116 EUR, achitata 2026.03.31 (numar cu un 1 suplimentar in extras). HN 26/30363, 2.640,22 EUR, fara debit identificat; scadenta 2026.05.05, 14 zile acordate de Capra. Taxa registru 493 EUR este deja inclusa in prima factura. Capra descrie somatia Stadt Wien drept 8 EUR; PDF-ul arata 7,59 EUR, fara a doua taxa identificata.','Confirmarea soldului 26/30363 dupa 2026.09.28; obtinerea decontului final Treuhand pentru 75.900 EUR taxe si 1.650.000 EUR pret. Predarea dosarului fostei administratii ramane de urmarit.')
update('ARTUS','2026.09.30: 1261556 = 360 EUR factura + 18,08 accesorii, achitat 378,08 EUR la 2026.06.05; 1263320 = 240 EUR achitat personal 2026.08.05; 1264971 = 144 EUR fara debit identificat. Impozit 186 EUR achitat 2026.06.11, valuta 2026.06.10, beneficiar AT62. PDF fiscal local criptat; 93 EUR Q3 si IBAN AT36 din Excelul vechi necesita reconfirmare. Decizia anuala este pentru 2026, 375 EUR.','Fisa de cont ARTUS si extras FinanzOnline actual; verificarea documentului criptat, soldului si IBAN-ului. Nu se dubleaza impozitul anual cu avansurile.')
update('MA6 - Taxe','2026.09.30: obligatia Q2 93,23 EUR este achitata personal la 2026.07.27, referinta 889970805056; cele doua randuri vechi descriu aceeasi obligatie. Aviz Q3 recuperat, 480960377428 din 2026.07.10: 186,46 EUR cumulativ, din care 93,23 pentru Q2 deja achitat; componenta Q3 93,23 EUR fara debit identificat. Cererea privind salubritatea se urmareste separat cu MA48.','Confirmarea alocarii platii Q2 si a soldului Q3 de 93,23 EUR; nu se achita din nou intregul aviz cumulativ.')
update('Wiener Wasser - MA31','2026.09.30: factura 310910909051 din 2026.09.07, 15,18 EUR, scadenta 2026.10.15, cont NOU 123482077 / WA120031905/06. Ordin anuntat in capturile 2026.09.30, debit neconfirmat. Somatia veche 7,59 EUR este pe cont 123003016 / contract /05; avizul vechi 15,18 si somatia nu se cumuleaza.','Confirmarea debitului de 15,18 EUR si transferului/inchiderii contului vechi. Urmarirea schimbarii titularului si soldurilor fara dublare.')
update('Sturm Energie','2026.09.30: gaz Top15, factura 29700, 115,36 EUR, achitat personal 2026.08.06. Alte opt facturi gaz insumeaza 883,98 EUR, fara debit identificat; perioada originala 2025.06.11–2026.04.08, vechiul UID. Noua facturi de curent din 2026.07.15, transmise de Capra la 2026.07.21: suma unica 1.817,73 EUR fara debit, perioade din 2023–2026. 41415 include 40809; 41422 aplica deja credit 122,41 EUR. Credit curent 29707 de 846,92 EUR, rambursare anuntata spre vechiul cont AT21...0800; nu apare incasare in AT22/AT13.','Confirmarea repartizarii intre proprietari, corectarea datelor si identificarea beneficiarului creditului. Fara compensare automata gaz–curent.')
old=m['Wien Energie']['status'];m['Wien Energie']['status']=old+' Reconciliere 2026.09.30: factura 5147075005, 109,02 EUR, fara debit identificat; perioada vechiului proprietar si repartizarea sunt de clarificat.' if '5147075005' not in old else old
m['Wien Energie']['surse']=list(dict.fromkeys(m['Wien Energie']['surse']+[audit]))
old=m['Donau Versicherung']['status'];m['Donau Versicherung']['status']=old+' Reconciliere: 3.734,34 EUR = 3.183,62 principal + 501,58 Inkasso + 30 cost creditor + 19,14 dobanda. Niciun debit identificat. Q4 1.591,81 EUR ramane estimare; PDF-ul etichetat Q4 este somatia din iunie.' if 'cost creditor' not in old else old
m['Donau Versicherung']['surse']=list(dict.fromkeys(m['Donau Versicherung']['surse']+[audit]))
update('MA29 - Geologie','2026.07.09: factura SHOP2026000439 pentru trei profile, 21,60 EUR, confirma plata electronica. Niciun debit identificat in cele doua extrase austriece; contul efectiv ramane necunoscut.','Pastrarea profilelor si identificarea contului platii pentru contabilizare, fara a plati din nou.')
P.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')

# Situatie noua de oferte, bazata pe mesajele Eva si sintezele verificate.
names=['BAU-WERTE','TOMS','Themis','Otis','Weigl','Schmitt + Sohn','Pech','Paknehad','LEBE Bau','OBENAUF','Novotny','Sedlak','SCHAUERLEUTE','STOLEX','HAZET','KALCON','Sanibau','Neid','Avunduk','Heissenberger','KPPK','KONE','TK Elevator','Markus Berger','Mattes Projektmanagement','LVR','Frigo','SMT Immobilien']
wb=Workbook();ws=wb.active;ws.title='Ultimele raspunsuri';ws.append(['Data actualizarii','Partener','Situatie verificata','Urmatorul pas','Surse'])
doc=Document();doc.add_heading(D+' | Ultimele raspunsuri ale ofertantilor',0);doc.add_paragraph('Verificat in mesajele disponibile Eva-Mail. Preturile si valabilitatea sunt cele din sursa indicata. Contractele pregatite sunt propuneri; ciornele nu reprezinta emailuri trimise.')
for k in names:
 if k not in m:continue
 v=m[k];ws.append([D,k,v['status'],v['urmatorul_pas'],'; '.join(v['surse'])]);doc.add_heading(k,1);doc.add_paragraph(v['status']);doc.add_paragraph('Urmatorul pas: '+v['urmatorul_pas']);doc.add_paragraph('Surse: '+'; '.join(v['surse']))
ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24486B')
for row in ws.iter_rows(min_row=2):
 ws.row_dimensions[row[0].row].height=110
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
for col,width in [('A',15),('B',25),('C',95),('D',65),('E',45)]:ws.column_dimensions[col].width=width
wb.save(R/'04. Firme + Executie'/f'{D} Ultimele raspunsuri ofertanti.xlsx');doc.save(R/'04. Firme + Executie'/f'{D} Ultimele raspunsuri ofertanti.docx')

ar=json.loads((R/'08. Corespondenta/2026.09.30 Arhiva Eva-Mail/2026.09.30 Registru atasamente.json').read_text(encoding='utf-8'))
themis=next((a for a in ar if a['id'].startswith('368aac5b')),None)
if themis and themis.get('cale'):
 dirs=[p for p in (R/'04. Firme + Executie').rglob('*') if p.is_dir() and 'themis' in p.name.lower()]
 dest=dirs[0] if dirs else R/'04. Firme + Executie/Themis'
 dest.mkdir(exist_ok=True);shutil.copy2(R/themis['cale'],dest/'2026.09.29 Themis oferta BB-2609-1-1 Bauwerksbuch.pdf')

notice='2026.09.30 | Situatia curenta verificata\n\nExcelurile mai vechi sunt pastrate ca istoric. Folositi: 10. Banci + Extrase de cont/2026.09.30 Reconciliere completa facturi si plati.xlsx si 08. Corespondenta/2026.09.30 De platit - Schallergasse 35 - verificat.xlsx.\n\nDeszapezirea Attensam este achitata cu Skonto; Q2 MA6 nu se dubleaza. PDF-ul Donau din iunie nu este aviz Q4. Pentru sume neclare vedeti foaia Clarificari inainte de plata si raportul din dosarul Audit facturi si plati.\n'
for folder in [R/'08. Corespondenta',R/'10. Banci + Extrase de cont']:
 (folder/f'{D} Citeste situatia actualizata.txt').write_text(notice,encoding='utf-8-sig')
print(json.dumps({'statusuri':len(m),'ofertanti':len(names),'themis_copiat':bool(themis)},ensure_ascii=False))
