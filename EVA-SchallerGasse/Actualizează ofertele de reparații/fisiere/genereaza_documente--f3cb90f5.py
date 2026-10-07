"""Generează Word și PDF din aceeași sursă contractuală; anexează oferta originală."""
from pathlib import Path
import re, json, hashlib, shutil, html
from copy import deepcopy
import fitz
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

WORK=Path(__file__).resolve().parent
OUT=WORK.parent
ROOT=OUT.parents[3]
LAYOUT=WORK/'layout';LAYOUT.mkdir(exist_ok=True)
SOURCE=WORK/'CONTRACT_DE.md'
OFFER=OUT.parent/'Honorarangebot 2026-08-11-AG26-0094.pdf'
NAVY=colors.HexColor('#20384E'); GREY=colors.HexColor('#627282'); LIGHT=colors.HexColor('#EBF0F4')
fonts=Path('C:/Windows/Fonts')
for n,f in [('Calibri','calibri.ttf'),('Calibri-Bold','calibrib.ttf'),('Calibri-Italic','calibrii.ttf'),('Calibri-BoldItalic','calibriz.ttf')]:pdfmetrics.registerFont(TTFont(n,str(fonts/f)))
pdfmetrics.registerFontFamily('Calibri',normal='Calibri',bold='Calibri-Bold',italic='Calibri-Italic',boldItalic='Calibri-BoldItalic')

def parse(text):
    out=[];lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line in {'---','***'}:i+=1;continue
        if line=='[PAGEBREAK]':out.append(('pagebreak',''));i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',c) for c in row):rows.append(row)
                i+=1
            out.append(('table',rows));continue
        heading=re.match(r'^(#{1,4})\s+(.*)',line)
        if heading:out.append(('h'+str(len(heading[1])),heading[2]));i+=1;continue
        if line.startswith('- '):out.append(('bullet',line[2:]));i+=1;continue
        if re.match(r'^\d+\.\s',line):out.append(('number',line));i+=1;continue
        if line.startswith('> '):out.append(('note',line[2:]));i+=1;continue
        para=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,4} |\||- |>|\[PAGEBREAK\]|\d+\.\s)',lines[i].strip()):
            para.append(lines[i].strip());i+=1
        out.append(('p',' '.join(para)))
    return out

def plain(s):return re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1',s).replace('**','').replace('`','')
def markup(s):
    s=html.escape(s).replace('`','')
    s=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',s)
    s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'<link href="\2" color="#24486B">\1</link>',s)
    return s

def field(p,name):
    r=OxmlElement('w:r');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),name);r.append(f);p._p.append(r)

def doc_header(doc,title):
    for s in doc.sections:
        s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
        s.header_distance=Cm(.85);s.footer_distance=Cm(.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.10
    for name,size in [('Title',23),('Heading 1',14),('Heading 2',11.5),('Heading 3',10.5)]:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('20384E');st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
    h=doc.sections[0].header.paragraphs[0];h.text='A&C WOHNART  /  BAU-WERTE';h.runs[0].font.size=Pt(8);h.runs[0].font.color.rgb=RGBColor.from_string('627282')
    f=doc.sections[0].footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('Schallergasse 35 · Entwurf 30.09.2026  |  Seite ');field(f,'PAGE');f.add_run(' / ');field(f,'NUMPAGES')
    for r in f.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('627282')
    settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.core_properties.title=title;doc.core_properties.subject='Planungs- und Baustellenkoordination · Schallergasse 35';doc.core_properties.author='A&C Wohnart Immobilien GmbH';doc.core_properties.keywords='BauKG, BAU-WERTE, Vertragsentwurf'

def add_runs(p,s):
    s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1 (\2)',s).replace('`','')
    for i,x in enumerate(re.split(r'\*\*(.*?)\*\*',s)):
        r=p.add_run(x)
        if i%2:r.bold=True

def word(blocks,destination,title,annex=False):
    doc=Document();doc_header(doc,title)
    for typ,value in blocks:
        if annex and typ=='h2' and value.startswith('Anlage 3'):
            sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
            sec.top_margin=Cm(.7);sec.bottom_margin=Cm(.7);sec.left_margin=Cm(.7);sec.right_margin=Cm(.7)
            sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
            for p in sec.header.paragraphs:p.text=''
            for p in sec.footer.paragraphs:p.text=''
            with fitz.open(OFFER) as offer:
                for i,page in enumerate(offer):
                    if i:doc.add_page_break()
                    img=LAYOUT/f'originalangebot_{i+1}.png';page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(img)
                    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                    p.add_run().add_picture(str(img),width=Cm(19.3))
            sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
            sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
            sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
            for dest,src in [(sec.header,doc.sections[0].header),(sec.footer,doc.sections[0].footer)]:
                for child in list(dest._element):dest._element.remove(child)
                for child in src._element:dest._element.append(deepcopy(child))
        elif annex and typ=='h2' and value.startswith('Anlage '):doc.add_page_break()
        if typ=='pagebreak':doc.add_page_break();continue
        if typ=='table':
            count=max(map(len,value));t=doc.add_table(rows=0,cols=count);t.style='Table Grid'
            for ri,row in enumerate(value):
                cells=t.add_row().cells
                for ci,txt in enumerate(row):
                    p=cells[ci].paragraphs[0];add_runs(p,txt)
                    for r in p.runs:r.font.size=Pt(9)
                    if ri==0:
                        for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor.from_string('FFFFFF')
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'20384E');cells[ci]._tc.get_or_add_tcPr().append(sh)
                trPr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            doc.add_paragraph();continue
        style={'h1':'Title','h2':'Heading 1','h3':'Heading 2','h4':'Heading 3','bullet':'List Bullet'}.get(typ)
        p=doc.add_paragraph(style=style);add_runs(p,value)
        if typ=='note':
            for r in p.runs:r.italic=True;r.font.color.rgb=RGBColor.from_string('627282')
        if typ=='number':p.paragraph_format.left_indent=Cm(.25)
        if 'Unterschrift' in value or 'Ort, Datum' in value:p.paragraph_format.keep_with_next=True
    if False:
        doc.add_page_break();doc.add_heading('Anlage 2 · Originalangebot BAU-WERTE',1)
        doc.add_paragraph('Angebot Nr. 26-0094 vom 11.08.2026 · vollständige Wiedergabe der drei Originalseiten. Der enthaltene historische Bestellschein begründet keine gesonderte Bestellung. Maßgeblich sind die Rangfolge und Abweichungen dieses Vertragsentwurfs.')
        sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.top_margin=Cm(.7);sec.bottom_margin=Cm(.7);sec.left_margin=Cm(.7);sec.right_margin=Cm(.7)
        sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
        for p in sec.header.paragraphs:p.text=''
        for p in sec.footer.paragraphs:p.text=''
        offer=fitz.open(OFFER)
        for i,page in enumerate(offer):
            if i:doc.add_page_break()
            img=LAYOUT/f'originalangebot_{i+1}.png';page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(img)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
            p.add_run().add_picture(str(img),width=Cm(19.5))
        offer.close()
    doc.save(destination)

STYLE={
    'p':ParagraphStyle('body',fontName='Calibri',fontSize=10.5,leading=13.6,spaceAfter=5.5,textColor=colors.HexColor('#17232E'),allowWidows=0,allowOrphans=0),
    'h1':ParagraphStyle('title',fontName='Calibri-Bold',fontSize=22,leading=26,spaceBefore=9,spaceAfter=13,textColor=NAVY,keepWithNext=True),
    'h2':ParagraphStyle('heading',fontName='Calibri-Bold',fontSize=13.3,leading=16,spaceBefore=12,spaceAfter=6,textColor=NAVY,keepWithNext=True),
    'h3':ParagraphStyle('subheading',fontName='Calibri-Bold',fontSize=11,leading=14,spaceBefore=8,spaceAfter=4,textColor=NAVY,keepWithNext=True),
    'h4':ParagraphStyle('subsub',fontName='Calibri-Bold',fontSize=10.5,leading=13,spaceBefore=7,spaceAfter=4,textColor=NAVY,keepWithNext=True),
    'bullet':ParagraphStyle('bullet',fontName='Calibri',fontSize=10.5,leading=13.6,leftIndent=11,firstLineIndent=-8,spaceAfter=4,allowWidows=0,allowOrphans=0),
    'number':ParagraphStyle('number',fontName='Calibri',fontSize=10.5,leading=13.6,leftIndent=8,spaceAfter=4,allowWidows=0,allowOrphans=0),
    'note':ParagraphStyle('note',fontName='Calibri-Italic',fontSize=9.4,leading=12,spaceAfter=8,textColor=GREY),
    'cell':ParagraphStyle('cell',fontName='Calibri',fontSize=9,leading=11.7,spaceAfter=0),
    'cellhead':ParagraphStyle('cellhead',fontName='Calibri-Bold',fontSize=9,leading=11.7,textColor=colors.white,spaceAfter=0),
}
def pdf(blocks,destination,title,annex=False):
    W,H=A4;story=[]
    for typ,value in blocks:
        if annex and typ=='h2' and value.startswith('Anlage '):story.append(PageBreak())
        if typ=='pagebreak':story.append(PageBreak());continue
        if typ=='table':
            n=max(map(len,value));widths=([.28,.72] if n==2 else [.22,.40,.38] if n==3 else [.10,.32,.33,.25] if n==4 else [1/n]*n)
            data=[[Paragraph(markup(c),STYLE['cellhead' if ri==0 else 'cell']) for c in row]+['']*(n-len(row)) for ri,row in enumerate(value)]
            t=Table(data,colWidths=[x*(W-112) for x in widths],repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),NAVY),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,LIGHT]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,0),(-1,0),.6,NAVY),('LINEBELOW',(0,1),(-1,-1),.25,colors.HexColor('#D3DCE3'))]))
            story.extend([t,Spacer(1,7)]);continue
        content=markup(value)
        if typ=='bullet':content='• '+content
        story.append(Paragraph(content,STYLE[typ]))
    if False:
        story.extend([PageBreak(),Paragraph('Anlage 2 · Originalangebot BAU-WERTE',STYLE['h2']),Paragraph('Angebot Nr. 26-0094 vom 11.08.2026',STYLE['h3']),Paragraph('Die folgenden drei Seiten geben das Originalangebot vollständig wieder. Die Rangfolge der Vertragsunterlagen und die ausdrücklich vereinbarten Änderungen des Hauptvertrags gehen vor. Der historische Bestellschein wird ausschließlich als Bestandteil des Originalangebots beigefügt und begründet keine gesonderte Bestellung.',STYLE['p'])])
    temp=LAYOUT/(destination.stem+'_text.pdf')
    def hf(canvas,doc):
        canvas.saveState();canvas.setFillColor(GREY);canvas.setFont('Calibri',8);canvas.drawString(56,H-30,'A&C WOHNART  /  BAU-WERTE');canvas.drawRightString(W-56,H-30,'SCHALLERGASSE 35 · 1120 WIEN');canvas.setStrokeColor(colors.HexColor('#C6D0D8'));canvas.setLineWidth(.4);canvas.line(56,H-37,W-56,H-37);canvas.restoreState()
    SimpleDocTemplate(str(temp),pagesize=A4,leftMargin=56,rightMargin=56,topMargin=52,bottomMargin=51,title=title,author='A&C Wohnart Immobilien GmbH').build(story,onFirstPage=hf,onLaterPages=hf)
    d=fitz.open(temp)
    contract_pages=len(d)
    insert_at=None
    if annex:
        insert_at=next(i for i,p in enumerate(d) if p.get_text().startswith('A&C WOHNART') and 'Anlage 3 – Bestellungs- und Nachweisblatt' in p.get_text())
        offer=fitz.open(OFFER);d.insert_pdf(offer,start_at=insert_at);offer.close()
    total=len(d)
    for i,page in enumerate(d):
        # Originalangebot bleibt unverändert; Fußzeile nur im neu verfassten Vertragskörper.
        if insert_at is None or not insert_at <= i < insert_at+3:
            page.insert_font(fontname='CalibriFooter',fontfile=str(fonts/'calibri.ttf'))
            page.insert_text((56,page.rect.height-25),'Vertragsentwurf · 30.09.2026',fontname='CalibriFooter',fontsize=8,color=(.39,.45,.51))
            page.insert_text((page.rect.width-125,page.rect.height-25),f'Seite {i+1} / {total}',fontname='CalibriFooter',fontsize=8,color=(.39,.45,.51))
    d.set_metadata({'title':title,'author':'A&C Wohnart Immobilien GmbH','subject':'Schallergasse 35 · Entwurf zur Abstimmung · 30.09.2026'})
    d.save(destination,garbage=4,deflate=True);d.close()
    return contract_pages,total

def main():
    blocks=parse(SOURCE.read_text(encoding='utf-8'))
    basename='01 Vertragsentwurf BauKG BAU-WERTE Schallergasse 35 2026-09-30'
    word(blocks,OUT/(basename+'.docx'),'Vertragsentwurf BauKG · Schallergasse 35',True)
    cp,total=pdf(blocks,OUT/(basename+'.pdf'),'Vertragsentwurf BauKG · Schallergasse 35',True)
    shutil.copy2(OFFER,OUT/'03 Anlage 2 Originalangebot BAU-WERTE 26-0094 2026-08-11.pdf')
    email=WORK/'EMAIL_DE.md'
    if email.exists():
        eb=parse(email.read_text(encoding='utf-8'))
        word(eb,OUT/'02 E-Mail Vertragsvorschlag BAU-WERTE.docx','E-Mail Vertragsvorschlag BAU-WERTE')
        (OUT/'02 E-Mail Vertragsvorschlag BAU-WERTE.txt').write_text('\n\n'.join(plain(v) for t,v in eb if isinstance(v,str)),encoding='utf-8-sig')
    manifest={'quelle_md':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'angebot_sha256':hashlib.sha256(OFFER.read_bytes()).hexdigest(),'vertrag_pdf_seiten':cp,'pdf_gesamtseiten':total,'anlagen_originalseiten':3,'stand':'2026-09-30','pdf_und_word':'Gleicher Vertragstext; Word Hauptvertrag editierbar, Originalangebot als Faksimile; PDF mit drei unveränderten Originalseiten.','artefakte':[]}
    for p in sorted(OUT.iterdir()):
        if p.is_file():manifest['artefakte'].append({'datei':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    (WORK/'manifest_erzeugung.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
