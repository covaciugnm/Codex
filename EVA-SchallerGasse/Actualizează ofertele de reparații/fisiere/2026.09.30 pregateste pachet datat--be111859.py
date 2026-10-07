from pathlib import Path
import shutil,json,hashlib,re,datetime
from email.message import EmailMessage
from email import policy
from email.utils import format_datetime
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parent.parent
PARTNER=ROOT/'04. Firme + Executie'/'02. Protectia Muncii (BauKG-Koordinator)'/'BAU-WERTE Lechner'
OLD=PARTNER/'2026-09-30 Propunere contract BauKG'
OUT=PARTNER/'2026.09.30 Propunere contract BauKG';OUT.mkdir(exist_ok=True)
mapping={}
for prefix,target in [('01 Vertragsentwurf','2026.09.30 BAU-WERTE - Contract propus BauKG'),('05 Audit intern','2026.09.30 BAU-WERTE - Audit intern'),('03 Anlage','2026.08.11 BAU-WERTE - Oferta 26-0094 original')]:
 for src in OLD.glob(prefix+'*'):
  p=OUT/(target+src.suffix);shutil.copy2(src,p);assert hashlib.sha256(src.read_bytes()).digest()==hashlib.sha256(p.read_bytes()).digest();mapping[src.name]=p.name
body=(OLD/'02 E-Mail Vertragsvorschlag BAU-WERTE.txt').read_text(encoding='utf-8-sig')
body=body[body.index('Sehr geehrter'):].split('Anlagen:')[0].strip()
body=re.sub(r'\b(\d{2})\.(\d{2})\.(20\d{2})\b',r'\3.\2.\1',body)
files=[OUT/'2026.09.30 BAU-WERTE - Contract propus BauKG.docx',OUT/'2026.09.30 BAU-WERTE - Contract propus BauKG.pdf']
body+='\n\nAnlagen:\n'+'\n'.join(p.name for p in files)
subject='2026.09.30 - Schallergasse 35 - Vertragsvorschlag Planungs- und Baustellenkoordination - Angebot 26-0094'
data={'data_redactarii':'2026.09.30','stare':'DRAFT - NETRIMIS','de_la':'Cosmin Adrian Covaciu <office@ac-wohnart.at>','catre':['baumeister@bau-werte.biz'],'cc':[],'subiect':subject,'corp':body,'atasamente':[p.name for p in files]}
stem='2026.09.30 DRAFT catre BAU-WERTE - Vertragsvorschlag BauKG'
txt='2026.09.30 | DRAFT - NETRIMIS\nDe la: '+data['de_la']+'\nCatre: '+', '.join(data['catre'])+'\nCC: —\nSubiect: '+subject+'\n\n'+body
(OUT/(stem+'.txt')).write_text(txt,encoding='utf-8-sig');(OUT/(stem+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
for i,p in enumerate(txt.split('\n\n')):
 if i==0:doc.add_heading('2026.09.30 | Email BAU-WERTE - netrimis',0)
 doc.add_paragraph(p)
doc.save(OUT/(stem+'.docx'))
m=EmailMessage(policy=policy.SMTP);m['From']=data['de_la'];m['To']=', '.join(data['catre']);m['Subject']=subject;m['Date']=format_datetime(datetime.datetime.now().astimezone());m['X-Unsent']='1';m['X-Project-Status']='DRAFT - NOT SENT';m['X-Document-Date']='2026.09.30';m.set_content(body)
for p in files:
 subtype='pdf' if p.suffix=='.pdf' else 'vnd.openxmlformats-officedocument.wordprocessingml.document'
 m.add_attachment(p.read_bytes(),maintype='application',subtype=subtype,filename=p.name)
(OUT/(stem+'.eml')).write_bytes(m.as_bytes())
(OUT/'2026.09.30 Citeste-ma.txt').write_text('2026.09.30 | Pachet curent BAU-WERTE\nContractul Word/PDF si oferta sunt copii identice cu documentele auditate anterior. Denumirile incep cu data YYYY.MM.DD. Emailul are data, subiectul si destinatarul completate si doua atasamente reale in EML. Stare: DRAFT, NETRIMIS. Nu s-a creat o comanda sau o numire semnata. Dosarul anterior ramane istoric.\nSursele juridice si matricea TOMS sunt in dosarul anterior, subfolderul _lucru.\n',encoding='utf-8-sig')
(OUT/'2026.09.30 Manifest documente.json').write_text(json.dumps({'data':'2026.09.30','documente':[{'nume':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'octeti':p.stat().st_size} for p in OUT.iterdir() if p.is_file()]},ensure_ascii=False,indent=2),encoding='utf-8')
print('Pachet datat si email EML cu Word/PDF anexate salvate; netrimis.')
