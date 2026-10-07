from pathlib import Path
import json,re,datetime,collections,hashlib,unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt
from email.utils import parseaddr
ROOT=Path(__file__).resolve().parent.parent
BASE=ROOT/'08. Corespondenta'/'2026.09.30 Arhiva Eva-Mail'
DATE='2026.09.30'
def safe(s,n=70):
 s=re.sub(r'[<>:"/\\|?*\r\n\t]',' ',s);return re.sub(r'\s+',' ',s).strip(' .')[:n]
def dates(s):return re.sub(r'\b(\d{2})\.(\d{2})\.(20\d{2})\b',r'\3.\2.\1',s)
def dateof(e):return (e.get('received_at') or '')[:10].replace('-','.')
own=re.compile(r'@(ac-wohnart\.at|ipec\.ro|cesiro\.(ro|com)|ig\.ro)$|covaciu.gnm@gmail.com|cesiro.horeca@gmail.com|anastasia.covaciu11@gmail.com',re.I)
known={'bau-werte.biz':('BAU-WERTE','Constructii - BauKG'),'lechner_kalender@outlook.com':('BAU-WERTE','Constructii - BauKG'),'toms.at':('TOMS','Constructii - Structura si verificare'),'otis.com':('Otis','Constructii - Lift'),'lebe-bau.at':('LEBE Bau','Constructii - Executie'),'obenauf.at':('OBENAUF','Constructii - Executie'),'novotny-bau.at':('Novotny','Constructii - Executie'),'sedlak.co.at':('Sedlak','Constructii - Executie'),'themis.co.at':('Themis','Constructii - BauKG si Bauwerksbuch'),'paknehad-bau.at':('Paknehad','Constructii - Consultanta'),'zt-pech.at':('Pech','Constructii - Consultanta'),'stolex.at':('STOLEX','Constructii - Executie'),'schauerleute.at':('SCHAUERLEUTE','Constructii - Executie'),'schmitt-aufzuege.at':('Schmitt + Sohn','Constructii - Lift'),'weigl.at':('Weigl','Constructii - Lift'),'woehrparking.at':('WOHR','Servicii - Parcari'),'adolf-tobias.at':('Adolf Tobias','Servicii - Garaj'),'peterhoenig.at':('Peter Honig','Servicii - Cosar si pod'),'attensam.at':('Attensam','Utilitati - Deratizare si deszapezire'),'sturmenergie.at':('Sturm Energie','Utilitati - Energie'),'wienenergie.at':('Wien Energie','Utilitati - Gaz si energie'),'wienernetze.at':('Wiener Netze','Utilitati - Retele'),'cerhahempel.com':('CERHA HEMPEL','Juridic si proprietate'),'hofhans.at':('Hofhans','Administrare imobil'),'frigo.at':('Frigo','Administrare imobil'),'raiffeisenbank.at':('Raiffeisen','Finantare'),'bankaustria.at':('Bank Austria','Finantare'),'unicreditgroup.at':('Bank Austria','Finantare')}
known['tobias.at']=('Adolf Tobias','Servicii - Garaj')
known['artus.at']=('ARTUS','Contabilitate si fiscalitate')
known['conrad.at']=('Conrad','Servicii - Garaj')
known['notarity.com']=('Notarity','Juridic - Notar')
known['cp-ag.at']=('CP Immobilien - Kornhofer','Administrare si predare imobil')
for domain,name,area in [('aufzug-heiszenberger.at','Heissenberger','Lift'),('bmberger.at','Markus Berger','Structura'),('hazet.at','HAZET','Executie'),('kalcon.at','KALCON','Executie'),('kone.com','KONE','Lift'),('neid.co.at','Neid','Verificare'),('sanibau.at','Sanibau','Executie'),('tkelevator.com','TK Elevator','Lift'),('zt-avunduk.at','Avunduk','Verificare'),('schindler.com','Schindler','Lift'),('vestner.at','Vestner','Lift')]:known[domain]=(name,'Constructii - '+area)
known['wienerstadtwerke.at']=('Wien Energie','Utilitati - Gaz si energie')
known['donauversicherung.at']=('Donau Versicherung','Asigurare imobil')
known['ksv.at']=('KSV1870','Creante - Contestatie')
known['maritczak.at']=('Maritczak - Commerz Inkasso','Creante - Asigurare')
known['commerz-inkasso.at']=('Maritczak - Commerz Inkasso','Creante - Asigurare')
for domain,name,area in [('kppk.at','KPPK','Constructii - Verificare'),('lv-r.at','LVR','Constructii - Proiectare'),('pm-mattes.at','Mattes Projektmanagement','Constructii - Proiectare si ÖBA'),('smt-immobilien.at','SMT Immobilien','Administrare imobil')]:known[domain]=(name,area)

def classify(e):
 subject=e.get('subject','');sender=parseaddr(e.get('from_address',''))[1].lower();tos=[parseaddr(x)[1].lower() for x in e.get('to',[])];body=e.get('text','')
 if 'lechner_kalender' in sender:return ('BAU-WERTE','Constructii - BauKG')
 addresses=[sender] if not own.search(sender) else [x.lower() for x in tos if not own.search(x)]
 for a in addresses:
  if 'wien.gv.at' in a:
   if 'ma31' in a or 'manda.jurkic' in a:return ('Wiener Wasser - MA31','Utilitati - Apa')
   if 'ma48' in a:return ('MA48 - Salubritate','Utilitati - Gunoi')
   if 'ma06' in a:return ('MA6 - Taxe','Autoritati - Taxe')
   if 'ma37' in a or 'harald.figar' in a:return ('MA37','Autoritati - Constructii')
   if 'ma29' in a or 'roman.hefter' in a:return ('MA29 - Geologie','Autoritati - Date teren')
   return ('Stadt Wien','Autoritati')
  for dom,pair in known.items():
   if dom in a:return pair
 if not addresses:
  low=(subject+' '+body[:400]).lower()
  for pat,pair in [('wasser',('Wiener Wasser - MA31','Utilitati - Apa')),('bau-werte',('BAU-WERTE','Constructii - BauKG')),('attensam',('Attensam','Utilitati - Deratizare si deszapezire')),('sturm',('Sturm Energie','Utilitati - Energie')),('wien energie',('Wien Energie','Utilitati - Gaz si energie')),('toms',('TOMS','Constructii - Structura si verificare')),('garaj',('Garaj - intern','Servicii - Garaj'))]:
   if pat in low:return pair
  return ('Coordonare interna','Proiect - Intern')
 a=addresses[0];domain=a.split('@')[-1]
 if any(x in a for x in ['uber.com','austrian.com','booking.com']):return ('Deplasari si logistica','Logistica - context')
 return (safe(domain,35),'Corespondenta proiect')
data={}
for p in sorted((BASE/'Surse').glob('*.json')):
 for e in json.loads(p.read_text(encoding='utf-8')):data[e['id']]=e
p=ROOT/'08. Corespondenta'/'BAU-WERTE'/'2026.09.30 Eva-Mail - export integral.json'
if p.exists():
 for e in json.loads(p.read_text(encoding='utf-8'))['emailuri']:data[e['id']]=e
emails=sorted(data.values(),key=lambda e:e.get('received_at',''))
expected=json.loads((ROOT/'folder map'/'2026.09.30 lista mesaje candidate.json').read_text(encoding='utf-8'))
complete=set(e['id'] for e in expected).issubset(data)
manualpath=ROOT/'folder map'/'2026.09.30 statusuri verificate.json'
manual=json.loads(manualpath.read_text(encoding='utf-8')) if manualpath.exists() else {}
attach={}
for p in [ROOT/'08. Corespondenta'/'BAU-WERTE'/'2026.09.30 Registru atasamente.json',BASE/'2026.09.30 Registru atasamente.json']:
 if p.exists():
  for r in json.loads(p.read_text(encoding='utf-8')):attach[r['id']]=r
grouped=collections.defaultdict(list);rows=[]
for e in emails:
 partner,area=classify(e);direction='TRIMIS' if e.get('is_sent') else 'PRIMIT'
 if own.search(e.get('from_address','')) and not e.get('is_sent'):direction='PRIMIT INTERN'
 if e.get('text','').lstrip().startswith('BEGIN:VCALENDAR'):direction='INVITATIE'
 stamp=dateof(e);time=(e.get('received_at','')[11:16] or '00:00').replace(':','-')
 title=safe(e.get('subject','Fara subiect'),40)
 folder=BASE/'Emailuri'/safe(partner,35);folder.mkdir(parents=True,exist_ok=True)
 name=f'{stamp} {time} {direction} {title} {e["id"][:8]}'
 p=folder/(name+'.txt')
 ats=[]
 for a in e.get('attachments',[]):
  inline=bool(re.fullmatch(r'(image\d+\.(jpg|png|jpeg|gif)|inline|smime\.p7s)',a.get('name',''),re.I) and (a.get('size') or 0)<40000 and (a.get('text_chars') or 0)<40)
  r=attach.get(a['id'],{});ats.append({'id':a['id'],'nume_original':a['name'],'cale':r.get('cale',''),'stare':r.get('stare','metadate; imagine mica probabila semnatura' if inline else 'de descarcat'),'octeti':a.get('size')})
 header=f'{stamp} {time.replace("-",":")} UTC | {direction} | {partner}\nData export: {DATE}\nDe la: {e.get("from_name","")} <{e.get("from_address","")}>\nCatre: {", ".join(e.get("to",[]))}\nCC: {", ".join(e.get("cc",[])) or "—"}\nSubiect: {e.get("subject","")}\nCasuta: {e.get("account_email","")}\nID Eva-Mail: {e["id"]}\nStare sursa: export text integral furnizat de Eva-Mail; formatul MIME/HTML original nu este disponibil in acest export.\n'
 header=header.replace('export text integral furnizat','export text disponibil furnizat')
 if e.get('truncated'):header+='LIMITA EVA: corp limitat la 20.000 caractere; restul lantului citat nu este disponibil prin instrumentul de export.\n'
 header+='\nAtasamente:\n'+'\n'.join(f'- {a["nume_original"]} | {a["stare"]} | {a["cale"]}' for a in ats)
 p.write_text(header+'\n\n--- TEXTUL FURNIZAT DE EVA, NEMODIFICAT ---\n\n'+e.get('text',''),encoding='utf-8-sig')
 meta={k:v for k,v in e.items() if k not in ('attachments','text')};meta.update(data_afisata=stamp,partener=partner,domeniu=area,atasamente=ats,text_fisier=p.name)
 (folder/(name+'.json')).write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
 if direction=='INVITATIE':(folder/(name+'.ics')).write_text(e['text'],encoding='utf-8')
 row={'data':stamp,'ora_utc':time.replace('-',':'),'partener':partner,'domeniu':area,'sens':direction,'subiect':e.get('subject','').replace('\r',' ').replace('\n',' '),'expeditor':e.get('from_address',''),'destinatari':'; '.join(e.get('to',[])),'cc':'; '.join(e.get('cc',[])),'id':e['id'],'cale':p.relative_to(ROOT).as_posix(),'atasamente':len(ats),'trunchiat':e.get('truncated',False)}
 rows.append(row);grouped[partner].append((e,row))
statuses=[]
for partner,items in sorted(grouped.items()):
 recv=[(e,r) for e,r in items if not e.get('is_sent') and r['sens']=='PRIMIT']
 last=items[-1];latest=recv[-1] if recv else None;m=manual.get(partner,{})
 excerpt=re.sub(r'\s+',' ',latest[0].get('text','')).strip()[:450] if latest else ''
 status=m.get('status',('Ultimul mesaj extern, '+latest[1]['data']+': '+excerpt+' [fragment automat; vezi mesajul integral]') if latest else 'Nu este identificat un raspuns extern in mesajele preluate. Vezi istoricul solicitarilor si comunicarea interna.')
 action=m.get('urmatorul_pas','Revizuire si raspuns pe baza ultimului mesaj; nicio actiune externa efectuata.')
 statuses.append({'partener':partner,'domeniu':last[1]['domeniu'],'status':status,'ultima_comunicare':last[1]['data'],'ultimul_primit':latest[1]['data'] if latest else '—','subiect_ultimul_primit':latest[1]['subiect'] if latest else '—','urmatorul_pas':action,'sursa_status':m.get('surse',[]),'mesaje':len(items),'nivel':'Sinteza verificata' if m else 'Index cronologic; continut disponibil prin Eva'})
out={'data_actualizarii':DATE,'acoperire':{'mesaje_candidate':len(expected),'mesaje_salvate':len(rows),'preluare_mesaje_completa':complete,'atasamente_salvate':sum(r.get('stare')=='salvat original integral' for r in attach.values()),'nota':'Cautari pe adresa proiectului si servicii conexe in casutele accesibile Eva-Mail; ultimele sincronizari difera. Mesajele originale isi pastreaza datele si textul. Orele din log sunt UTC.'},'statusuri':statuses,'comunicatii':rows}
out['acoperire']['corpuri_trunchiate_de_Eva']=sum(bool(e.get('truncated')) for e in emails)
out['acoperire']['atasamente_identificate']=len({a['id'] for e in emails for a in e.get('attachments',[])})
allats={a['id']:a for e in emails for a in e.get('attachments',[])}
inlineids={a['id'] for a in allats.values() if re.fullmatch(r'(image\d+\.(jpg|png|jpeg|gif)|inline|smime\.p7s)',a.get('name',''),re.I) and (a.get('size') or 0)<40000 and (a.get('text_chars') or 0)<40}
savedids={k for k,v in attach.items() if v.get('cale') and (ROOT/v['cale']).is_file() and v.get('stare')=='salvat original integral'}
out['acoperire'].update(atasamente_documentare=len(set(allats)-inlineids),atasamente_documentare_salvate=len((set(allats)-inlineids)&savedids),imagini_mici_numai_metadate=len(inlineids-savedids),mesaje_suplimentare_fata_de_lista=len(set(data)-{x['id'] for x in expected}))
out['acoperire']['nota']+=' Unele corpuri sunt limitate de API la 20.000 caractere; randurile respective sunt marcate trunchiat. Salvarea tuturor ID-urilor nu inseamna export MIME integral.'
out['acoperire']['nota']+=' Imaginile mici cu nume generic, fara text, probabile semnaturi, au numai metadate; unele emailuri furnizeaza numai avertismentul de securitate. Anexele prin link temporar expirat nu sunt considerate descarcate.'
(BASE/'2026.09.30 Registru comunicatii.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
wb=Workbook();ws=wb.active;ws.title='Status pe scurt'
cols=list(statuses[0]) if statuses else []
if cols:
 ws.append(cols)
 for s in statuses:ws.append(['; '.join(s[c]) if isinstance(s[c],list) else s[c] for c in cols])
for name,values in [('Comunicatii',rows),('Atasamente',list(attach.values()))]:
 sh=wb.create_sheet(name)
 if values:
  cols=list(dict.fromkeys(k for r in values for k in r if k not in ['text','error']))
  sh.append(cols)
  for r in values:sh.append([str(r.get(c,'')) if isinstance(r.get(c), (dict,list)) else r.get(c,'') for c in cols])
  if 'cale' in cols:
   ci=cols.index('cale')+1
   for ri in range(2,sh.max_row+1):
    val=sh.cell(ri,ci).value
    if val:sh.cell(ri,ci).hyperlink=(ROOT/val).as_uri()
sh=wb.create_sheet('Acoperire');sh.append(['Camp','Valoare'])
for k,v in out['acoperire'].items():sh.append([k,str(v)])
for sh in wb:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24486B')
 for row in sh.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=24 if col[0].value not in ['status','subiect','urmatorul_pas','cale'] else 65
wb.save(ROOT/'2026.09.30 Log progres proiect.xlsx')
def makedoc(title,paras,path):
 d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10.5);d.add_heading(title,0)
 for kind,text in paras:
  if kind=='h':d.add_heading(text,1)
  else:d.add_paragraph(text)
 d.save(path)
summary=[('p',f'{DATE} | Actualizare proiect Schallergasse 35. Preluare emailuri: '+('completa pentru lista cautata' if complete else 'IN CURS')+f'. {len(rows)} mesaje salvate; {len(expected)} in lista cautata, {out["acoperire"]["mesaje_suplimentare_fata_de_lista"]} suplimentare.')]
summary.append(('p',f'Anexe documentare salvate: {out["acoperire"]["atasamente_documentare_salvate"]} / {out["acoperire"]["atasamente_documentare"]}. Imagini mici pastrate numai ca metadate: {out["acoperire"]["imagini_mici_numai_metadate"]}. Corpuri limitate de Eva la 20.000 caractere: {out["acoperire"]["corpuri_trunchiate_de_Eva"]}. Orele sunt UTC. Exportul este textul disponibil prin API, uneori numai avertisment de securitate; nu este export MIME integral.'))
summary.append(('h','Status pe scurt'))
for s in statuses:
 if s['nivel']=='Sinteza verificata':summary.append(('p',f'{s["partener"]}: {s["status"]}\nUrmatorul pas: {s["urmatorul_pas"]}'))
summary += [('h','Progresul documentarii'),('p',f'{DATE} — Contractul BAU-WERTE a fost pregatit in Word/PDF, cu oferta originala integrala si audit. Emailul de propunere este DRAFT, NETRIMIS.\n{DATE} — Au fost create jurnale pe partener si arhiva locala a comunicatiilor, cu subiect, expeditor, destinatari, CC, ID, data si atasamente.\n{DATE} — Regula persistenta de denumire: YYYY.MM.DD la inceputul numelor noi. Originalele si metadatele standard raman neschimbate.'),('h','Trasabilitate'),('p','Registrul Excel contine toate mesajele, legaturile locale, atasamentele si statusurile. Fisele partenerilor sunt in 08. Corespondenta/2026.09.30 Arhiva Eva-Mail/Parteneri. Situatiile istorice raman in arhiva.')]
makedoc(DATE+' | Log progres proiect',summary,ROOT/'2026.09.30 Log progres proiect.docx')
brief=summary[:2]+[('h','Financiar'),('p','62 miscari bancare inregistrate; 49 pozitii documentare. 1.538,12 EUR plati personale pentru obligatii asociate firmei. 2.877,45 EUR documente fara debit identificat: CERHA 2.640,22 + ARTUS 144 + MA6 Q3 93,23. Sumele contestate, estimarile si ordinul apa de 15,18 sunt tratate separat. Extrase generale pana la 2026.09.28; zilele 29–30 nu sunt acoperite integral. Registru: 10. Banci + Extrase de cont/2026.09.30 Reconciliere completa facturi si plati.xlsx.')]
for s in statuses:
 if s['partener'] in ['BAU-WERTE','TOMS','Attensam','Donau Versicherung','Themis','Peter Honig','Sturm Energie']:brief.append(('p',s['partener']+': '+s['status']))
brief.append(('p','Detalii si urmatorii pasi: 2026.09.30 Log progres proiect.xlsx; toate jurnalele partenerilor: 08. Corespondenta/2026.09.30 Arhiva Eva-Mail/Parteneri. Oferte: 04. Firme + Executie/2026.09.30 Ultimele raspunsuri ofertanti.xlsx.'))
(ROOT/'2026.09.30 Status proiect.txt').write_text('\n\n'.join(t for _,t in brief),encoding='utf-8-sig')
pd=BASE/'Parteneri';pd.mkdir(exist_ok=True)
for s in statuses:
 items=grouped[s['partener']]
 lines=[('h','Status pe scurt'),('p',DATE+' | '+s['status']),('p','Ultimul mesaj primit: '+s['ultimul_primit']+' | '+s['subiect_ultimul_primit']),('p','Urmatorul pas: '+s['urmatorul_pas']),('h','Istoric discutii — cele mai recente primele')]
 for e,r in reversed(items):
  lines.append(('p',r['data']+' '+r['ora_utc']+' UTC | '+r['sens']+'\nSubiect: '+r['subiect']+'\nDe la: '+r['expeditor']+'\nCatre: '+r['destinatari']+(' | CC: '+r['cc'] if r['cc'] else '')+'\nText disponibil: '+r['cale']+'\nID: '+r['id']))
 if s['partener']=='BAU-WERTE':
  lines.insert(5,('p',DATE+' | DRAFT, NETRIMIS — Propunere contract 8.250 EUR net, 12+3 luni; catre baumeister@bau-werte.biz. Word/PDF si EML cu atasamente in dosarul 2026.09.30 Propunere contract BauKG.'))
 name=DATE+' Log discutii - '+safe(s['partener'],38)
 (pd/(name+'.txt')).write_text('\n\n'.join(t for _,t in lines),encoding='utf-8-sig')
 if s['partener'] in manual or s['partener']=='BAU-WERTE':makedoc(name,lines,pd/(name+'.docx'))
 if s['partener']=='BAU-WERTE':
  target=ROOT/'04. Firme + Executie'/'02. Protectia Muncii (BauKG-Koordinator)'/'BAU-WERTE Lechner'
  import shutil
  for ext in ['.txt','.docx']:shutil.copy2(pd/(name+ext),target/(DATE+' Log discutii BAU-WERTE'+ext))
# Elimina numai exporturile noastre devenite redundante dupa corectarea clasificarii.
validpaths={str((ROOT/r['cale']).with_suffix(ext)) for r in rows for ext in ['.txt','.json','.ics']}
for p in (BASE/'Emailuri').rglob('*'):
 if p.is_file() and p.suffix in ['.txt','.json','.ics'] and re.search(r' [a-f0-9]{8}$',p.stem) and str(p) not in validpaths:p.unlink()
print(json.dumps({'mesaje':len(rows),'candidate':len(expected),'parteneri':len(statuses),'preluare_completa':complete,'atasamente':len(attach),'acoperire':out['acoperire']},ensure_ascii=False))
