from pathlib import Path
import json, re, datetime
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document

BASE=Path(__file__).resolve().parent;ROOT=BASE.parent
FIRME=ROOT/'04. Firme + Executie'
DEST=FIRME/'Actualizare oferte 2026.09.30'
DEST.mkdir(exist_ok=True)
inventory=json.loads((BASE/'inventar.json').read_text(encoding='utf-8'))
def src(name):
    matches=[r['cale'] for r in inventory if name in Path(r['cale']).name]
    if not matches:raise ValueError(name)
    return matches[0]
def link(path,label=None):return f'[{label or Path(path).name}]({quote(str(ROOT/path).replace(chr(92),chr(47)),safe="/:")})'

T28=src('2026-09-28 Fragekatalog');T23=src('2026-09-23 Anlage 5');GU=src('06 - Comparativ GU');REG=src('Registru Comunicari si Oferte - Schallergasse 35 - 2026-09-17.xlsx')
R=[]
def add(service,firm,date,status,price,reply,nextstep,sources,level='Document / email local verificat'):
    R.append(dict(serviciu=service,firma=firm,data_ultimei_dovezi=date,status=status,pret_net=price,ultimul_raspuns=reply,pas_urmator=nextstep,surse=sources,nivel_dovada=level,verificare_mail='Eva-Mail neconectat; fără verificare live la 30.09.2026'))
add('Statică + Prüfingenieur','TOMS','2026-09-28','Ofertă și clarificări; limite de prestație deschise','34.500 EUR (26.500 + 8.000)',
    'Chestionar 2 semnat: calculele nodurilor oțel/lemn și unele ancore/suduri contra cost; detalii desenate și alte elemente excluse. Termenele propuse rămân nebifate. D1: 3 săptămâni după contract semnat; D2 de stabilit de comun acord.',
    'Cere prețuri pentru suplimente, responsabil pentru detaliile excluse, termen ferm pe pachet și contract consolidat. Valabilitatea ANG 839 prelungită la 03.10 în răspunsul din 23.09; ANG 855 până la 03.10.',[T28,T23,src('ANG 855_Schallergasse'),src('ANG 839_Schallergasse')])
add('Bauwerksbuch','TOMS','2026-09-28','Ofertă; fără reducere confirmată','3.400 EUR',
    'ANG 891 din 23.09; răspunsul 28.09 include pozițiile 1–7 și 10–12, exclude documentarea foto/fisuri înainte și după lucrări și inventarul planurilor/staticilor/avizelor. Câmpurile de onorariu redus și termen sunt goale.',
    'Clarifică prețul final, excluderile și data predării. Oferta ANG 891 valabilă până la 31.12.2026.',[T28,src('2026-09-23 ANG 891')])
add('BauKG','TOMS','2026-09-23','Ofertă pentru contractare ulterioară','11.040 EUR la 12 luni',
    'Răspunsul semnat la întrebarea 20 prelungește ANG 840 până la 31.12.2026; vizite săptămânale, două întâlniri de coordonare incluse, vizită suplimentară 200 EUR; peste 12 luni se facturează o sumă lunară.',
    'Solicită versiunea consolidată ANG 840 și valoarea lunară pentru prelungire.',[T23,src('ANG 840_Schallergasse')])
add('Lift','Otis','2026-09-24','Ofertă primită; compatibilitate tehnică de clarificat','38.680 EUR',
    'Gen3 Core, 630 kg / 8 persoane, 6 opriri, 1 m/s; puț 1600×1750 mm, cap superior 3600 mm, groapă 1000 mm. Livrare 12–14 săptămâni după acordul tehnic/comercial final. Plată 40/40/20.',
    'Clarifică 3600 mm față de aproximativ 3350 mm din cererea proiectului, sarcinile pentru structurist, mentenanța și oferta pentru Treppenlift.',[src('2026-09-24 Otis Angebot'),src('2026-09-24 10-52 Otis'),src('Email-Text_Anfrage_Aufzug_MASTER_v3')])
add('Execuție GU','LEBE Bau','2026-09-24','Prețuri orientative; fără ofertă globală fermă','10 poziții orientative; fără total',
    'Comparativul local consemnează răspunsul de la 17:38: prețuri foarte aproximative, fără planuri/detalii suficiente. Vizita 25.09 ora 10 era confirmată; început realist noiembrie/decembrie.',
    'Recuperează emailul original 24.09 și eventualul atașament; verifică rezultatul vizitei și cere ofertă pe cantități și planuri.',[GU,src('2026-09-10 LEBE Bau Office')], 'Răspunsul 24.09 este consemnat într-un raport local; originalul nu a fost identificat separat')
add('Execuție GU','OBENAUF','2026-09-22','Interes și prețuri orientative completate','10 poziții; GU-Zuschlag auf Sub 20%',
    'Anlage A completată există din 16.09. Confirmarea întâlnirii 25.09 ora 13 este consemnată în comparativ la 22.09. Adaosul 20% este etichetat pe subcontractări, nu automat pe întregul proiect.',
    'Verifică rezultatul vizitei; cere ofertă globală și baza exactă a adaosului de 20%.',[src('Anlage_A_Richtpreisliste_Obenauf.pdf'),GU], 'Prețuri din PDF original; confirmarea întâlnirii din raport local')
add('Execuție GU','Ing. Felix Novotny','2026-09-24','Interes condiționat; fără prețuri', '—',
    'Conform comparativului, cere răspuns scris, vizită, apoi prețuri; întreabă dacă LV este funcțional sau detaliat și se retrage în cazul unei descrieri funcționale. Nu are Anlage A/B completată. Vizita de vineri nu mai era valabilă.',
    'Recuperează mesajele 24.09, clarifică tipul LV și propune un nou termen după verificarea corespondenței.',[GU,src('Raspuns_Novotny_2026-09-21')], 'Răspunsul 24.09 din raport local; răspunsul 21.09 din rezumat local')
add('Execuție GU','DI Wilhelm Sedlak','2026-09-21','Solicită documentația; fără ofertă','—',
    'Michael Hollinger cere Ausschreibungsunterlagen pentru evaluare. Raportul din 24.09 nu identifica încă transmiterea linkului. Novotny aparține aceluiași grup, conform documentelor locale.',
    'Verifică dacă dosarul a fost transmis după 24.09; pregătește trimiterea dacă lipsește.',[src('Raspuns_Sedlak_2026-09-21'),GU])
for firm,pattern,reason in [
    ('KALCON','Raspuns_KALCON','Refuz la 16.09: nu preia GU la acest volum pe baza unei descrieri funcționale.'),
    ('Sanibau','Raspuns_Sanibau','Refuz la 16.09 din lipsă de capacitate.'),
    ('HAZET','2026-09-16 Hazet','Refuz la 16.09 din lipsă de capacitate.')]:
    add('Execuție GU',firm,'2026-09-16','Refuz', '—',reason,'Închide ca refuz pentru această rundă; urmărește doar dacă există mesaj nou.',[src(pattern)])
for firm in ['Nems Bau','Giefing Bau','Ing. Kurt Hammerl','Sattler Bau','Ing. Adolf Klein','StoleX','DEZET Bau','König Heinrich','Baumeister Jovicic','GERSTL BAU']:
    add('Execuție GU',firm,'2026-09-24','Fără răspuns identificat în raportul local','—','Raportul local din 24.09 consemnează lipsa răspunsului. Nu confirmă situația la 30.09.','Verifică Eva-Mail pentru perioada după raport înainte de revenire.',[GU],'Raport local, fără verificare live')
add('BauKG','BAU-WERTE','2026-09-22','Ofertă istorică; valabilitate de reînnoit','8.250 EUR',
    '750 EUR planificare + 7.500 EUR șantier pentru 12 luni. Oferta inițială a expirat la 10.09. Dosarul de întâlnire consemnează mutarea confirmată la 25.09, 09:30–10:00.',
    'Verifică rezultatul întâlnirii; cere reînnoirea scrisă, frecvența vizitelor și tariful peste 12 luni.',[src('Honorarangebot 2026-08-11'),src('02 - 09-30 BAU-WERTE')],'Ofertă originală; corespondența 22.09 descrisă în dosarul local')
add('BauKG','Themis','2026-08-10','Ofertă / estimare de cost','3.500 EUR + 300 EUR/săptămână; baza TVA de confirmat',
    'SiGe-Plan 1.900, notificare 300, documentație lucrări ulterioare 1.300; coordonare 300 EUR/săptămână. Valabilitate trei luni. Condițiile generale descriu și alte servicii, fără a le include automat în această ofertă BauKG.',
    'Confirmă durata, baza TVA și serviciile suplimentare înainte de comparația totalurilor.',[src('Angebot_E-2608')])
add('BauKG','DI Paknehad','2026-08-29','Indicație neangajantă','3.500–6.900 EUR + 375 EUR/vizită',
    'Prezentarea precizează explicit că nu este ofertă fermă; fără durată și fără total de contract.',
    'Cere ofertă fermă pe dosarul complet și calendar.',[src('260829_Preisindikation')])
add('Statică + Prüfingenieur','Markus Berger','2026-08-17','Ofertă combinată','45.000 EUR',
    '13.000 calcule + 20.000 planuri + 2.000 asistență (maximum 8 vizite) + 10.000 Prüfingenieur. Nu este preț pentru statică singură. Lucrările de sprijinire a excavației sunt excluse. Valabilitate 8 săptămâni.',
    'Compară pe același pachet cu TOMS; confirmă termenul și limitele; contractarea parțială necesită acord.',[src('2026_08_17_Anbot_Statik_Detail.pdf')])
add('Statică','Dr. PECH','2026-09-17','Ofertă primită','69.300 EUR',
    '42.000 execuție statică + 6.000 goluri în existent + 5.000 sprijinire excavație + 800 teren + 14.000 asistență (12 recepții) + 1.500 documentație finală. Recepție în plus: 450 EUR. Valabilitate 90 zile.',
    'Compară includerile cu TOMS/Berger și clarifică rolul Prüfingenieur separat.',[src('Angebot 2026-470-AN01')])
add('Statică / Prüfingenieur','DI Remzi Avunduk','2026-09-10','Refuz','—',
    'A refuzat după acceptarea a două proiecte noi; capacitate depășită. Registrul vechi îl păstra eronat în așteptare.',
    'Înregistrează refuzul pentru această rundă.',[src('2026-09-10 Dipl.-Ing. Remzi')])
add('Prüfingenieur','KPPK','2026-08-11','Refuz','—','Refuz din lipsă de capacitate.','Fără acțiune până la un eventual mesaj nou.',[src('Email-Text_Absage_2026-08-11')])
for firm,service in [('DI Janka Neid','Statică / Prüfingenieur'),('PCD ZT','Prüfingenieur'),('Potyka & Partner','Prüfingenieur'),('BK Baumanagement Kazda','BauKG'),('SSB','BauKG')]:
    add(service,firm,'2026-09-17','Necesită reconfirmare în Eva-Mail','—',
        'Registrul vechi și auditurile nu oferă o confirmare recentă suficientă. Pentru Neid există informații istorice contradictorii despre interes și refuz.',
        'Citește ultimul fir complet înainte de a marca activ sau fără răspuns.',[REG],'Registru istoric; stare neconfirmată')
for firm in ['Schindler','KONE','TK Aufzüge','Schmitt + Sohn','Vestner','Füglister','Heissenberger','WEIGL']:
    add('Lift',firm,'2026-09-24','Nicio ofertă identificată în registrul local','—','Există cerere și/sau draft; câmpurile de răspuns din registrul lift sunt goale. Aceasta nu dovedește absența răspunsului în email.','Verifică Eva-Mail; păstrează separat liftul principal și Treppenlift.',[src('Registru_Oferte_Lift')],'Registru local, fără verificare live')
add('Administrare','Frigo','2026-09-10','Ofertă actualizată retransmisă; acceptare neconfirmată','305,85 EUR/lună + 490 EUR/an (registru istoric)',
    'Emailul 10.09 spune că trimite oferta actualizată și materialul informativ. Este ulterior închiderii din 08.09; eticheta folderului INCHIS nu descrie singură ultima corespondență.',
    'Confirmă dacă oferta mai este disponibilă și dacă există o decizie ulterioară.',[src('2026-09-10 Frigo Peter'),REG],'Email original local; preț preluat din registrul istoric')
add('Administrare','IMV','2026-09-17','Ofertă financiară neconfirmată','—','PDF-ul GBA este extras de carte funciară, nu dovadă de ofertă financiară.','Caută oferta reală în Eva-Mail sau cere retransmitere după verificare.',[REG],'Corecție consemnată în registrul istoric')

corrections=[
('KALCON','Interes','Refuz 16.09, descriere funcțională neacceptată.'),
('Avunduk','În așteptare','Refuz explicit 10.09 din lipsă de capacitate.'),
('OBENAUF','Richtpreisliste necompletată','PDF completat cu 10 poziții; 20% este etichetat GU-Zuschlag auf Sub.'),
('LEBE','Prețuri urmează','Raportul local 24.09 consemnează 10 prețuri; originalul emailului lipsește separat.'),
('Novotny / Sedlak','Fără răspuns','Interes condiționat / cerere de documente, fără ofertă fermă.'),
('Berger','45.000 EUR statică','45.000 EUR include 10.000 EUR pentru Prüfingenieur; baza statică + planuri + asistență este 35.000 EUR, fără drept automat la contractare separată.'),
('TOMS ANG 839 / ANG 840','Valabilitate din oferta inițială','Răspunsul 23.09 prelungește ANG 839 până la 03.10 și ANG 840 până la 31.12.'),
('TOMS termene','Calendarul propus tratat ca acceptat','Răspunsul 28.09 lasă confirmările nebifate și introduce termene relative/de negociat.'),
('TOMS / independență','Interdicție categorică a cumulului statică + PI','TOMS declară în răspunsul 23.09 personalunion și confirmarea independenței; registrul actual consemnează afirmația ofertantului, fără concluzie juridică proprie.'),
('Frigo','Închis','Închidere 08.09 urmată de ofertă actualizată transmisă la 10.09; decizia ulterioară necunoscută.'),
('Total consultanți','45.290 EUR pentru 26.500 + 8.000 + 8.250','Suma aritmetică este 42.750 EUR. Cu ÖBA 98.000 rezultă 140.750 EUR. Acestea nu sunt bugete aprobate sau pachete cu servicii echivalente.'),
('Întâlniri 25.09','Acțiuni prezentate ca viitoare','Datele au trecut; trebuie verificat dacă întâlnirile s-au ținut și ce s-a decis.')]

data={'data_actualizarii':'2026-09-30','acoperire':'Surse locale; Eva-Mail nu este disponibil. Nu s-au descărcat oferte noi din email.','inregistrari':R,'corectii':corrections}
(DEST/'situatie_structurata.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
wb=Workbook();ws=wb.active;ws.title='Situație curentă'
columns=['serviciu','firma','data_ultimei_dovezi','status','pret_net','ultimul_raspuns','pas_urmator','nivel_dovada','surse','verificare_mail']
ws.append(columns)
for row in R:ws.append(['\n'.join(row[k]) if isinstance(row[k],list) else row[k] for k in columns])
for i,row in enumerate(R,2):ws.cell(i,9).hyperlink=(ROOT/row['surse'][0]).as_uri()
c=wb.create_sheet('Corecții față de 17 sept');c.append(['Subiect','Registrul anterior','Corecție documentată'])
for row in corrections:c.append(row)
v=wb.create_sheet('Valabilități');v.append(['Ofertă','Data PDF','Valabilitate','Dovadă ulterioară','Acțiune'])
for row in [
('TOMS ANG 855','03.09.2026','03.10.2026','Discuții până la 28.09','Confirmare ofertă consolidată'),
('TOMS ANG 839','20.08.2026','19.09 inițial','Prelungire la 03.10 în răspuns 23.09','Confirmare contract consolidat'),
('TOMS ANG 840','20.08.2026','19.09 inițial','Prelungire la 31.12 în răspuns 23.09','Clarificare tarif lunar'),
('TOMS ANG 891','23.09.2026','31.12.2026','Fără discount completat la 28.09','Clarificare excluderi'),
('BAU-WERTE 26-0094','11.08.2026','10.09.2026','Nicio prelungire scrisă identificată','Reînnoire'),
('Otis 32KB1294-01','24.09.2026','30 zile; calcul calendaristic 24.10.2026','Preț fix până la 31.12.2027 este o clauză distinctă','Clarificări tehnice'),
('Berger','17.08.2026','8 săptămâni; calcul calendaristic 12.10.2026','Nicio modificare identificată','Confirmare condiții'),
('Dr. PECH','17.09.2026','90 zile; calcul calendaristic 16.12.2026','Nicio modificare identificată','Comparare includeri'),
('Themis','10.08.2026','3 luni','Nicio modificare identificată','Confirmare cost la durata reală')]:v.append(row)
gu=wb.create_sheet('Prețuri GU orientative');gu.append(['Poziție','Lucrare','Unitate','OBENAUF net','LEBE net','Dovada LEBE','Observație'])
works=[('P01','Demolare pereți ≤15 cm','m²',45,85),('P02','Gips-carton dublu placat','m²',90,160),('P03','Radier C25/30 30 cm','m²',300,384.75),('P04','Oțel HEB 120–180','kg',7.5,12.35),('P05','Structură lemn mansardă','m³',850,1650),('P06','Acoperiș DA01','m²',450,450),('P07','Șapă 6 cm','m²',65,43),('P08','Electricitate apartament','apartament',10500,10500),('P09','Baie fără placări','bucată',3500,9500),('P10','Schelă 3 luni','m²',11.5,18.5)]
for row in works:gu.append([*row,'Raport local 24.09; email original de recuperat','Fără cantități și fără total. Adaos OBENAUF 20% pe subcontractări: bază de clarificat.'])
n=wb.create_sheet('Acoperire');n.append(['Element','Situație']);n.append(['Actualizare','30.09.2026']);n.append(['Email Eva-Mail','Neconectat. Actualizarea live și descărcarea ofertelor noi rămân de făcut.']);n.append(['Gmail verificat','covaciu.gnm@gmail.com; căutare Schallergasse după 20.09 a returnat documente bancare/călătorie, fără oferte relevante.']);n.append(['Sume','EUR fără TVA unde sursa indică net; costurile și includerile diferă.']);n.append(['Istoric','Fișierul 17.09 este păstrat separat și conține informații depășite.']);n.append(['Prețuri vechi execuție și ÖBA','Ofertele Schwarz, Gschirtz, LVR/Mattes și Hofhans rămân în dosarele istorice; nu sunt prezentate drept răspunsuri recente.'])
for sh in wb:
    sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
    for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24486B')
    for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=24 if col[0].column<6 else 65
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,str):cell.data_type='s'
    for i in range(2,sh.max_row+1):sh.row_dimensions[i].height=95
out=FIRME/'Registru Comunicari si Oferte - Schallergasse 35 - 2026-09-30.xlsx';wb.save(out)

intro='''# Situația ofertelor Schallergasse 35 la 30 septembrie 2026

Actualizare pe baza arhivei locale. Ultima clarificare documentată este răspunsul TOMS semnat la 28 septembrie; ultima ofertă de lift identificată este Otis din 24 septembrie. Eva-Mail nu este disponibil în această sesiune: mesajele noi și atașamentele de după documentele existente nu sunt verificate sau descărcate.

## Ce necesită atenție acum

1. **TOMS:** prețul de bază statică + Prüfingenieur este 34.500 EUR net, dar răspunsul din 28.09 conține suplimente fără preț și excluderi de detalii. Calendarul propus nu este confirmat. ANG 855 și prelungirea ANG 839 ajung la 03.10.2026.
2. **Otis:** 38.680 EUR net, cu cap de puț cerut de 3.600 mm. Cererea proiectului indică aproximativ 3.350 mm; diferența de 250 mm trebuie rezolvată în proiect înaintea deciziei.
3. **LEBE și OBENAUF:** există prețuri orientative, fără cantități și fără ofertă totală fermă comparabilă. Prețurile LEBE din 24.09 sunt păstrate în comparativul local; emailul original trebuie recuperat.
4. **Novotny și Sedlak:** există răspunsuri, dar fără ofertă. Novotny condiționează participarea de tipul documentației; Sedlak cere dosarul de licitație.
5. **BAU-WERTE:** oferta de 8.250 EUR avea valabilitate până la 10.09; o prelungire scrisă nu a fost identificată. Întâlnirile din 25.09 sunt date trecute, iar rezultatele lor nu sunt confirmate de simpla existență a programului.

## Clarificările TOMS din 28 septembrie

Verificare vizuală a căsuțelor din PDF, paginile 2–4:

- **Contra cost:** calculul nodurilor și îmbinărilor la oțel și lemn, ancore/dibluri în zidărie sau beton, grosimi de table și suduri (A1, A2, A5, A6). Nu este indicată suma suplimentară.
- **Incluse:** detaliile de reazem pe zidăria existentă (A9).
- **Excluse sau nepredate în forma cerută:** desenele detaliate ale îmbinărilor, tipul/numărul/dispunerea conectorilor, listele pe element în plan, îmbinările contravântuirilor, specificațiile de protecție, planurile de atelier (A3, A4, A8, A10–A12). A7 este bifat „nu” pentru forțele înscrise în planuri, cu precizarea că acestea sunt în calculul static.
- **Calendar:** căsuțele de confirmare sunt goale. La D1 apare „3 săptămâni după contract semnat”; la D2 apare că termenul va fi convenit ulterior. Nu se poate păstra calendarul propus ca angajament acceptat.
- **Formate:** PDF, DWG, DXF, DWG la revizii, IFC, listă de planuri și link de descărcare sunt bifate. La DWG se adaugă „scara 1:50”, deși cererea solicita model 1:1. Nu sunt confirmate straturile separate, modelul de calcul, Excel/CSV pentru armături, BTL/BVX și NC/DSTV.
- **Bauwerksbuch:** pozițiile 1–7 și 10–12 sunt incluse; documentarea foto/fisuri și inventarul documentelor sunt excluse (E8–E9). Câmpurile pentru reducere și termen sunt necompletate. Oferta de 3.400 EUR net nu are reducere confirmată.

Răspunsul din 23.09 este relevant pentru prelungirea valabilităților și condițiile deja clarificate. El declară menținerea cumulului Tragwerksplanung/Prüfingenieur și confirmarea independenței; aceasta este poziția ofertantului. Registrul vechi formula o interdicție categorică fără analiza necesară. Prezenta situație nu stabilește o concluzie juridică.

## Otis și costurile recurente

Oferta 32KB1294-01 are 14 pagini și cuprinde 630 kg / 8 persoane, șase opriri, 1 m/s, puț 1.600 × 1.750 mm, cap 3.600 mm și groapă 1.000 mm. Livrarea este de 12–14 săptămâni după acordul tehnic și comercial final; montajul este condiționat de verificarea/aprobarea pozitivă indicată în ofertă. Plăți: 40% la contract, 40% la eliberarea în producție, 20% la recepția tehnică.

Pagina 12 listează mentenanță de bază 1.620 EUR/an sau completă 2.540 EUR/an în perioada garanției și 3.260 EUR/an după aceasta, control de funcționare 186 EUR/an, comunicație apel de urgență 1.355 EUR/an și modul mobil 512 EUR/an. Dacă sunt cumulative, scenariile aritmetice sunt 3.673 / 4.593 / 5.313 EUR net pe an. Trebuie confirmat pachetul aplicabil; verificările periodice independente nu au preț inclus identificat. Prețul fix până la 31.12.2027 nu prelungește automat valabilitatea de 30 de zile a ofertei.

## Ultimele răspunsuri pe ofertant
'''
parts=[intro]
for r in R:
    if r['firma'] in ['TOMS','Otis','LEBE Bau','OBENAUF','Ing. Felix Novotny','DI Wilhelm Sedlak','BAU-WERTE','KALCON','Sanibau','HAZET','DI Remzi Avunduk','Frigo']:
        parts.append(f"### {r['firma']} — {r['serviciu']}\n\n**Dovadă: {r['data_ultimei_dovezi']} · {r['status']} · {r['pret_net']}.**\n\n{r['ultimul_raspuns']}\n\nUrmătorul pas: {r['pas_urmator']}\n\nSursă: "+' · '.join(link(s) for s in r['surse'])+f"\n\nCalitatea dovezii: {r['nivel_dovada']}.\n")
parts.append('## Corecții ale situației anterioare\n')
parts.extend(f'- **{a}:** {c}' for a,b,c in corrections)
parts.append('''

## Ce rămâne de verificat în Eva-Mail

- Caută atât office@ac-wohnart.at, cât și covaciu.gnm@gmail.com, după Schallergasse / 1120 / numele ofertanților și referințele ANG 855, 839, 840, 891, 32KB1294-01 și 2026-470/AN01. Citește firul complet pentru răspunsurile relevante.
- Verifică mesajele după 24.09 pentru constructori și lift, după 28.09 pentru TOMS, precum și orice revizie anterioară care nu este încă în arhivă.
- Recuperează emailul LEBE din 24.09, 17:38, mesajele Novotny din 24.09 și eventualele rezultate ale întâlnirilor 25.09.
- Salvează fiecare original în subfolderul firmei, cu data și referința ofertei. Înregistrează ID-ul mesajului, expeditorul, data, numele exact al atașamentului, dimensiunea și SHA-256. Un nume de atașament găsit într-un fir nu dovedește existența unei oferte completate.
- Compară hash-ul cu inventarul înainte de a crea o copie. Păstrează reviziile mai vechi ca istoric. Nu marca oferta drept acceptată doar pentru că există un chestionar semnat ori un proiect de contract.

În această actualizare nu s-au trimis mesaje către ofertanți și nu s-au modificat contractele. Subfolderele create conțin fișe de situație și trimiteri la documentele deja existente.
''')
report='\n\n'.join(parts)
(DEST/'Situatie oferte 2026-09-30.md').write_text(report,encoding='utf-8')
doc=Document();doc.add_heading('Situația ofertelor Schallergasse 35',0);doc.add_paragraph('30 septembrie 2026 · Surse locale · Actualizarea Eva-Mail rămâne deschisă')
for line in report.splitlines():
    if line.startswith('# '):continue
    if line.startswith('## '):doc.add_heading(line[3:],1)
    elif line.startswith('### '):doc.add_heading(line[4:],2)
    elif line.strip():doc.add_paragraph(re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',line).replace('**',''))
doc.save(DEST/'Situatie oferte 2026-09-30.docx')
for firm in ['TOMS','Otis','LEBE Bau','OBENAUF','Ing. Felix Novotny','DI Wilhelm Sedlak','BAU-WERTE']:
    fd=DEST/firm;fd.mkdir(exist_ok=True);texts=[f'# {firm} situație la 30 septembrie 2026','', 'Surse locale. Eva-Mail nu a fost verificat live.']
    for r in R:
        if r['firma']!=firm:continue
        texts.extend(['',f"## {r['serviciu']}",'',f"{r['data_ultimei_dovezi']} · {r['status']} · {r['pret_net']}",'',r['ultimul_raspuns'],'','Următorul pas: '+r['pas_urmator'],'','Surse: '+ ' · '.join(link(s) for s in r['surse'])])
    (fd/'Situatie si surse.md').write_text('\n'.join(texts),encoding='utf-8')
print(json.dumps({'registru':str(out),'inregistrari':len(R),'raport':str(DEST/'Situatie oferte 2026-09-30.md')},ensure_ascii=False))
