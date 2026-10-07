from pathlib import Path
import json,fitz,sys
from openpyxl import load_workbook
R=Path(__file__).resolve().parent.parent
D=R/'10. Banci + Extrase de cont'/'2026.09.30 Audit facturi si plati';D.mkdir(exist_ok=True)
mode=sys.argv[1] if len(sys.argv)>1 else 'bank'
if mode=='bank':
 for p in (R/'10. Banci + Extrase de cont').glob('Extrase*/*.pdf'):
  d=fitz.open(p);text='\n\n'.join('PAGINA '+str(i+1)+'\n'+page.get_text(sort=True) for i,page in enumerate(d))
  (D/('2026.09.30 Extras text '+('A&C' if 'AT22' in p.name else 'Cosmin')+'.txt')).write_text(text,encoding='utf-8')
  print('\nSURSA',p.relative_to(R),'PAGINI',len(d),'\n',text)
else:
 paths=list((R/'10. Banci + Extrase de cont').glob('*.xlsx'))+list((R/'08. Corespondenta').glob('De platit*.xlsx'))
 all=[]
 for p in paths:
  w=load_workbook(p,data_only=False);obj={'cale':p.relative_to(R).as_posix(),'foi':{}}
  print('\nWORKBOOK',p.name)
  for sh in w:
   rows=[[v.isoformat() if hasattr(v,'isoformat') else v for v in row] for row in sh.values]
   obj['foi'][sh.title]=rows
   print(sh.title,sh.max_row,sh.max_column)
   if len(sys.argv)>2 and sys.argv[2].lower() in p.name.lower():
    for i,row in enumerate(rows):print(i+1,json.dumps(row,ensure_ascii=False))
  all.append(obj)
 (D/'2026.09.30 Surse Excel.json').write_text(json.dumps(all,ensure_ascii=False,indent=2),encoding='utf-8')
