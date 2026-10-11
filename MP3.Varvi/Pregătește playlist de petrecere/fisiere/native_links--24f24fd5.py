from pathlib import Path
import json,zipfile,xml.etree.ElementTree as ET,collections
root=Path(__file__).parent;path=root/'Playlist_MP3_gratuite_24_ore.xlsx'
data=json.loads((root/'hyperlinks.json').read_text('utf8'))
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';rn='http://schemas.openxmlformats.org/officeDocument/2006/relationships';pn='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('',ns);ET.register_namespace('r',rn)
with zipfile.ZipFile(path) as z:files={n:z.read(n) for n in z.namelist()}
after={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
groups=collections.defaultdict(list)
for row in data:groups[row['sheet']].append(row)
for i,rows in groups.items():
 sheet=f'xl/worksheets/sheet{i}.xml';rel=f'xl/worksheets/_rels/sheet{i}.xml.rels'
 tree=ET.fromstring(files[sheet]);rels=ET.fromstring(files[rel]) if rel in files else ET.Element('{'+pn+'}Relationships');links=ET.Element('{'+ns+'}hyperlinks')
 for j,row in enumerate(rows):
  rid=f'rIdFreeMP3{j}';ET.SubElement(links,'{'+ns+'}hyperlink',{'ref':row['cell'],'{'+rn+'}id':rid});ET.SubElement(rels,'{'+pn+'}Relationship',{'Id':rid,'Type':rn+'/hyperlink','Target':row['url'],'TargetMode':'External'})
 pos=next((j for j,c in enumerate(tree) if c.tag.split('}')[-1] in after),len(tree));tree.insert(pos,links)
 files[sheet]=ET.tostring(tree,encoding='utf-8',xml_declaration=True);files[rel]=ET.tostring(rels,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in files.items():z.writestr(n,b)
from openpyxl import load_workbook
w=load_workbook(path,data_only=True)
assert sum(1 for s in w for row in s for c in row if c.hyperlink)==len(data)
assert w['Sinteză']['B19'].value==len(data)
assert w['Sinteză']['C19'].value.total_seconds()==86400
print('Validated',len(data),'links, 24 hours,',len(w.sheetnames),'tabs,',path.stat().st_size,'bytes')
