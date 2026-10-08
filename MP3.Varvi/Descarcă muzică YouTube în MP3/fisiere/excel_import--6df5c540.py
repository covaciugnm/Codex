"""Read link columns from XLSX without executing workbook formulas."""
import io
import re
import unicodedata
import zipfile

from openpyxl import load_workbook


def normalized(value):
    text = unicodedata.normalize('NFKD', str(value or '').strip().lower())
    return ' '.join(''.join(c for c in text if not unicodedata.combining(c)).split())


def read_excel(content, validate_url):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('Fișierul depășește limita de 10 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('Fișierul Excel este prea mare după decomprimare (maximum 40 MB).')
        workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Fișierul nu este un Excel .xlsx valid sau este protejat/parolat.') from exc
    entries, issues, sheets, seen = [], [], [], {}
    try:
        for sheet in workbook:
            if sheet.max_row > 20000 or sheet.max_column > 200:
                raise ValueError(f'Foaia „{sheet.title}” depășește limita de 20.000 rânduri / 200 coloane.')
            headers = None
            for row in sheet.iter_rows(max_row=min(50, sheet.max_row)):
                columns = {normalized(c.value): c.column for c in row if c.value is not None}
                link_column = next((columns[k] for k in ('link youtube', 'youtube', 'youtube url', 'url', 'link') if k in columns), None)
                if link_column:
                    headers = row[0].row, link_column
                    break
            if not headers:
                sheets.append({'name': sheet.title, 'count': 0, 'skipped': True})
                continue
            count = 0
            for row in sheet.iter_rows(min_row=headers[0] + 1):
                cell = row[headers[1] - 1]
                values = {str(sheet.cell(headers[0], c.column).value): str(c.value) for c in row if c.value is not None}
                raw = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else cell.value
                if raw is None or not str(raw).strip():
                    if any(c.value is not None and str(c.value).strip() for c in row):
                        issues.append({'sheet': sheet.title, 'row': cell.row, 'reason': 'Link YouTube lipsă.'})
                    continue
                raw = str(raw).strip()
                if raw.startswith('='):
                    match = re.fullmatch(r'=\s*HYPERLINK\(\s*"((?:[^"]|"")*)"\s*(?:[,;].*)?\)', raw, re.I)
                    if not match:
                        issues.append({'sheet': sheet.title, 'row': cell.row, 'reason': 'Formula nu conține un link HYPERLINK literal. Lipește URL-ul direct.'})
                        continue
                    raw = match[1].replace('""', '"')
                try:
                    url = validate_url(raw)
                except ValueError as exc:
                    issues.append({'sheet': sheet.title, 'row': cell.row, 'reason': str(exc)})
                    continue
                if url in seen:
                    issues.append({'sheet': sheet.title, 'row': cell.row, 'reason': f'Duplicat omis; prima apariție: {seen[url]}.'})
                    continue
                seen[url] = f'{sheet.title}, rândul {cell.row}'
                entries.append({'url': url, 'sheet': sheet.title, 'row': cell.row, 'values': values})
                count += 1
                if len(entries) > 2000:
                    raise ValueError('Maximum 2.000 de linkuri distincte per import.')
            sheets.append({'name': sheet.title, 'count': count, 'skipped': False})
    finally:
        workbook.close()
    return {'entries': entries, 'issues': issues, 'sheets': sheets}
