import io
import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
from contextlib import nullcontext
from openpyxl import Workbook
import app
from excel_import import read_excel


def fixture_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Jazz & lounge'
    sheet.append(['Model'])
    sheet.cell(7, 1, 'Nume artist')
    sheet.cell(7, 7, 'Link YouTube')
    sheet.cell(8, 1, 'Artist')
    sheet.cell(8, 7, 'https://youtu.be/abcdefghijk')
    sheet.cell(9, 7, 'Deschide').hyperlink = 'https://www.youtube.com/watch?v=12345678901'
    sheet.cell(10, 7, '=HYPERLINK("https://youtu.be/zyxwvutsrqp","Video")')
    sheet.cell(11, 7, 'https://www.youtube.com/watch?v=abcdefghijk&t=20')
    sheet.cell(12, 7, 'https://example.com/video')
    sheet.cell(13, 1, 'Fără link')
    sheet.cell(14, 7, '=HYPERLINK(A1,"Video")')
    other = book.create_sheet('Românești')
    other.append(['Link YouTube'])
    other.append(['https://www.youtube.com/playlist?list=PL123456789012'])
    book.create_sheet('Sinteză').append(['Sumar'])
    buffer = io.BytesIO(); book.save(buffer)
    return buffer.getvalue()


class ImportTests(unittest.TestCase):
    def tearDown(self):
        app.jobs.clear(); app.imports.clear()

    def test_hyperlinks_formulas_duplicates_and_row_numbers(self):
        result = read_excel(fixture_bytes(), app.youtube_url)
        self.assertEqual(len(result['entries']), 4)
        self.assertEqual([r['row'] for r in result['entries']], [8, 9, 10, 2])
        self.assertEqual(len(result['issues']), 4)
        self.assertTrue(result['sheets'][-1]['skipped'])
        self.assertEqual(result['entries'][0]['values']['Nume artist'], 'Artist')

    def test_user_model_is_empty_and_recognized(self):
        model = app.ROOT / 'outputs/party_20261008/Model_playlist_masa_si_petrecere.xlsx'
        if not model.exists():
            self.skipTest('Model local indisponibil')
        result = read_excel(model.read_bytes(), app.youtube_url)
        self.assertEqual(len(result['entries']), 0)
        self.assertEqual(len(result['issues']), 0)
        self.assertEqual(sum(not sheet['skipped'] for sheet in result['sheets']), 10)

    def test_invalid_file(self):
        with self.assertRaises(ValueError):
            read_excel(b'not an xlsx', app.youtube_url)

    def test_preview_does_not_download_and_start_is_serial(self):
        client = app.app.test_client()
        headers = {'X-MP3-App': 'local'}
        with patch.object(app.pool, 'submit') as submit:
            response = client.post('/api/imports', data={'file': (io.BytesIO(fixture_bytes()), 'test.xlsx')}, headers=headers)
            self.assertEqual(response.status_code, 200)
            submit.assert_not_called()
            token = response.json['id']
            started = client.post(f'/api/imports/{token}/start', json={'quality': '192'}, headers=headers)
            self.assertEqual(started.status_code, 202)
            submit.assert_called_once()
            blocked = client.post('/api/jobs', json={'url': 'https://youtu.be/abcdefghijk'}, headers=headers)
            self.assertEqual(blocked.status_code, 409)
            self.assertEqual(client.post(f'/api/imports/{token}/start', json={}, headers=headers).status_code, 404)

    def test_empty_import_cannot_start(self):
        app.imports['empty'] = {'entries': [], 'filename': 'empty.xlsx'}
        response = app.app.test_client().post('/api/imports/empty/start', json={}, headers={'X-MP3-App': 'local'})
        self.assertEqual(response.status_code, 400)

    def test_batch_continues_after_unavailable_link(self):
        imported = read_excel(fixture_bytes(), app.youtube_url)
        imported['filename'] = 'test.xlsx'
        with app.lock:
            job = app.new_job()
        # Each unavailable source fails before media download; all four are attempted.
        with tempfile.TemporaryDirectory() as tmp, patch.object(app, 'OUTPUT', Path(tmp)), patch.object(app, 'paced_request', return_value=nullcontext()), patch.object(app, 'YoutubeDL') as ydl:
            ydl.return_value.__enter__.return_value.extract_info.return_value = None
            app.run_import(job, imported, '192')
            self.assertEqual(ydl.return_value.__enter__.return_value.extract_info.call_count, 4)
        self.assertEqual(app.jobs[job]['status'], 'failed')
        self.assertEqual(len(app.jobs[job]['errors']), 4)
        self.assertIn('rând 8', app.jobs[job]['errors'][0]['source'])

    def test_stable_folders_without_summary(self):
        imported = read_excel(fixture_bytes(), app.youtube_url)
        imported['filename'] = 'test.xlsx'
        with tempfile.TemporaryDirectory() as tmp, patch.object(app, 'OUTPUT', Path(tmp)):
            first = app.import_folders(imported, create=True)
            self.assertEqual(first, app.import_folders(imported, create=True))
            self.assertEqual(len(first), 2)
            self.assertNotIn('Sinteză', first)
            self.assertTrue(all((Path(tmp) / p).is_dir() for p in first.values()))
