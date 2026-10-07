from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'C:\Users\User\Documents\Codex\2026-09-30\in')
OUT = ROOT / 'outputs'
NAME = 'Anexa 1 la factura nr. 1 din 29.09.2026'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.6), Cm(1.5)
s.left_margin = s.right_margin = Cm(1.8)
s.header_distance = s.footer_distance = Cm(.7)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.05
for sty in ['Heading 1', 'Heading 2']:
    doc.styles[sty].font.name = 'Arial'
    doc.styles[sty].font.color.rgb = RGBColor.from_string('254B43')
    doc.styles[sty].font.size = Pt(12)
    doc.styles[sty].paragraph_format.space_before = Pt(10)
    doc.styles[sty].paragraph_format.space_after = Pt(5)

def p(text='', bold=False, size=None):
    par = doc.add_paragraph()
    r = par.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    return par

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
        sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E8EFEC'); c._tc.get_or_add_tcPr().append(sh)
        for r in c.paragraphs[0].runs: r.bold = True
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = text
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            if widths: c.width = Cm(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.space_before = Pt(3)
                for run in para.runs: run.font.size = Pt(9)
    return t

header = s.header.paragraphs[0]
header.text = 'ÎNTREȚINERE BAZĂ SPORTIVĂ  |  ANEXĂ LA FACTURĂ'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('66746E')
foot = s.footer.paragraphs[0]
foot.alignment = 2
foot.add_run('Factura nr. 1 / 29.09.2026  •  Pagina ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)

p('ANEXA 1', True, 21)
p('la factura nr. 1 din 29.09.2026', True, 13)
p('Descrierea lucrărilor manuale de întreținere și calcul orientativ al suprafețelor', size=10)

doc.add_heading('1. Datele părților și ale facturii', 1)
table(['PRESTATOR / FURNIZOR', 'BENEFICIAR'], [[
    'REV0 5C SRL\nCIF: 52226710\nNr. registrul comerțului: J2025056106001\nSediu: Alba Iulia, str. Septimius Severus nr. 47, bl. TOR 3, ap. 13, jud. Alba',
    'ASOCIAȚIA CLUB SPORTIV ACADEMIA DE FOTBAL OPTIMUM ALBA IULIA\nCIF: 40804594\nSediu: Alba Iulia, Bld. Transilvaniei nr. 13, jud. Alba'
]], [8.7, 8.7])
p('Data facturii: 29.09.2026  |  Scadență: 29.09.2026  |  Monedă: RON', size=9)
table(['Poziția din factură', 'Cantitate', 'Valoare'], [[
    '1 – ÎNTREȚINERE CONFORM ANEXEI 1', '1', '30.000,00 lei'
]], [11.1, 2.3, 4])

doc.add_heading('2. Obiectul anexei', 1)
p('Întreținerea manuală a bazei sportive aflate în administrarea beneficiarului: teren de fotbal cu gazon natural, garduri, porți de acces, tribune, vestiare, alei și spații anexe.')
p('Operațiunile de mai jos reprezintă lucrări care se pot executa, în funcție de starea amplasamentului. Lucrările efectiv realizate, perioada și cantitățile se confirmă de părți la finalul anexei.', size=9)

doc.add_heading('3. Lucrări manuale de întreținere', 1)
table(['Zonă', 'Operațiuni posibile'], [
    ('Gazon – curățare', 'Strângerea deșeurilor, pietrelor și crengilor; greblarea frunzelor și a resturilor vegetale; plivirea manuală a buruienilor.'),
    ('Gazon – refaceri locale', 'Reașezarea și tasarea ușoară a brazdelor desprinse; completarea locală a denivelărilor cu amestec adecvat de nisip și pământ; supraînsămânțarea manuală a zonelor rare; udarea locală.'),
    ('Gazon – margini', 'Finisarea marginilor cu foarfecă manuală; aerarea punctuală cu furca în zonele compactate; distribuirea manuală a materialelor de întreținere, după necesitate.'),
    ('Garduri și acces', 'Îndepărtarea vegetației de la baza gardului; curățarea plasei și stâlpilor; strângerea prinderilor; refixarea locală a plasei; curățarea ruginii și retușuri cu pensula; ungerea balamalelor.'),
    ('Tribune', 'Măturarea treptelor și zonelor de acces; spălarea băncilor sau scaunelor; colectarea deșeurilor; îndepărtarea buruienilor; strângerea prinderilor accesibile; curățarea balustradelor, șlefuire manuală și retușuri de vopsire.'),
    ('Vestiare și anexe', 'Măturarea și spălarea pardoselilor; curățarea băncilor, dulapurilor, ușilor și geamurilor; igienizarea dușurilor și grupurilor sanitare; strângerea accesoriilor și retușuri locale de zugrăveală.'),
    ('Alei și dotări', 'Măturarea aleilor; plivirea rosturilor; colectarea și transportul manual al resturilor la punctul de colectare; curățarea băncilor de rezerve și a porților de joc; refixarea locală a plaselor.')
], [3.4, 14])

doc.add_page_break()
doc.add_heading('4. Calcul orientativ – terenul de fotbal', 1)
p('Ipoteză de calcul: teren dreptunghiular cu lungimea de 105 m și lățimea de 68 m, dimensiuni recomandate de FIFA. Dimensiunile reale ale bazei sportive nu sunt precizate în factura consultată.')
table(['Element', 'Calcul', 'Rezultat orientativ'], [
    ('Suprafața terenului de joc', '105 m × 68 m', '7.140 m²'),
    ('Suprafața în hectare', '7.140 m² ÷ 10.000', '0,714 ha'),
    ('Perimetrul terenului de joc', '2 × (105 m + 68 m)', '346 m liniari'),
    ('Exemplu: bandă exterioară de 2 m', '(105 + 4) × (68 + 4)', '7.848 m² în total'),
    ('Suprafața benzii de 2 m', '7.848 − 7.140', '708 m²'),
    ('Conturul exterior al benzii', '2 × (109 m + 72 m)', '362 m liniari')
], [6.1, 6.3, 5])
p('Banda de 2 m este un exemplu geometric pentru estimarea întreținerii. Lățimea efectivă se stabilește prin măsurare.', size=9)
p('Suprafața întregului stadion se determină separat: teren de joc + zone perimetrale + alei + amprenta vestiarelor și a celorlalte construcții + tribune și alte suprafețe, după caz. Suprafața de 7.140 m² se referă numai la terenul de joc.', size=9)

doc.add_heading('5. Calculul lucrărilor la garduri și vestiare', 1)
table(['Element de măsurat', 'Mod de calcul / exemplu'], [
    ('Lungimea împrejmuirii', 'Se măsoară traseul real al gardului. Cei 346 m reprezintă conturul terenului de joc; nu confirmă lungimea gardului existent.'),
    ('Suprafața geometrică a gardului', 'Lungime gard × înălțime. Exemplu ipotetic: 346 m × 2 m = 692 m² pe o față; 1.384 m² pentru două fețe. La gardul din plasă, consumul de vopsea depinde de tipul plasei și de stâlpi.'),
    ('Pardoseli vestiare', 'Pentru fiecare încăpere: lungime × lățime; se însumează suprafețele încăperilor.'),
    ('Pereți pentru retușuri / zugrăveli', 'Perimetrul încăperii × înălțimea, minus suprafața ușilor și ferestrelor; se măsoară separat zonele efectiv lucrate.')
], [6.1, 11.3])

doc.add_heading('6. Completarea lucrărilor efective', 1)
p('Amplasamentul bazei sportive: ........................................................................................')
p('Perioada executării: .........................................................................................................')
p('Lucrări executate și cantități măsurate: ...........................................................................')
p('.........................................................................................................................................')
p('.........................................................................................................................................')
p('Valoarea poziției de întreținere din factura nr. 1: 30.000,00 lei. Factura prezintă o singură poziție; nu conține o defalcare a valorii pe operațiuni sau tarife pe m² / metru liniar.', size=9)
p('Calculele orientative și lista operațiunilor posibile nu atestă executarea acestora. Confirmarea privește lucrările efective completate mai sus.', size=9)
table(['PRESTATOR', 'BENEFICIAR'], [[
    'Reprezentant: .....................................\nSemnătură: ..........................................\nData: ....................................................',
    'Reprezentant: .....................................\nSemnătură: ..........................................\nData: ....................................................'
]], [8.7, 8.7])
p('Referință pentru dimensiuni: FIFA – Stadium Guidelines, secțiunea 5.3, Pitch dimensions and surrounding areas. https://inside.fifa.com/innovation/stadium-guidelines/technical-guidelines/stadiums-guidelines/pitch-dimensions-and-surrounding-areas', size=7)
doc.save(OUT / (NAME + '.docx'))
print(OUT / (NAME + '.docx'))
