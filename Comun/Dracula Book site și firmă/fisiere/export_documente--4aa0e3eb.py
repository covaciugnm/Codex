from pathlib import Path
import sys,re,html,json
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,LongTable,TableStyle
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
ROOT=Path(__file__).resolve().parent.parent
files=['00_CONDUCERE/MANUAL_ATELIER.md','00_CONDUCERE/DECIZIE_EDITORIALA_001_SURSE.md','00_CONDUCERE/TRASABILITATE.md','00_CONDUCERE/ROADMAP.md','01_ECHIPA/00_ECHIPA_SI_RESPONSABILITATI.md','00_CONDUCERE/RUBRICI.md','00_CONDUCERE/PROTOCOL_ARHIVARE.md']
def clean(s):
    s=re.sub(r'(?<=[a-zăâîșț])(?=\d)', ' ',s)
    s=re.sub(r'(?<=\d)(?=[a-jl-zăâîșț])', ' ',s)
    return s.replace('**','').replace(chr(96),'')
def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch('[-: ]+',x) for x in row):rows.append(row)
                i+=1
            yield 'table',rows;continue
        if s.startswith('#'):
            n=len(s)-len(s.lstrip('#'));yield 'heading',(n,clean(s[n:].strip()))
        else:yield 'paragraph',clean(s)
        i+=1
FONT=Path(r'C:\Windows\Fonts')
for n,f in [('DB','arial.ttf'),('DBB','arialbd.ttf')]:pdfmetrics.registerFont(TTFont(n,str(FONT/f)))
burg=colors.HexColor('#681E35');ink=colors.HexColor('#202635');gold=colors.HexColor('#A97832')
styles=getSampleStyleSheet()
for n,size,lead,font,color in [('BodyDB',9.1,13,'DB',ink),('H1DB',18,23,'DBB',burg),('H2DB',13,17,'DBB',burg),('H3DB',10.5,14,'DBB',ink),('CellDB',7.6,10,'DB',ink),('HeadDB',7.8,10,'DBB',colors.white)]:
    styles.add(ParagraphStyle(name=n,fontName=font,fontSize=size,leading=lead,textColor=color,spaceAfter=7 if n=='BodyDB' else 9,spaceBefore=10 if n.startswith('H') else 0,keepWithNext=n.startswith('H'),splitLongWords=True,allowWidows=0,allowOrphans=0))
styles.add(ParagraphStyle(name='BodyCompactDB',parent=styles['BodyDB'],spaceAfter=4,leading=12.5))
def p(s,style='BodyDB'):return Paragraph(html.escape(str(s)),styles[style])
def widths(n):return [174*mm*x for x in ({2:[.25,.75],3:[.22,.18,.60],4:[.20,.44,.14,.22],5:[.06,.22,.31,.18,.23]}.get(n,[1/n]*n))]
def footer(c,d):
    c.saveState();c.setStrokeColor(gold);c.line(18*mm,18*mm,A4[0]-18*mm,18*mm)
    c.setFont('DB',7);c.setFillColor(ink);c.drawString(18*mm,13*mm,'DRACULA BOOK | Atelier editorial | 24.09.2026')
    c.drawRightString(A4[0]-18*mm,13*mm,str(d.page));c.restoreState()
flow=[Spacer(1,25*mm),p('DRACULA BOOK','H1DB'),Spacer(1,12*mm),p('ATELIER EDITORIAL','H1DB'),p('Romane originale de minimum 50.000 de cuvinte','H2DB'),Spacer(1,10*mm),p('Echipă de agenți • Roadmap • Audit independent • Remediere • Arhivare','H2DB'),Spacer(1,18*mm),p('Prag strict: fiecare criteriu >9,50/10, doi auditori separați și metaaudit. Nu se publică și nu se avansează pe baza unei note declarate de autor.'),p('Acesta este manualul sistemului de lucru, nu un roman finalizat. Starea verificată și rapoartele reale sunt păstrate separat în registru și arhivă.'),p('Fișele individuale ale celor 33 de roluri sunt în directorul 01_ECHIPA. Modelele editabile sunt în 03_MODELE.')]
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sty=doc.styles['Normal'];sty.font.name='Calibri';sty.font.size=Pt(10);sty.paragraph_format.space_after=Pt(7)
for name,size in [('Title',28),('Heading 1',16),('Heading 2',12)]:
    s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('681E35')
doc.add_paragraph('DRACULA BOOK','Subtitle');doc.add_paragraph('Atelier editorial','Title');doc.add_paragraph('Romane originale de minimum 50.000 de cuvinte','Subtitle')
doc.add_paragraph('Echipă de agenți • Roadmap • Audit independent • Remediere • Arhivare')
doc.add_paragraph('Manual de lucru — 24.09.2026. Starea verificată și rapoartele reale sunt separate în registru. Nu certifică finalizarea unui roman.')
fo=sec.footer.paragraphs[0];fo.text='Dracula Book | Atelier editorial | ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fo._p.append(field)
for f in files:
    flow.append(PageBreak());doc.add_page_break()
    for kind,v in blocks((ROOT/f).read_text(encoding='utf-8')):
        if kind=='heading':
            n,t=v;flow.append(p(t,'H1DB' if n==1 else 'H2DB' if n==2 else 'H3DB'));doc.add_heading(t,1 if n<=2 else 2)
        elif kind=='paragraph':
            compact=f.endswith('DECIZIE_EDITORIALA_001_SURSE.md')
            flow.append(p(v,'BodyCompactDB' if compact else 'BodyDB'))
            para=doc.add_paragraph(v)
            if compact:para.paragraph_format.space_after=Pt(4)
        else:
            cells=[[p(x,'HeadDB' if i==0 else 'CellDB') for x in row] for i,row in enumerate(v)]
            tab=LongTable(cells,colWidths=widths(len(v[0])),repeatRows=1,hAlign='LEFT',spaceAfter=10)
            tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),burg),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F6F1EE')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));flow.append(tab)
            t=doc.add_table(rows=1,cols=len(v[0]));t.autofit=False;t.style='Light Shading Accent 1'
            for j,x in enumerate(v[0]):t.rows[0].cells[j].text=x
            rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
            for row in v[1:]:
                c=t.add_row().cells
                for j,x in enumerate(row):c[j].text=x
            for i,row in enumerate(t.rows):
                for j,c in enumerate(row.cells):
                    c.width=Cm(widths(len(v[0]))[j]/mm/10)
                    for para in c.paragraphs:
                        para.paragraph_format.space_after=Pt(3)
                        for run in para.runs:run.font.size=Pt(8);run.bold=(i==0)
            doc.add_paragraph()
pdf=ROOT/'Manual_operational_atelier_Dracula_Book.pdf'
SimpleDocTemplate(str(pdf),pagesize=A4,leftMargin=18*mm,rightMargin=18*mm,topMargin=20*mm,bottomMargin=24*mm,title='Dracula Book — Atelier editorial',author='Dracula Book — document de lucru').build(flow,onFirstPage=footer,onLaterPages=footer)
doc.save(ROOT/'Manual_operational_atelier_Dracula_Book.docx')
print(json.dumps({'pdf':str(pdf),'docx':str(ROOT/'Manual_operational_atelier_Dracula_Book.docx')},ensure_ascii=False))
