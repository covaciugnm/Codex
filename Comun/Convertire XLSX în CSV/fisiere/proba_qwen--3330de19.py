# Proba traducere EVA Learn cu Qwen local (Ollama) + glosar din dictionarele EVA_Learn_Import_35
import json, random, sqlite3, sys, time, urllib.request, collections
import openpyxl

OLLAMA = "http://192.168.100.160:11436/api/chat"
MODEL = "qwen3.8:27b-q4_K_M"
XLSX = r"Z:\00.Roboti\EVA.Pro\Learn\EVA-traducere-RO-TOT-2026-10-03.xlsx"
DB_RO = r"Z:\00.Roboti\EVA.Pro\Learn\EVA_Learn_Import_35\05_Romana_ro\DB\eva_ro.sqlite"
DB_EN = r"Z:\00.Roboti\EVA.Pro\Learn\EVA_Learn_Import_35\01_Engleza_en\DB\eva_en.sqlite"

LIMBI = {"de": "German (Deutsch)", "bn": "Bengali (বাংলা)", "vi": "Vietnamese (Tiếng Việt)",
         "ha": "Hausa", "yue": "Cantonese (粵語, Traditional Chinese characters)"}

SYSTEM = """You are a professional translator for EVA Learn, an app that teaches ENGLISH to speakers of {lang}.
The app was written for Romanian speakers; you localize it for {lang} speakers.

You receive a JSON list of segments. Each has: id, src ("ro" or "en"), text, ctx (where it appears), and optionally gloss (dictionary hints).
Translate each "text" into {lang}.

RULES
1. src="ro": translate the Romanian into natural {lang}. Any ENGLISH words/sentences inside the Romanian text are the material being taught: keep them EXACTLY as they are, untranslated.
   Example: "E-mail formal: The manager ___ I sent the report has already replied." -> only "E-mail formal:" is translated.
   Example: "„Mistake” = greșeală." -> keep "Mistake", translate only "greșeală".
2. src="en": give the {lang} version of the English text (hint, card, test item) so the learner understands it.
3. Keep unchanged: "/", "___", "|", emoji, the name EVA, digits, ${{...}} placeholders, HTML tags like <b> <br>, markdown (**bold**, `code`), line breaks.
4. If the Romanian text explains something specific to Romanian speakers (e.g. a Romanian word, Romanian pronunciation), translate it faithfully anyway; do not invent new content.
5. The tutor EVA is female; use a friendly, simple tone, informal "you" where natural in {lang}.
6. "gloss" are dictionary suggestions; use them only if they fit the context.
7. Output ONLY JSON: {{"t":[{{"id":<id>,"tr":"<translation>"}}, ...]}} with exactly one item per input segment, same ids.
"""

def lookup(db, word, lang, n=3):
    rows = db.execute("select text,count(*) c from entries where word=? and target_lang=? and kind='translation' "
                      "group by text order by c desc limit ?", (word, lang, n)).fetchall()
    return [r[0] for r in rows]

def load_sample(k_per_sheet=4, seed=7):
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    random.seed(seed)
    segs = []
    for ws in wb.worksheets[1:]:
        rows = list(ws.iter_rows(values_only=True))
        h = [str(x) for x in rows[0]]
        iro = h.index("Text RO")
        ien = h.index("Text EN") if "Text EN" in h else h.index("Text EN (de învățat)")
        ictx = h.index("Câmp") if "Câmp" in h else (h.index("Câmp EN") if "Câmp EN" in h else h.index("Tip / id"))
        cand = [r for r in rows[1:] if r[iro] or r[ien]]
        for r in random.sample(cand, min(k_per_sheet, len(cand))):
            if r[iro]:
                segs.append({"src": "ro", "text": str(r[iro]), "ctx": f"{ws.title} / {r[ictx]}", "en": r[ien]})
            else:
                segs.append({"src": "en", "text": str(r[ien]), "ctx": f"{ws.title} / {r[ictx]} / {r[-1]}"})
    for i, s in enumerate(segs):
        s["id"] = i
    return segs

def add_gloss(segs, lang):
    ro, en = sqlite3.connect(f"file:{DB_RO}?mode=ro", uri=True), sqlite3.connect(f"file:{DB_EN}?mode=ro", uri=True)
    for s in segs:
        g = {}
        if len(s["text"].split()) <= 3:
            db = ro if s["src"] == "ro" else en
            hit = lookup(db, s["text"].strip().lower(), lang)
            if hit: g[s["text"]] = hit
            if s["src"] == "ro" and s.get("en") and len(str(s["en"]).split()) <= 3:
                hit = lookup(en, str(s["en"]).strip(), lang)
                if hit: g[str(s["en"])] = hit
        if g: s["gloss"] = g

def ask(lang, batch):
    payload = [{k: s[k] for k in ("id", "src", "text", "ctx", "gloss") if k in s} for s in batch]
    body = {"model": MODEL, "stream": False, "think": False, "format": "json",
            "messages": [{"role": "system", "content": SYSTEM.format(lang=LIMBI[lang])},
                         {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}],
            "options": {"temperature": 0.2, "num_predict": 8000, "num_ctx": 16384}}
    req = urllib.request.Request(OLLAMA, json.dumps(body).encode(), {"Content-Type": "application/json"})
    t = time.time()
    r = json.load(urllib.request.urlopen(req, timeout=1800))
    dt = time.time() - t
    return json.loads(r["message"]["content"]), dt, r.get("prompt_eval_count"), r.get("eval_count"), r.get("eval_duration", 0) / 1e9

if __name__ == "__main__":
    langs = sys.argv[1:] or ["de", "bn"]
    segs = load_sample()
    out = {}
    for lang in langs:
        add_gloss(segs, lang)
        res, dt, pin, pout, evs = ask(lang, segs)
        tr = {x["id"]: x["tr"] for x in res.get("t", [])}
        print(f"== {lang}: {len(segs)} segmente, {dt:.0f}s, tokeni in {pin} out {pout}, {pout/max(evs,1e-9):.1f} tok/s, primite {len(tr)}")
        out[lang] = [{**s, "tr": tr.get(s["id"])} for s in segs]
        for s in segs: s.pop("gloss", None)
    json.dump(out, open(r"Z:\00.Roboti\EVA.Pro\Learn\Traducator\proba_rezultat.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
