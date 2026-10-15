import docx, glob, os, shutil, re
SRC = r"Z:\00. Proiecte 2025\2025.06.12 Ionel Capota SES\3. Achizitie\2. Revo 5C SRL\Achizitie veche"
DST = r"Z:\00. Proiecte 2025\2025.06.12 Ionel Capota SES\3. Achizitie\2. Revo 5C SRL\REVO 5C -CFMoto"
REPL = [
 ("SC DANCOR-PROIECT SRL", "SC REVO 5C S.R.L."),
 ("DANCOR-PROIECT SRL", "REVO 5C S.R.L."),
 ("DANCOR-PROIECT", "REVO 5C"),
 ("DANCOR", "REVO 5C"),
 ("Municipiul Alba Iulia, Str. ȘTEFAN LUCHIAN, Nr. 3B, Județ Alba", "Municipiul Alba Iulia, Str. SEPTIMIUS SEVERUS, Nr. 47, Bl. TOR 3, Et. 3, Ap. 13, Județ Alba"),
 ("Municipiul Alba Iulia, Str. ȘTEFAN LUCHIAN, nr. 3B, județul Alba", "Municipiul Alba Iulia, Str. SEPTIMIUS SEVERUS, nr. 47, bl. TOR 3, et. 3, ap. 13, județul Alba"),
 ("EUID ROONRC.J01/410/1999", "EUID ROONRC.J2025056106001"),
 ("J01/410/1999", "J2025056106001"),
 ("12445967", "52226710"),
 ("telefon 058/822200; 0742749605", "telefon ______________"),
 ("dl COVACIU COSMIN ADRIAN", "dna GIURGIU IULIA-ELISABETA"),
 ("COVACIU COSMIN ADRIAN", "GIURGIU IULIA-ELISABETA"),
]
def fix_par(p):
    n=0
    for r in p.runs:
        t=r.text
        for a,b in REPL:
            if a in t: t=t.replace(a,b); n+=1
        if t!=r.text: r.text=t
    # fallback: text split across runs
    full=p.text
    if any(a in full for a,_ in REPL):
        for a,b in REPL: full=full.replace(a,b)
        if p.runs:
            p.runs[0].text=full
            for r in p.runs[1:]: r.text=""
            n+=1
    return n
def walk(d):
    n=0
    for p in d.paragraphs: n+=fix_par(p)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: n+=fix_par(p)
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            for p in part.paragraphs: n+=fix_par(p)
            for t in part.tables:
                for row in t.rows:
                    for c in row.cells:
                        for p in c.paragraphs: n+=fix_par(p)
    return n
for f in sorted(os.listdir(SRC)):
    out = f.replace("DANCOR_PROIECT_SRL","REVO_5C_SRL").replace("DANCOR-PROIECT_SRL","REVO_5C_SRL").replace("DANCOR","REVO_5C")
    if f.endswith(".docx"):
        d=docx.Document(os.path.join(SRC,f)); n=walk(d); d.save(os.path.join(DST,out)); print(f"{n:3d} repl  {out}")
    elif f.endswith(".pdf") and "DANCOR" not in f:
        shutil.copy2(os.path.join(SRC,f), os.path.join(DST,out)); print("copy      ", out)
    else: print("skip      ", f)
